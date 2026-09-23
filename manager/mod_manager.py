#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
FFXIV Mod 管理工具
=====================================================================
扫描 Mod 目录 -> 建立本地索引库(SQLite) -> 一键导出 Excel
（每个大分类一个工作表，预览图内嵌，Excel / WPS 都能直接显示）

命令行用法：
    python mod_manager.py            打开图形界面（默认）
    python mod_manager.py scan       只扫描目录，更新索引库
    python mod_manager.py export     用索引库生成 Excel
    python mod_manager.py run        扫描 + 导出（最常用）
    python mod_manager.py check      检查：缺地址 / 缺预览图 / 序号重复 等
    python mod_manager.py stats      打印统计
    python mod_manager.py backup     打包备份（Mod 文件夹 + 索引库 + 汇总表 -> 一个 zip）
    python mod_manager.py backups    列出已有的备份
    python mod_manager.py restore --src "<备份.zip>" [--mode skip|overwrite|rename]

整理规则（与手工整理时完全一致）：
    分类      = Mod 根目录下的一级目录（衣服 / 饰品 / 武器 / 皮肤 ...）
    序号      = 文件夹名前缀数字，如 "1.[Arte] Neolithe Bodystocking" -> 1
    作者      = 文件夹名中括号内的内容
    Mod名称   = 中括号后面的名称（去掉开头多余的 "- "）
    NSFW/SFW  = 路径中出现的 SFW / NSFW 目录名；没有则默认 SFW（如 武器\双蛇\...）
    Mod地址   = 文件夹内 地址.txt 的第一行；找不到就在子文件夹里递归找
    预览图    = 优先"与文件夹同名"的图片（同级目录 > 文件夹内），
                其次 [作者] 名称.jpg / 名称.jpg，最后才用文件夹里任意图片
"""
from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import time
import traceback
import urllib.parse
import urllib.request
import zipfile
from html import unescape
from pathlib import Path

from app_version import APP_VERSION                 # 版本号唯一来源（同目录 app_version.py）

try:                      # 跟随控制台编码输出中文，避免乱码/报错
    sys.stdout.reconfigure(errors="replace")
    sys.stderr.reconfigure(errors="replace")
except Exception:
    pass

# 管理器版本号见同目录 app_version.py（APP_VERSION），这里不再单独定义

if getattr(sys, "frozen", False):        # PyInstaller 打包后，程序目录 = exe 所在目录
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent
CONFIG_PATH = APP_DIR / "mod_manager.json"
DB_PATH = APP_DIR / "mod_manager.db"
THUMB_DIR = APP_DIR / ".thumb_cache"
BACKUP_DIR = APP_DIR / "backup"
LOG_PATH = APP_DIR / "mod_manager.log"

DEFAULT_CONFIG = {
    "root": "",                 # Mod 根目录（首次运行会让你选）
    "excel": "",                # 汇总表输出（留空=自动放在 Mod 根目录旁边）
    "category_order": ["衣服", "饰品", "武器", "皮肤"],
    "thumb_width": 800,          # 缩略图宽度(像素)；原图比它窄就不放大
    "embed_images": True,        # 是否把预览图内嵌进 Excel
    "preview_px": [178, 100],    # Excel 里预览图的显示尺寸(像素)
    "autofilter": False,         # 表头是否加筛选按钮
    "download_dir": "",          # 浏览器下载目录（留空=自动用 用户\Downloads）
    "inbox_dir": "",             # 暂存目录（留空=自动用 <Mod根目录>\..\_待导入）
    "browser_port": 9333,        # 内置浏览器调试端口
    "browser_path": "",          # 浏览器 exe（留空=自动探测，优先 CentBrowser/Chrome/Edge）
    "browser_dir": "",           # 内置浏览器配置目录（留空=程序目录\browser_profile）
    "install_dir": "",           # Mod 安装目录（Penumbra 的 Mods 目录，留空=自动探测）
    "backup_dir": "",            # 备份存放目录（留空=程序目录\backup）
}

MOD_RE = re.compile(r"^\s*(\d+)\s*\.")
NAME_RE = re.compile(r"\s*\[([^\]]*)\]\s*(.*)$")


def split_author_name(text: str):
    """从 '[作者] 名称' 里拆出 (作者, 名称)。
    关键：**按括号配对**找结尾，而不是遇到第一个 ']' 就停 ——
    作者名本身可能带方括号，例如 '[Scherana [The Moon] 3D] 某 Mod'。
    括号不配对时视为"没有作者"，整串当名称。"""
    s = (text or "").strip()
    if not s.startswith("["):
        return "", s
    depth = 0
    for i, ch in enumerate(s):
        if ch == "[":
            depth += 1
        elif ch == "]":
            depth -= 1
            if depth == 0:
                return s[1:i].strip(), s[i + 1:].strip()
    return "", s
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif")
ADDR_NAME = "地址.txt"
HEADERS = ["序号", "作者", "NSFW/SFW", "影响/替换", "更新时间", "Mod名称", "Mod地址", "预览图"]
COL_WIDTHS = [7.7, 15.9, 15.9, 30.0, 17.0, 48.0, 63.0, 28.0]
PREVIEW_COL = 8                  # 预览图所在列(H)
PREVIEW_COL_PX = 201
ROW_HEIGHT_PT = 80
EMU_PER_PX = 9525

SCHEMA = """
CREATE TABLE IF NOT EXISTS mods (
    folder      TEXT PRIMARY KEY,
    category    TEXT NOT NULL,
    seq         INTEGER,
    author      TEXT,
    nsfw        TEXT,
    subcat      TEXT,
    name        TEXT,
    addr        TEXT,
    addr_source TEXT,
    img         TEXT,
    img_source  TEXT,
    img_hash    TEXT,
    file_mtime  REAL,
    file_size   INTEGER,
    updated_at  TEXT,
    affects     TEXT,
    -- ---- 云存储（夸克归档，2026-09 加）----
    cloud_backend TEXT,       -- 'quark' / ''（没启用）
    cloud_path    TEXT,       -- 云端目录（相对 root 的镜像路径，如 皮肤/NSFW/纹身/7.xxx）
    cloud_state   TEXT,       -- ''(未归档) / archived(本地无载荷、云端有) / uploading / missing(云端找不到)
    cloud_size    INTEGER,    -- 载荷合计字节
    cloud_synced  TEXT,       -- 最近一次归档/校验时间
    payload_mtime REAL,       -- 归档前的载荷最新时间（顶替「本地文件时间」当更新基线）
    -- ---- 站点元信息 / 本地描述（2026-09 加）----
    races       TEXT,         -- 站点上的 Races（种族）
    genders     TEXT,         -- 站点上的 Genders（性别）
    released    TEXT,         -- 站点上的 Original Release Date（首发日期）
    desc        TEXT,         -- Mod 内容描述（站点有就带过来，也能自己编）
    -- ---- 「检查更新」用（2026-09 加）----
    site_updated TEXT,        -- 站点上 Last Version Update（= 我们这次下载/更新时的版本时间，基线）
    site_latest  TEXT,        -- 最近一次检查时站点上的最后更新时间
    site_version TEXT,        -- 站点上的最新版本号（版本历史里的 version_new）
    site_checked TEXT,        -- 最近一次检查时间
    update_avail INTEGER      -- 1 = 站点上有比本地新的文件
);
CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT);
-- 标签单独一张表：按 folder 关联，重扫描/upsert 都不可能冲掉它
CREATE TABLE IF NOT EXISTS mod_tags (
    folder TEXT NOT NULL,
    tag    TEXT NOT NULL,
    PRIMARY KEY (folder, tag)
);
CREATE INDEX IF NOT EXISTS idx_mod_tags_tag ON mod_tags(tag);
-- 载荷清单：一条 mod 可能有几十个包、多层子目录（实测最多 27 个），逐文件记账
CREATE TABLE IF NOT EXISTS payload_files (
    folder    TEXT NOT NULL,   -- mod 文件夹（与 mods.folder 同一个键）
    rel_path  TEXT NOT NULL,   -- 相对 mod 文件夹的路径（含子目录）
    size      INTEGER,
    md5       TEXT,
    sha1      TEXT,
    cloud_fid TEXT,
    state     TEXT,            -- local(本地有) / archived(已上传、本地已删) / missing
    checked   TEXT,
    PRIMARY KEY (folder, rel_path)
);
"""


# --------------------------------------------------------------------- basic
def _excepthook(exc_type, exc, tb):
    text = "".join(traceback.format_exception(exc_type, exc, tb))
    log(text)
    try:
        import ctypes
        ctypes.windll.user32.MessageBoxW(None, text[-1800:], "Mod 管理工具 - 出错", 0x10)
    except Exception:
        pass


sys.excepthook = _excepthook


def log(msg=""):
    line = str(msg)
    print(line)
    try:
        with LOG_PATH.open("a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except Exception:
        pass


def load_config() -> dict:
    cfg = dict(DEFAULT_CONFIG)
    if CONFIG_PATH.exists():
        try:
            cfg.update(json.loads(CONFIG_PATH.read_text(encoding="utf-8")))
        except Exception:
            pass
    return cfg


def save_config(cfg: dict):
    CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")


def md5_of(path: Path) -> str:
    h = hashlib.md5()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def open_path(path):
    try:
        os.startfile(str(path))          # noqa: S606  (Windows)
    except Exception:
        subprocess.Popen(["explorer", os.path.normpath(str(path))])


# ------------------------------------------------------------------- scanning
def find_addr(folder: Path):
    """返回 (地址, 来源说明)"""
    p = folder / ADDR_NAME
    if p.is_file():
        try:
            lines = [l.strip() for l in p.read_text(encoding="utf-8-sig", errors="replace").splitlines() if l.strip()]
            return (lines[0] if lines else ""), ADDR_NAME
        except Exception:
            return "", ADDR_NAME + "(读取失败)"
    for dirpath, dirnames, filenames in os.walk(folder):
        dirnames.sort()
        if ADDR_NAME in filenames:
            sub = Path(dirpath) / ADDR_NAME
            try:
                lines = [l.strip() for l in sub.read_text(encoding="utf-8-sig", errors="replace").splitlines() if l.strip()]
            except Exception:
                lines = []
            return (lines[0] if lines else ""), str(sub.relative_to(folder))
    return "", "(缺失)"


def _images_in(d: Path):
    try:
        return sorted(f for f in os.listdir(d)
                      if (d / f).is_file() and f.lower().endswith(IMG_EXT))
    except OSError:
        return []


def _pick_image(d: Path, base: str, author: str, name: str):
    files = _images_in(d)
    for want in (base, "[%s] %s" % (author, name) if author else name, name):
        for f in files:
            if os.path.splitext(f)[0] == want:
                return d / f
    for f in files:
        if name and name.lower() in f.lower():
            return d / f
    return (d / files[0]) if files else None


def find_image(folder: Path, base: str, author: str, name: str):
    """返回 (图片绝对路径, 说明)"""
    for ext in IMG_EXT:                      # 1) 同级目录里与文件夹同名的图
        p = folder.parent / (base + ext)
        if p.is_file():
            return p, "同级同名"
    p = _pick_image(folder, base, author, name)
    if p:
        return p, "文件夹内"
    for dirpath, dirnames, filenames in os.walk(folder):     # 2) 子文件夹
        dirnames.sort()
        p = _pick_image(Path(dirpath), base, author, name)
        if p:
            return p, "子文件夹"
    return None, "(缺失)"


def find_mod_dirs(root: Path) -> list:
    """遍历目录，找出所有 Mod 文件夹（只找路径，不读内容）"""
    out = []
    for dirpath, dirnames, filenames in os.walk(root):
        here = Path(dirpath)
        if here == root:
            continue
        if any(mp in here.parents for mp in out):            # 已是某个 Mod 文件夹的子目录
            dirnames[:] = []
            continue
        if MOD_RE.match(here.name):
            out.append(here)
    return out


# 注意：文件下面还有个老的 parse_mod_folder（返回元组，给 rename_mod 用），
# 所以这里必须换个名字，别撞车。
def parse_mod_record(root: Path, here: Path):
    """把一个 Mod 文件夹解析成记录；名字不以数字开头就返回 None"""
    base = here.name
    m = MOD_RE.match(base)
    if not m:
        return None
    rel = here.relative_to(root).parts
    category = rel[0]
    zone = "SFW"
    mid = []
    for part in rel[1:-1]:
        if part.upper() in ("SFW", "NSFW"):
            zone = part.upper()
        else:
            mid.append(part)
    subcat = os.sep.join(mid)
    rest = base[m.end():]
    author, _nm = split_author_name(rest)
    name = re.sub(r"^-\s*", "", _nm)
    addr, addr_src = find_addr(here)
    img, img_src = find_image(here, base, author, name)
    return {
        "folder": str(here), "rel": os.path.relpath(here, root),
        "category": category, "seq": int(m.group(1)), "author": author,
        "nsfw": zone, "subcat": subcat, "name": name, "addr": addr,
        "addr_source": addr_src,
        "img": str(img) if img else "", "img_source": img_src,
    }


def scan_root(root: Path) -> list:
    """遍历目录，解析出所有 Mod 记录"""
    return [r for r in (parse_mod_record(root, p) for p in find_mod_dirs(root)) if r]


def scan_root_progress(root: Path, progress=None, should_cancel=None) -> list:
    """结果和 scan_root 完全一样，但会逐个回调进度（进度条 / Web 版用）"""
    root = Path(root)
    dirs = find_mod_dirs(root)
    mods = []
    for i, p in enumerate(dirs, 1):
        if should_cancel and should_cancel():
            raise BackupCancelled()
        rec = parse_mod_record(root, p)
        if rec:
            mods.append(rec)
        if progress:
            progress(i, len(dirs), p.name)
    return mods


def first_run_setup(win, cfg) -> bool:
    """首次使用：让用户指定 Mod 根目录 / Excel 输出（返回 False = 用户取消）"""
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    dlg = tk.Toplevel(win)
    dlg.title("首次使用：选择 Mod 目录")
    dlg.resizable(False, False)
    v_root = tk.StringVar(value=cfg.get("root") or "")
    v_excel = tk.StringVar(value=cfg.get("excel") or "")

    def pick_root():
        d = filedialog.askdirectory(parent=dlg, title="选择 Mod 根目录（里面是 衣服/饰品/武器/皮肤 这些分类）")
        if d:
            v_root.set(d)
            if not v_excel.get().strip():
                v_excel.set(str(Path(d).resolve().parent / "Mod信息汇总表.xlsx"))

    def pick_excel():
        f = filedialog.asksaveasfilename(parent=dlg, defaultextension=".xlsx",
                                         filetypes=[("Excel", "*.xlsx")])
        if f:
            v_excel.set(f)

    ttk.Label(dlg, text="欢迎使用 FFXIV Mod 管理工具",
              font=("Microsoft YaHei", 11, "bold")).grid(
        row=0, column=0, columnspan=3, padx=14, pady=(14, 2), sticky="w")
    ttk.Label(dlg, foreground="#666",
              text="第一次使用，请先指定 Mod 根目录（例如  ...\\MOD\\202609 ）。").grid(
        row=1, column=0, columnspan=3, padx=14, sticky="w")
    ttk.Label(dlg, text="Mod 根目录").grid(row=2, column=0, sticky="e", padx=8, pady=6)
    ttk.Entry(dlg, textvariable=v_root, width=58).grid(row=2, column=1, sticky="we", padx=6)
    ttk.Button(dlg, text="浏览…", command=pick_root).grid(row=2, column=2, padx=8)
    ttk.Label(dlg, text="Excel 输出").grid(row=3, column=0, sticky="e", padx=8, pady=6)
    ttk.Entry(dlg, textvariable=v_excel, width=58).grid(row=3, column=1, sticky="we", padx=6)
    ttk.Button(dlg, text="浏览…", command=pick_excel).grid(row=3, column=2, padx=8)
    dlg.columnconfigure(1, weight=1)
    res = {"ok": False}

    def ok():
        r = v_root.get().strip()
        if not r or not Path(r).is_dir():
            messagebox.showwarning("提示", "请选择一个存在的 Mod 根目录", parent=dlg)
            return
        cfg["root"] = r
        cfg["excel"] = (v_excel.get().strip()
                        or str(Path(r).resolve().parent / "Mod信息汇总表.xlsx"))
        save_config(cfg)
        res["ok"] = True
        dlg.destroy()

    bf = ttk.Frame(dlg)
    bf.grid(row=4, column=0, columnspan=3, pady=12)
    ttk.Button(bf, text="开始使用", width=12, command=ok).pack(side="left", padx=8)
    ttk.Button(bf, text="退出", width=12, command=dlg.destroy).pack(side="left", padx=8)
    dlg.bind("<Return>", lambda *_: ok())
    dlg.bind("<Escape>", lambda *_: dlg.destroy())
    dlg.update_idletasks()
    sw, sh = dlg.winfo_screenwidth(), dlg.winfo_screenheight()
    dlg.geometry("+%d+%d" % (max(8, (sw - dlg.winfo_reqwidth()) // 2),
                            max(8, (sh - dlg.winfo_reqheight()) // 3)))
    dlg.lift()
    try:
        dlg.focus_force()
    except Exception:
        pass
    dlg.wait_window()
    return res["ok"]


# ------------------------------------------------------------- 下载 / 待导入
ARCHIVE_EXT = (".pmp", ".zip", ".7z", ".rar", ".pcp", ".ttmp2", ".ttmp", ".penumbra")
PARTIAL_EXT = (".crdownload", ".part", ".partial", ".tmp", ".download", ".opdownload", ".bak")


def resolve_dirs(cfg):
    """返回 (下载目录, 暂存目录)"""
    dl = cfg.get("download_dir") or ""
    if not dl:
        dl = str(Path.home() / "Downloads")
    ib = cfg.get("inbox_dir") or ""
    if not ib:
        ib = str(Path(cfg["root"]).resolve().parent / "_待导入")
    return dl, ib


def list_pending(cfg, all_files=False):
    """列出下载目录 / 暂存目录里还没入库的文件（默认只认常见 Mod 文件格式）"""
    dl, ib = resolve_dirs(cfg)
    out = []
    for tag, d in (("下载", dl), ("暂存", ib)):
        p = Path(d)
        if not p.is_dir():
            continue
        items = []
        for f in p.iterdir():
            if not f.is_file():
                continue
            ext = f.suffix.lower()
            if ext in PARTIAL_EXT:            # 没下完的临时文件
                continue
            if all_files or ext in ARCHIVE_EXT:
                items.append(f)
        items.sort(key=lambda f: f.stat().st_mtime, reverse=True)
        for f in items:
            st = f.stat()
            a, n = split_mod_name(f.stem)
            out.append({"src": str(f), "where": tag, "file": f.name,
                        "size": st.st_size, "mtime": st.st_mtime,
                        "author": a, "name": n or f.stem})
    return out


# ------------------------------------------------- 安装检查 / 查重 / 序号重排
def norm_name(s) -> str:
    """名字规范化：去掉空格标点、统一小写（中英文都保留）"""
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", (s or "").lower())


def find_install_dir(cfg) -> str:
    """Mod 安装目录（Penumbra 的 Mods 目录）：配置优先，其次自动探测"""
    p = (cfg.get("install_dir") or "").strip()
    if p and Path(p).is_dir():
        return p
    ap = os.environ.get("APPDATA") or str(Path.home() / "AppData" / "Roaming")
    for c in (Path(ap) / "XIVLauncher/pluginConfigs/Penumbra/Mods",
              Path(ap) / "XIVLauncher/pluginConfigs/Penumbra/mods",
              Path(ap) / "XIVLauncher/pluginConfigs/Penumbra",
              Path.home() / "Documents/Penumbra/Mods"):
        try:
            if c.is_dir():
                return str(c)
        except Exception:
            pass
    return p


def list_installed(cfg):
    """返回 (规范化文件夹名集合, 原始文件夹名列表, 目录路径)"""
    d = find_install_dir(cfg)
    names, raw = set(), []
    if d and Path(d).is_dir():
        try:
            for e in sorted(Path(d).iterdir()):
                if e.is_dir():
                    raw.append(e.name)
                    names.add(norm_name(e.name))
        except Exception:
            pass
    return names, raw, d


def installed_match(m, names, raw=None) -> bool:
    """按安装目录里的文件夹名判断这个 Mod 是否已安装"""
    if not names:
        return False
    name = m.get("name") or ""
    author = m.get("author") or ""
    cand = {norm_name(name), norm_name("%s_%s" % (name, author)),
            norm_name("%s %s" % (name, author)), norm_name("%s-%s" % (name, author)),
            norm_name("%s [%s]" % (name, author)), norm_name("%s(%s)" % (name, author))}
    cand = {c for c in cand if c}
    if cand & names:
        return True
    base = norm_name(name)
    if base and len(base) >= 6:
        for n in names:
            if n and (base in n or n in base):
                return True
    return False


def mark_installed(cfg, mods) -> dict:
    """给每条记录标上 installed 状态，返回 {folder: bool}"""
    names, raw, d = list_installed(cfg)
    out = {}
    for m in mods:
        out[m["folder"]] = installed_match(m, names, raw)
    return out


def find_duplicates(mods):
    """找重复 Mod：同一个 Mod 地址 / 相同的 作者+名称。返回 [(原因, 说明, [记录...])]"""
    out, seen = [], set()
    by_addr = {}
    for m in mods:
        a = (m.get("addr") or "").strip().lower()
        if a:
            by_addr.setdefault(a, []).append(m)
    for a, g in sorted(by_addr.items()):
        if len(g) > 1:
            key = tuple(sorted(x["folder"] for x in g))
            if key not in seen:
                seen.add(key)
                out.append(("地址相同", a, g))
    by_key = {}
    for m in mods:
        k = (norm_name(m.get("name")), norm_name(m.get("author")))
        if k[0]:
            by_key.setdefault(k, []).append(m)
    for k, g in sorted(by_key.items()):
        if len(g) > 1:
            key = tuple(sorted(x["folder"] for x in g))
            if key not in seen:
                seen.add(key)
                out.append(("名称+作者相同", "%s / %s" % (g[0].get("author") or "-",
                                                         g[0].get("name") or "-"), g))
    return out


def renumber_plan(cfg, category="", subcat="", start=1):
    """算出重排计划 [(记录, 新序号)]；category 为空或「全部分类」时，按分类各自编号"""
    allm = Store().all()
    order = list(cfg.get("category_order") or [])
    if category and category != "全部分类":
        cats = [category]
    else:
        cats = sorted({m["category"] for m in allm},
                      key=lambda c: (order.index(c) if c in order else 99, c))
    plan = []
    for cat in cats:
        ms = [m for m in allm if m["category"] == cat
              and (not subcat or (m.get("subcat") or "") == subcat)]
        ms.sort(key=lambda m: (m["seq"], m["name"]))
        plan += [(m, i) for i, m in enumerate(ms, start=start) if m["seq"] != i]
    return plan


def renumber(cfg, category="", subcat="", start=1, dry_run=True, export=False):
    """把序号重排成连续的 start..N（每个分类各自编号）。

    dry_run=True 只返回计划；否则真的改文件夹名（先临时号再正式号，避免撞名）。
    """
    plan = renumber_plan(cfg, category, subcat, start)
    if dry_run or not plan:
        return plan
    tmp = []
    for i, (m, new_seq) in enumerate(plan):
        t = rename_mod(cfg, m["folder"], seq=900000 + i)
        tmp.append((str(t), new_seq))
    done = []
    for folder, new_seq in tmp:
        done.append(str(rename_mod(cfg, folder, seq=new_seq)))
    cmd_scan(cfg, quiet=True)
    if export:
        cmd_export(cfg)
    return done


def ensure_preview(cfg, folder, author="", name="", local_cover="", addr="") -> str:
    """确保 Mod 文件夹里有预览图，并返回图片路径（没有则返回 ""）。

    优先级：
      1. 文件夹里已经有预览图 → 直接用，不覆盖
      2. 用本次抓好的本地封面（local_cover）复制成 <文件夹名>.jpg
      3. 按 Mod 地址（地址里有 /modid/ 且内置浏览器已开）去站点抓封面
    """
    folder = Path(folder)
    if not folder.is_dir():
        return ""
    base = folder.name
    try:
        found = find_image(folder, base, author, name)[0]
    except Exception:
        found = None
    if found:
        return str(found)                       # 已经有预览图，别动
    dest = folder / (base + ".jpg")
    if local_cover and Path(local_cover).is_file():
        try:
            shutil.copy2(local_cover, dest)
            log("已放入封面图：%s" % dest)
            return str(dest)
        except Exception as e:
            log("封面放入失败：%s" % e)
    got = re.search(r"/modid/(\d+)", addr or "")
    if got and browser_running(cfg):
        try:
            browser_wait_ready(cfg, timeout=25)
            cur = browser_eval(cfg, "location.href") or ""
            if ("/modid/%s" % got.group(1)) not in cur:
                browser_eval(cfg, "location.href=%s" % json.dumps(addr))
                browser_wait_ready(cfg, timeout=25)
            info = browser_capture(cfg)
            if info.get("cover"):
                browser_fetch(cfg, info["cover"], dest)
                log("已按地址抓取封面：%s" % dest)
                return str(dest)
        except SystemExit as e:
            log("按地址抓封面跳过：%s" % e)
        except Exception as e:
            log("按地址抓封面失败：%s" % e)
    return ""


def has_preview(folder, author="", name="") -> bool:
    try:
        return bool(find_image(Path(folder), Path(folder).name, author, name)[0])
    except Exception:
        return False


def fmt_size(n) -> str:
    n = float(n or 0)
    if n < 1048576:
        return "%.0f KB" % (n / 1024)
    if n < 1073741824:
        return "%.1f MB" % (n / 1048576)
    return "%.2f GB" % (n / 1073741824)


def norm_addr(text: str) -> str:
    """把粘贴进来的网址补全（支持 xivmodarchive.com/modid/123 这种写法）"""
    a = (text or "").strip()
    if not a or a.lower().startswith(("http://", "https://")):
        return a
    if "xivmodarchive.com" in a or "heliosphere.app" in a:
        return "https://" + a.lstrip("/")
    return a


# ------------------------------------------------- 内置浏览器（Edge / Chrome + CDP）
# 思路：用独立配置目录启动一个「真实浏览器」——站点的人机验证能正常通过；
#      工具通过调试端口(CDP)读取当前页面的 网址/名称/作者/封面/下载直链，
#      再借用浏览器的 Cookie 抓封面，并把文件直接下载到暂存目录。
BROWSER_HINT = ("需要一个基于 Chromium 的浏览器（Microsoft Edge 或 Google Chrome）。\n"
                "工具会用独立配置目录启动它，登录一次即可长期保持。")
BROWSER_HOME = "https://www.xivmodarchive.com/"

PAGE_INFO_JS = (
    "JSON.stringify((function(){"
    "function abs(h){try{return h?(new URL(h,location.href)).href:''}catch(e){return h||''}}"
    "var links=[].slice.call(document.querySelectorAll('a'));"
    # 候选下载链接：优先 #mod-download-link，其次 #download-mod-button 所在的那个 <a>，
    # 再到 href 里带 /files/ 的链接；**只取 href 非空的那个** ——
    # 站点改版/登录状态下 #mod-download-link 有时挂在 <button> 上而没有 href，
    # 以前直接用它就会拿到空值，兜底也不会跑（这就是"找不到下载链接"的来源）。
    "var cands=[];"
    "var byId=document.querySelector('#mod-download-link'); if(byId)cands.push(byId);"
    "var btn=document.querySelector('#download-mod-button');"
    "if(btn&&btn.closest&&btn.closest('a'))cands.push(btn.closest('a'));"
    "links.forEach(function(e){if((e.getAttribute('href')||'').indexOf('/files/')>-1)cands.push(e);});"
    "links.forEach(function(e){if(e.hasAttribute('download'))cands.push(e);});"
    "var dl=null,href='';"
    "for(var i=0;i<cands.length;i++){var hh=cands[i].getAttribute('href')||'';"
    "  if(hh&&hh!=='#'&&hh.indexOf('javascript:')!==0){dl=cands[i];href=hh;break;}}"
    "var img=document.querySelector('img[src*=\"/mod-images/\"]');"
    "var author='';"
    "links.forEach(function(e){"
    "  if(!author&&(e.getAttribute('href')||'').indexOf('/user/')===0){author=e.innerText.trim();}});"
    "var h1=document.querySelector('h1');"
    "var text=(document.body?document.body.innerText:'').slice(0,400);"
    "var low=text.toLowerCase();"
    # Files 标签页是页内 Bootstrap 标签（href="#files"）—— 只能点，不能跳 URL
    "var filesTabEl=document.querySelector('a[data-toggle=\"tab\"][href$=\"#files\"], a[href=\"#files\"]');"
    "if(!filesTabEl){links.forEach(function(e){"
    "  if(!filesTabEl&&(e.innerText||'').trim().toLowerCase()==='files')filesTabEl=e;});}"
    "var filesTab=filesTabEl?'1':'';"
    "var tags=[];"
    "[].slice.call(document.querySelectorAll('.mod-meta-block,div,p,li')).forEach(function(e){"
    "  if(tags.length||e.children.length>3)return;"
    "  var t2=(e.innerText||'').trim();"
    "  var m2=t2.match(/^Tags\\s*:\\s*([\\s\\S]*)$/i);"
    "  if(m2){m2[1].split(',').forEach(function(x){x=x.trim(); if(x&&tags.indexOf(x)<0)tags.push(x);});}});"
    "var races='',genders='',mtype='',affects='',lastUp='',firstRel='';"
    "[].slice.call(document.querySelectorAll('.mod-meta-block,div,p,li')).forEach(function(e){"
    "  if(e.children.length>3)return;var t3=(e.innerText||'').trim();var m3;"
    "  if(!races&&(m3=t3.match(/^Races?\\s*:\\s*(.+)$/i)))races=m3[1].trim();"
    "  if(!genders&&(m3=t3.match(/^Genders?\\s*:\\s*(.+)$/i)))genders=m3[1].trim();"
    "  if(!mtype&&(m3=t3.match(/^Type\\s*:\\s*(.+)$/i)))mtype=m3[1].trim();"
    # XIV Mod Archive 的「Affects / Replaces」：可能是单件"Eerie Tights"，
    # 也可能是多件（逗号/斜杠分隔），原样读回来给用户改
    "  if(!affects&&(m3=t3.match(/^Affects\\s*\\/\\s*Replaces\\s*:\\s*(.+)$/i)))affects=m3[1].trim();"
    # 「检查更新」用：NSFW 的 Mod 只有真浏览器读得到页面，所以这两栏也一起带回来
    "  if(!lastUp&&(m3=t3.match(/^Last Version Update\\s*:\\s*(.+)$/i)))lastUp=m3[1].trim();"
    "  if(!firstRel&&(m3=t3.match(/^Original Release Date\\s*:\\s*(.+)$/i)))firstRel=m3[1].trim();});"
    "var dlc=[];"
    "links.forEach(function(e){var h2=abs(e.getAttribute('href')||'');"
    "  if(h2&&/\\/files\\/|\\.pmp|\\.ttmp2?|\\.zip|\\.7z/i.test(h2)&&dlc.indexOf(h2)<0)dlc.push(h2);});"
    "if(dlc.length<5){var all=[].slice.call(document.querySelectorAll('[href],[onclick],[data-href],[data-url]'));"
    "  all.forEach(function(e){if(dlc.length>=5)return;"
    "    var s=(e.getAttribute('href')||'')+' '+(e.getAttribute('onclick')||'')+' '+(e.getAttribute('data-href')||'')+' '+(e.getAttribute('data-url')||'');"
    "    var r2=s.match(/[^\\s\"']*\\/(?:private|files)\\/(?:[^\\s\"']+)/);"
    "    if(r2&&dlc.indexOf(r2[0])<0)dlc.push(abs(r2[0]));});}"
    "return {url:location.href,title:document.title,"
    "name:h1?h1.innerText.trim():'',author:author,cover:img?img.src:'',"
    "dl:dl?abs(href):'',dlRaw:href,dlCands:dlc.slice(0,5),filesTab:filesTab,"
    "tags:tags.slice(0,40),races:races,genders:genders,mtype:mtype,affects:affects,"
    "lastUpdate:lastUp,firstRelease:firstRel,"
    "login:(location.href.indexOf('/login')>-1)||(low.indexOf('log in')>-1)||(low.indexOf('sign in')>-1),"
    "hidden:(low.indexOf('hidden')>-1)||(low.indexOf('not authorized')>-1)||(low.indexOf('no files')>-1),"
    "ready:document.readyState,body:text};"
    "})())")

DL_CLICK_JS = (
    "(function(){"
    "var l=[].slice.call(document.querySelectorAll('a')).filter(function(e){"
    "return (e.getAttribute('href')||'').indexOf('/files/')>-1;});"
    "if(l.length){l[0].click();return 'link';}"
    "var b=document.querySelector('#download-mod-button');"
    "if(b){b.click();return 'button';}"
    "return '';})()")

NEXT_PAGE_JS = (
    "(function(){"
    "var a=document.querySelector('a[rel=\"next\"],.pagination .next a,li.next a');"
    "if(!a){var l=[].slice.call(document.querySelectorAll('a')).filter(function(e){"
    "return /下一页|next|›|»/i.test(e.innerText||'');});a=l[0];}"
    "if(a){a.click();return a.getAttribute('href')||'ok';}return '';})()")


def detect_browsers():
    """扫描本机 Chromium 内核浏览器，返回 [(名称, exe 路径)]"""
    la = os.environ.get("LOCALAPPDATA") or os.path.join(
        os.environ.get("USERPROFILE", str(Path.home())), "AppData", "Local")
    pf = os.environ.get("ProgramFiles", r"C:\Program Files")
    pf86 = os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")
    def j(*a):
        return os.path.join(*[x for x in a if x])
    cands = [
        ("CentBrowser 百分浏览器", [j(la, "CentBrowser/Application/chrome.exe"),
                                 j(pf, "CentBrowser/Application/chrome.exe"),
                                 j(pf86, "CentBrowser/Application/chrome.exe")]),
        ("Google Chrome", [j(pf, "Google/Chrome/Application/chrome.exe"),
                           j(pf86, "Google/Chrome/Application/chrome.exe"),
                           j(la, "Google/Chrome/Application/chrome.exe")]),
        ("Microsoft Edge", [j(pf86, "Microsoft/Edge/Application/msedge.exe"),
                            j(pf, "Microsoft/Edge/Application/msedge.exe")]),
        ("Brave", [j(la, "BraveSoftware/Brave-Browser/Application/brave.exe"),
                   j(pf, "BraveSoftware/Brave-Browser/Application/brave.exe")]),
        ("Vivaldi", [j(la, "Vivaldi/Application/vivaldi.exe")]),
        ("Opera", [j(la, "Programs/Opera/opera.exe"), j(pf, "Opera/opera.exe")]),
        ("360 极速浏览器", [j(la, "360Chrome/Chrome/Application/360chrome.exe"),
                        j(pf86, "360Chrome/Chrome/Application/360chrome.exe")]),
        ("Chromium", [j(la, "Chromium/Application/chrome.exe")]),
    ]
    out, seen = [], set()
    for name, paths in cands:
        for q in paths:
            if q and Path(q).is_file() and q.lower() not in seen:
                seen.add(q.lower())
                out.append((name, q))
                break
    return out


def find_browser(cfg):
    p = (cfg.get("browser_path") or "").strip()
    if p and Path(p).is_file():
        return p
    found = detect_browsers()
    return found[0][1] if found else None


def browser_port(cfg) -> int:
    try:
        return int(str(cfg.get("browser_port") or 9333))
    except Exception:
        return 9333


def browser_profile(cfg) -> Path:
    """每个浏览器单独一个配置目录（不同内核互读配置会崩）"""
    d = (cfg.get("browser_dir") or "").strip()
    if d:
        d = Path(d)
    else:
        exe = find_browser(cfg) or "browser"
        slug = re.sub(r"[^0-9A-Za-z]+", "_", Path(exe).stem)[:12] or "browser"
        tag = hashlib.md5(str(exe).lower().encode("utf-8")).hexdigest()[:6]
        d = APP_DIR / "browser_profile" / ("%s_%s" % (slug, tag))
    d.mkdir(parents=True, exist_ok=True)
    return d


def _http_json(url, timeout=3):
    import urllib.request
    with urllib.request.urlopen(url, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8", "replace"))


def browser_running(cfg) -> bool:
    try:
        return bool(_http_json("http://127.0.0.1:%d/json/version" % browser_port(cfg))
                    .get("webSocketDebuggerUrl"))
    except Exception:
        return False


def launch_browser(cfg, url=BROWSER_HOME):
    """启动（或复用）内置浏览器，返回调试端口"""
    import time
    exe = find_browser(cfg)
    if not exe:
        raise SystemExit(BROWSER_HINT)
    port = browser_port(cfg)
    if browser_running(cfg):
        return port
    Path(resolve_dirs(cfg)[1]).mkdir(parents=True, exist_ok=True)
    subprocess.Popen([exe,
                      "--user-data-dir=%s" % browser_profile(cfg),
                      "--remote-debugging-port=%d" % port,
                      "--remote-allow-origins=*",
                      "--no-first-run", "--no-default-browser-check",
                      "--disable-session-crashed-bubble",
                      url], close_fds=True)
    for _ in range(60):
        time.sleep(0.5)
        if browser_running(cfg):
            break
    return port


class CDP:
    """极简 Chrome DevTools Protocol 客户端"""

    def __init__(self, ws_url, timeout=25):
        import websocket
        self.ws = websocket.create_connection(ws_url, timeout=timeout, suppress_origin=True)
        self.n = 0

    def call(self, method, **params):
        self.n += 1
        self.ws.send(json.dumps({"id": self.n, "method": method, "params": params}))
        while True:
            msg = json.loads(self.ws.recv())
            if msg.get("id") == self.n:
                if "error" in msg:
                    raise RuntimeError("%s -> %s" % (method, msg["error"]))
                return msg.get("result", {})

    def js(self, expr, await_promise=False):
        r = self.call("Runtime.evaluate", expression=expr,
                      returnByValue=True, awaitPromise=await_promise)
        return r.get("result", {}).get("value")

    def close(self):
        try:
            self.ws.close()
        except Exception:
            pass


_CDP_CACHE = {}


def _cdp_new(cfg, page_only=True, prefer="xivmodarchive"):
    """新建一个 CDP 连接（内部用）"""
    port = browser_port(cfg)
    if page_only:
        targets = [t for t in _http_json("http://127.0.0.1:%d/json/list" % port)
                   if t.get("type") == "page" and (t.get("url") or "").startswith(("http", "file"))]
        if not targets:
            raise SystemExit("浏览器里没有打开中的网页")
        if prefer:
            hit = [t for t in targets if prefer in (t.get("url") or "")]
            if hit:
                targets = hit + [t for t in targets if t not in hit]
        last = None
        for t in targets:
            try:
                return CDP(t["webSocketDebuggerUrl"])
            except Exception as e:
                last = e
        raise SystemExit("连接浏览器失败：%s" % last)
    return CDP(_http_json("http://127.0.0.1:%d/json/version" % port)["webSocketDebuggerUrl"])


def _cdp_get(cfg, page_only=True, prefer="xivmodarchive"):
    """拿一个可复用的 CDP 连接（断了自愈重建）——比每次新建快很多"""
    port = browser_port(cfg)
    key = (port, bool(page_only), prefer or "")
    cdp = _CDP_CACHE.get(key)
    if cdp is not None:
        try:
            cdp.call("Runtime.evaluate", expression="1", returnByValue=True)
            return cdp
        except Exception:
            try:
                cdp.close()
            except Exception:
                pass
            _CDP_CACHE.pop(key, None)
    if not browser_running(cfg):
        raise SystemExit("内置浏览器没有运行，请先点「打开浏览器」")
    cdp = _cdp_new(cfg, page_only, prefer)
    _CDP_CACHE[key] = cdp
    return cdp


def browser_connect(cfg, page_only=True, prefer="xivmodarchive"):
    """连到内置浏览器（默认复用缓存连接，不要 close）"""
    return _cdp_get(cfg, page_only, prefer)


def browser_eval(cfg, expr, prefer="xivmodarchive", tries=3, await_promise=False):
    """在浏览器页面里执行 JS（复用连接，断了自动重连重试）"""
    import time as _t
    last = None
    for _ in range(max(1, tries)):
        try:
            return _cdp_get(cfg, True, prefer).js(expr, await_promise=await_promise)
        except SystemExit:
            raise
        except Exception as e:
            last = e
            for k in [k for k in list(_CDP_CACHE) if k[1]]:
                try:
                    _CDP_CACHE.pop(k).close()
                except Exception:
                    pass
            _t.sleep(0.3)
    raise SystemExit("读取浏览器页面失败：%s" % last)


CF_TITLES = ("请稍候", "just a moment", "checking your browser", "attention required", "验证")

PAGE_STATE_JS = ("JSON.stringify({t:document.title,u:location.href,"
                 "h:(document.querySelector('h1')||{}).innerText||'',"
                 "i:(document.querySelector('img[src*=\"/mod-images/\"]')||{}).src||''})")


def browser_page_state(cfg, prefer="xivmodarchive"):
    """返回 (标题, 网址, 是否还在人机验证) —— 一次查询拿全部"""
    try:
        v = json.loads(browser_eval(cfg, PAGE_STATE_JS, prefer=prefer, tries=1) or "{}")
    except SystemExit:
        raise
    except Exception:
        return "", "", True
    title, url = v.get("t", ""), v.get("u", "")
    low = title.strip().lower()
    challenge = (not low) or any(k in low for k in CF_TITLES)
    return title, url, challenge


def browser_wait_ready(cfg, timeout=30, prefer="xivmodarchive", on_tick=None):
    """等 Cloudflare 人机验证自动通过（页面一好立刻返回）"""
    import time
    t0 = time.time()
    while True:
        try:
            v = json.loads(browser_eval(cfg, PAGE_STATE_JS, prefer=prefer, tries=1) or "{}")
        except SystemExit:
            raise
        except Exception:
            v = {}
        title, url = v.get("t", ""), v.get("u", "")
        low = title.strip().lower()
        if any(k in low for k in CF_TITLES) or not title.strip():
            ready = False
        elif "/modid/" in (url or ""):
            ready = bool(v.get("h")) or bool(v.get("i"))     # Mod 详情页：要有内容才算好
        else:
            ready = ("xiv mod archive" in low) or bool(v.get("h")) or bool(v.get("i"))
        if ready:
            return title, url
        if time.time() - t0 >= timeout:
            return title, url
        if on_tick:
            try:
                on_tick(int(time.time() - t0))
            except Exception:
                pass
        time.sleep(0.6)


# ------------------------------------------------------ 站点更新时间 / 检查更新
SITE_MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
               "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
SITE_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"


def _local_offset() -> _dt.timedelta:
    return _dt.datetime.now() - _dt.datetime.utcnow()


def site_time_iso(text) -> str:
    """站点上的日期文本 -> 本机时区的 'YYYY-MM-DD HH:MM:SS'。

    形如 'Tue Aug 18 2026 21:41:17 GMT+0000 (Coordinated Universal Time)'。
    不用 strptime 是因为 %a/%b 跟系统语言有关，中文环境下会解析失败。
    """
    s = str(text or "").strip()
    # 登录后站点用的是 <code class="server-date">2026/9/21 08:34:00</code> —— 实测那一栏
    # 已经是**浏览器本地时区**（拿已知 UTC 值反推确认过），所以直接当本地时间存
    m0 = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})\s+(\d{1,2}):(\d{2}):(\d{2})\s*$", s)
    if m0:
        try:
            return _dt.datetime(*[int(x) for x in m0.groups()]).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return ""
    m = re.search(r"(\w{3})\s+(\w{3})\s+(\d{1,2})\s+(\d{4})\s+(\d{2}):(\d{2}):(\d{2})\s*GMT\s*([+-]\d{4})",
                  s)
    if not m:
        return ""
    mon = SITE_MONTHS.get(m.group(2).lower())
    if not mon:
        return ""
    try:
        t = _dt.datetime(int(m.group(4)), mon, int(m.group(3)),
                         int(m.group(5)), int(m.group(6)), int(m.group(7)))
        off = m.group(8)
        d = _dt.timedelta(hours=int(off[1:3]), minutes=int(off[3:5]))
        utc = t - (d if off[0] == "+" else -d)
        return (utc + _local_offset()).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


def site_key(addr, modid="") -> str:
    """把 Mod 地址归一成可比对的键：'xma:<modid>' / 'helio:<shortId>' / ''（认不出来）。

    用它来判断「这次下载的是不是库里已有的那条」——识别到就覆盖更新，不再新建一条。
    """
    a = str(addr or "").strip()
    mid = str(modid or "").strip()
    if "xivmodarchive" in a:
        m = re.search(r"/modid/(\d+)", a)
        return "xma:" + (m.group(1) if m else mid)
    if "heliosphere" in a:
        m = re.search(r"/mod/([A-Za-z0-9]+)", a)
        return "helio:" + (m.group(1) if m else mid)
    if mid and mid.isdigit():
        return "xma:" + mid
    if a.isdigit():
        return "xma:" + a
    return ""


def _helio_iso(text) -> str:
    """heliosphere 的 ISO 时间（2026-09-21T05:36:25.342711+00:00）→ 本机时区字符串"""
    s = str(text or "").strip()
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})[T ](\d{2}):(\d{2}):(\d{2})", s)
    if not m:
        return ""
    try:
        t = _dt.datetime(*[int(x) for x in m.groups()])
    except ValueError:
        return ""
    off = _dt.timedelta(0)
    mm2 = re.search(r"([+-])(\d{2}):(\d{2})$", s)
    if mm2:
        d = _dt.timedelta(hours=int(mm2.group(2)), minutes=int(mm2.group(3)))
        off = d if mm2.group(1) == "+" else -d
    return (t - off + _local_offset()).strftime("%Y-%m-%d %H:%M:%S")


def fetch_helio_update(cfg, addr, cookies="") -> dict:
    """读 heliosphere 的版本信息（页面 SSR 数据里就带 version / updatedAt / downloadSize）。

    注意：它的「下载」是页面上的按钮（点了才调接口，接口没公开），所以这里只做版本检测，
    更新要靠「打开页面手动下载 → 上传新文件替换」。
    """
    out = {"ok": False, "meta_ok": False, "modid": "", "updated": "", "version": "",
           "patch": "", "tags": [], "affects": "", "source": "", "error": "", "kind": "helio"}
    m = re.search(r"/mod/([A-Za-z0-9]+)", str(addr or ""))
    if not m:
        out["error"] = "不是 heliosphere 的 Mod 链接"
        return out
    out["modid"] = m.group(1)
    page = "https://heliosphere.app/mod/%s" % out["modid"]
    try:
        html = _site_get(page, cookies).decode("utf-8", "replace")
    except Exception as e:
        out["error"] = "页面读取失败：%s" % str(e)[:80]
        return out
    flat = html.replace('\\"', '"').replace("\\/", "/")      # SSR 数据里的 JSON 是转义的
    v = re.search(r'"version"\s*:\s*"([^"]+)"', flat)
    up = re.search(r'"updatedAt"\s*:\s*"([^"]+)"', flat)
    rel = re.search(r'"releasedAt"\s*:\s*"([^"]+)"', flat)
    tgs = re.findall(r'\{"category":(?:true|false),"slug":"([^"]+)"', flat)
    aff = re.search(r'"affects"\s*:\s*\[([^\]]*)\]', flat)
    if tgs:
        out["tags"] = [x.strip() for x in dict.fromkeys(tgs) if x.strip()][:30]
        out["meta_ok"] = True
    if aff is not None:
        items = [unescape(x).strip() for x in re.findall(r'"([^"]*)"', aff.group(1))]
        out["affects"] = ", ".join(x for x in items if x)
        out["meta_ok"] = True
    dl = re.search(r'"downloadSize"\s*:\s*(\d+)', flat)
    if v:
        out["version"] = v.group(1)
    iso = (up.group(1) if up else (rel.group(1) if rel else ""))
    if iso:
        out["updated"] = _helio_iso(iso)
    if dl and dl.group(1).isdigit():
        out["size_mb"] = round(int(dl.group(1)) / 1048576.0, 1)
    out["ok"] = bool(out["version"] or out["updated"])
    if not out["ok"]:
        out["error"] = "页面里没找到版本信息"
    else:
        out["source"] = "helio-page"
    return out


def _site_get(url, cookies="", timeout=25) -> bytes:
    req = urllib.request.Request(url, headers={
        "User-Agent": SITE_UA, "Referer": BROWSER_HOME,
        "X-Requested-With": "XMLHttpRequest",
        "Accept": "text/html,application/json,*/*"})
    if cookies:
        req.add_header("Cookie", cookies)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def parse_page_meta(html: str) -> dict:
    """从 Mod 页 HTML 里解析元信息：Tags / Affects-Replaces / Races / Genders / 更新时间。

    匿名与登录态的标记不一样（登录态每栏是 `<div class="mod-meta-block">标题 : <br><code>值</code>`
    且值是链接；匿名态是纯文本），统一做法：每个 meta 块**先去标签压空白**，再按「标题 : 值」切。
    """
    out = {"ok": False, "tags": [], "affects": "", "races": "", "genders": "",
           "last_update": "", "first_release": ""}
    for m in re.finditer(r'<div[^>]*class="[^"]*mod-meta-block[^"]*"[^>]*>([\s\S]*?)</div>', html or ""):
        txt = re.sub(r"<[^>]+>", " ", m.group(1))
        txt = unescape(re.sub(r"\s+", " ", txt)).strip()
        mo = re.match(r"^([A-Za-z][A-Za-z /]*?)\s*:\s*(.*)$", txt)
        if not mo:
            continue
        key, val = mo.group(1).strip().lower(), mo.group(2).strip()
        if key.startswith("tags"):
            out["tags"] = [x.strip() for x in re.split(r"[,、]", val) if x.strip()]
            out["ok"] = True
        elif key.startswith("affects"):
            out["affects"] = val.strip().strip(",;")
            out["ok"] = True
        elif key.startswith("races"):
            out["races"] = val
        elif key.startswith("genders"):
            out["genders"] = val
        elif key.startswith("last version update"):
            out["last_update"] = site_time_iso(val)
        elif key.startswith("original release date"):
            out["first_release"] = site_time_iso(val)
    return out


def fetch_site_update(cfg, addr, cookies="") -> dict:
    """读站点上这条 Mod 的「最后更新时间 / 最新版本 / 更新说明 / 标签 / 影响替换」。

    顺序：**先读页面**（Tags、Affects / Replaces、Last Version Update 都在这上面），
    再问一次版本历史接口补充版本号与更新说明（读不到也无所谓）。

    返回 {ok, meta_ok, modid, updated, version, patch, tags, affects, races, genders,
          released, history, source, error}
      meta_ok=True 表示页面上那两块元信息是**可靠读到的** → 调用方可以拿它覆盖本地标签/影响替换；
      False（只有接口读通）时**不要**动本地那两项，免得把主人填的值清空。
    """
    out = {"ok": False, "meta_ok": False, "modid": "", "updated": "", "version": "",
           "patch": "", "tags": [], "affects": "", "races": "", "genders": "",
           "released": "", "history": [], "source": "", "error": ""}
    if "heliosphere" in str(addr or ""):
        return fetch_helio_update(cfg, addr, cookies)
    m = re.search(r"/modid/(\d+)", str(addr or ""))
    if not m:
        out["error"] = "这条 Mod 没有站点地址（缺 modid）"
        return out
    out["modid"] = m.group(1)
    page = "https://www.xivmodarchive.com/modid/%s" % out["modid"]
    errs = []
    try:
        html = _site_get(page, cookies).decode("utf-8", "replace")
        meta = parse_page_meta(html)
        if meta["ok"] or meta["last_update"]:
            out["tags"] = meta["tags"]
            out["affects"] = meta["affects"]
            out["races"] = meta.get("races") or ""
            out["genders"] = meta.get("genders") or ""
            out["released"] = meta.get("first_release") or ""
            out["meta_ok"] = meta["ok"]
            out["updated"] = meta["last_update"] or meta["first_release"]
            out["source"] = "page"
            out["ok"] = True
        else:
            errs.append("页面里没找到「Last Version Update」")
    except Exception as e:
        errs.append("页面：%s" % str(e)[:60])
    try:                      # 版本号 / 更新说明（有版本历史才有；读不到不影响上面的结果）
        j = json.loads(_site_get("https://www.xivmodarchive.com/api/mod/update_history?modid=%s"
                                 % out["modid"], cookies).decode("utf-8", "replace") or "{}")
        hist = j.get("version_history") or []
        if hist:
            out["history"] = [{"version": str(h.get("version_new") or ""),
                               "prev": str(h.get("version_previous") or ""),
                               "time": _dt.datetime.fromtimestamp(int(h["timestamp"]) / 1000
                                     ).strftime("%Y-%m-%d %H:%M:%S") if str(h.get("timestamp") or "").isdigit() else "",
                               "notes": str(h.get("patch_notes") or "")[:600]}
                              for h in sorted(hist, key=lambda h: -int(h.get("timestamp") or 0))][:30]
            last = max(hist, key=lambda h: int(h.get("timestamp") or 0))
            out["version"] = str(last.get("version_new") or "")
            out["patch"] = str(last.get("patch_notes") or "")[:400]
            if not out["updated"]:
                out["updated"] = _dt.datetime.fromtimestamp(
                    int(last["timestamp"]) / 1000).strftime("%Y-%m-%d %H:%M:%S")
                out["source"] = "api"
                out["ok"] = True
    except Exception as e:
        if not out["ok"]:
            errs.append("版本接口：%s" % str(e)[:60])
    if not out["ok"]:
        if errs and "页面" in errs[0] and "403" in errs[0]:
            out["error"] = ("页面需要登录（多半是 NSFW）：到「设置 → 打开内置浏览器」登录一次 "
                            "XIVModArchive 就能读能更新")
        else:
            out["error"] = "；".join(errs) or "读不到站点信息"
    return out


def fetch_history(cfg, addr, cookies="") -> dict:
    """读站点的**版本历史**（XMA 的 update_history 接口），给详情里的「历史」页签用。

    返回 {ok, modid, items: [{version, prev, time, notes}], error}
    items 按时间倒序。读不到就把原因放 error（NSFW 未登录会 403，界面据此提示去登录）。
    """
    out = {"ok": False, "modid": "", "items": [], "error": ""}
    a = str(addr or "").strip()
    if "heliosphere" in a:
        info = fetch_helio_update(cfg, a, cookies)
        if info.get("ok"):
            out["ok"] = True
            out["items"] = [{"version": info.get("version") or "", "prev": "",
                             "time": info.get("updated") or "", "notes": ""}]
        else:
            out["error"] = info.get("error") or "heliosphere 读不到版本历史"
        return out
    m = re.search(r"/modid/(\d+)", a)
    if not m:
        out["error"] = "这条 Mod 没有站点地址（缺 modid）"
        return out
    out["modid"] = m.group(1)
    try:
        info = fetch_site_update(cfg, a, cookies)
        out["items"] = info.get("history") or []
        out["ok"] = bool(out["items"])
        if not out["ok"]:
            out["error"] = info.get("error") or "站点上没有版本历史（只有一条首发记录）"
    except Exception as e:
        out["error"] = str(e)[:120]
    return out


def replace_mod_payload(cfg, folder, new_file, mode="same_name", to_recycle=True) -> dict:
    """把 Mod 文件夹里的旧载荷换成新文件（手动上传 / 从站点更新都走这里）。

    保留：文件夹本身（编号/作者/名称不变）、地址.txt、所有图片（预览图/画廊图）。
    mode:
      same_name   只换掉与新文件同名（或同主名）的旧文件 —— 多功能变体的文件夹不受影响（推荐）
      all_payload 把除 地址.txt / 图片 之外的全部旧文件换掉
    """
    folder, new_file = Path(folder), Path(new_file)
    if not folder.is_dir():
        raise SystemExit("Mod 文件夹不存在：%s" % folder)
    if not new_file.exists():
        raise SystemExit("新文件不存在：%s" % new_file)
    olds = [f for f in sorted(folder.iterdir())
            if f.is_file() and f.name != ADDR_NAME and f.suffix.lower() not in IMG_EXT]
    if mode == "all_payload":
        targets = olds
    else:
        stem = new_file.stem.lower()
        targets = [f for f in olds if f.name.lower() == new_file.name.lower() or f.stem.lower() == stem]
        if mode == "auto" and not targets and olds:
            # 自动模式：名字对不上（作者改了文件名）就当整份替换，免得库里留两个版本
            targets = olds
            mode = "all_payload"
        elif mode == "auto":
            mode = "same_name"
    removed, kept = [], []
    for f in targets:
        try:
            if not (to_recycle and send_to_recycle_bin(f)):
                f.unlink()
            removed.append(f.name)
        except OSError as e:
            raise SystemExit("替换失败：%s 处理不了（%s）" % (f.name, e))
    kept = [f.name for f in olds if f not in targets]
    dst = folder / new_file.name
    if new_file.is_dir():
        shutil.copytree(str(new_file), str(dst), dirs_exist_ok=True)
    else:
        shutil.copy2(str(new_file), str(dst))
    log("  %s：换掉 %s -> %s（保留 %s）"
        % (folder.name, "、".join(removed) or "（无同名旧文件）", new_file.name,
           "、".join(kept) or "无"))
    return {"folder": str(folder), "mode": mode, "removed": removed, "kept": kept,
            "added": [dst.name] if not new_file.is_dir() else ["（目录）" + dst.name]}


def browser_capture(cfg) -> dict:
    """读取当前页面的 网址/名称/作者/封面/下载直链（以及站点上的更新时间）"""
    info = json.loads(browser_eval(cfg, PAGE_INFO_JS) or "{}")
    for k in ("lastUpdate", "firstRelease"):        # 转成本机时区的标准写法
        if info.get(k):
            info[k + "_iso"] = site_time_iso(info[k])
    m = re.search(r"/modid/(\d+)", info.get("url") or "")
    info["modid"] = m.group(1) if m else ""
    info["addr"] = ("https://www.xivmodarchive.com/modid/%s" % m.group(1)) if m else ""
    info["is_mod"] = bool(m)
    low = (info.get("title") or "").strip().lower()
    info["challenge"] = (not low) or any(k in low for k in CF_TITLES)
    return info


def browser_fetch(cfg, url, dest) -> int:
    """借用浏览器 Cookie 抓资源（封面图等），保存到 dest"""
    import urllib.request
    if not (url or "").strip():
        raise SystemExit("没有抓到封面地址（页面可能还没加载好）")
    dest = Path(dest)
    cdp = _cdp_get(cfg, True)
    ua = cdp.js("navigator.userAgent") or "Mozilla/5.0"
    try:
        cks = cdp.call("Network.getCookies",
                       urls=["https://www.xivmodarchive.com/",
                             "https://static.xivmodarchive.com/"])["cookies"]
    except Exception:
        cdp.call("Network.enable")
        cks = cdp.call("Network.getCookies",
                       urls=["https://www.xivmodarchive.com/",
                             "https://static.xivmodarchive.com/"])["cookies"]
    hdr = {"User-Agent": ua,
           "Cookie": "; ".join("%s=%s" % (c["name"], c["value"]) for c in cks),
           "Referer": BROWSER_HOME,
           "Accept": "image/avif,image/webp,image/*,*/*;q=0.8"}
    req = urllib.request.Request(url, headers=hdr)
    with urllib.request.urlopen(req, timeout=30) as r:
        data = r.read()
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_bytes(data)
    return len(data)


def user_data_dir_of(exe):
    """从浏览器 exe 推出它的 "User Data" 配置目录"""
    out = []
    if exe:
        q = Path(exe)
        if q.parent.name.lower() == "application":
            out.append(q.parent.parent / "User Data")
        out.append(q.parent.parent / "User Data")
    la = os.environ.get("LOCALAPPDATA") or os.path.join(
        os.environ.get("USERPROFILE", str(Path.home())), "AppData", "Local")
    for rel in ("CentBrowser/User Data", "Google/Chrome/User Data", "Microsoft/Edge/User Data",
                "BraveSoftware/Brave-Browser/User Data", "360Chrome/Chrome/User Data",
                "Vivaldi/User Data", "Chromium/User Data"):
        out.append(Path(la) / rel)
    for d in out:
        try:
            if d.is_dir() and ((d / "Default").is_dir() or list(d.glob("Profile *"))):
                return d
        except OSError:
            pass
    return None


def _cookie_file_in(profile_dir):
    for rel in ("Network/Cookies", "Cookies"):
        f = Path(profile_dir) / rel
        if f.is_file():
            return f
    return None


def browser_download_dir(cfg):
    """从你自己的浏览器配置里读出"默认下载目录"（只读 Preferences，不碰 Cookie）。

    为什么要它：方案是"让浏览器自己下载"（用你已登录的身份），但文件会落到
    浏览器设置的下载目录里，管理器得知道去哪儿接住它。Preferences 是普通 JSON，
    浏览器开着也能读。
    """
    ud = user_data_dir_of(find_browser(cfg))
    home = os.environ.get("USERPROFILE") or str(Path.home())
    guess = Path(home) / "Downloads"
    if not ud:
        return str(guess)
    for pd in ([Path(ud) / "Default"] + sorted(Path(ud).glob("Profile *"))):
        pref = pd / "Preferences"
        if not pref.is_file():
            continue
        try:
            j = json.loads(pref.read_text(encoding="utf-8", errors="ignore") or "{}")
            d = ((j.get("download") or {}).get("default_directory") or "").strip()
            if d and Path(d).is_dir():
                return d
        except Exception:
            pass
    return str(guess)


def browser_login_state(cfg):
    """看内置浏览器里有没有 xivmodarchive 的 Cookie（= 登录状态）"""
    try:
        if not browser_running(cfg):
            return {"running": False, "count": 0, "names": [], "logged": False}
        cdp = _cdp_get(cfg, True)
        try:
            cks = cdp.call("Network.getCookies",
                           urls=["https://www.xivmodarchive.com/",
                                 "https://static.xivmodarchive.com/"])["cookies"]
        except Exception:
            cdp.call("Network.enable")
            cks = cdp.call("Network.getCookies",
                           urls=["https://www.xivmodarchive.com/",
                                 "https://static.xivmodarchive.com/"])["cookies"]
        names = sorted({c.get("name", "") for c in cks})
        low = [n.lower() for n in names]
        logged = any(("session" in n) or n.startswith("xf_") or ("remember" in n) for n in low)
        return {"running": True, "count": len(cks), "names": names[:12], "logged": logged}
    except SystemExit:
        raise
    except Exception as e:
        return {"running": True, "count": 0, "names": [], "logged": False, "error": str(e)[:80]}


def sync_login_from_user_browser(cfg):
    """把你平时用的浏览器里的 Cookie 搬进内置浏览器的配置目录。

    为什么需要：管理器只能靠调试端口(CDP)读页面，而你自己那个浏览器没开调试端口，
    所以读不到你"已登录"的页面 —— NSFW 的 Mod 在内置浏览器里没登录就是看不到下载。
    同一台电脑、同一个浏览器程序 → 加密密钥一致，Cookie 复制过去就能直接解开用。
    """
    import time
    exe = find_browser(cfg)
    src_ud = user_data_dir_of(exe)
    if not src_ud:
        raise SystemExit("没找到你的浏览器配置目录（先在设置里选好浏览器程序）")
    if Path(src_ud).resolve() == Path(browser_profile(cfg)).resolve():
        raise SystemExit("内置浏览器用的就是你的浏览器配置目录，不用同步")
    srcs = [d for d in ([Path(src_ud) / "Default"] + sorted(Path(src_ud).glob("Profile *")))
            if _cookie_file_in(d)]
    if not srcs:
        raise SystemExit("你的浏览器里没有找到 Cookie 文件：%s" % src_ud)
    best = max(srcs, key=lambda d: _cookie_file_in(d).stat().st_mtime)
    was = browser_running(cfg)
    if was:                                  # 先关掉内置浏览器，免得它退出时把 Cookie 写回去
        try:
            _cdp_get(cfg, False).call("Browser.close")
        except Exception:
            pass
        for _ in range(25):
            if not browser_running(cfg):
                break
            time.sleep(0.4)
    for k in list(_CDP_CACHE):               # 连接缓存清掉（浏览器已经关了）
        try:
            _CDP_CACHE.pop(k).close()
        except Exception:
            pass
    # 加密密钥要一致（Local State 的 os_crypt），Cookie 才解得开
    src_ls = Path(src_ud) / "Local State"
    if src_ls.is_file():
        try:
            a = json.loads(src_ls.read_text(encoding="utf-8", errors="ignore") or "{}")
            dst_ls = Path(browser_profile(cfg)) / "Local State"
            b = json.loads(dst_ls.read_text(encoding="utf-8", errors="ignore") or "{}") if dst_ls.is_file() else {}
            if a.get("os_crypt"):
                b["os_crypt"] = a["os_crypt"]
                dst_ls.parent.mkdir(parents=True, exist_ok=True)
                dst_ls.write_text(json.dumps(b), encoding="utf-8")
        except Exception:
            pass
    n = 0
    locked = ""
    dst_net = Path(browser_profile(cfg)) / "Network"
    dst_net.mkdir(parents=True, exist_ok=True)
    for name in ("Cookies", "Cookies-journal", "Cookies-wal", "Cookies-shm"):
        s = _cookie_file_in(best).parent / name
        if s.is_file():
            for _try in range(4):          # 有些浏览器是"短暂占用"，多试几次
                try:
                    shutil.copy2(str(s), str(dst_net / name))
                    n += 1
                    break
                except OSError as e:
                    if _try < 3:
                        time.sleep(0.6)
                    elif not locked:
                        locked = "%s（%s）" % (s, e.__class__.__name__)
    if n == 0:
        if locked:      # Chromium 把 Cookie 库独占占用，只有关掉它才能读
            raise SystemExit(
                "你自己的浏览器正在运行，它的 Cookie 文件被占用，读不出来：%s。"
                "两种做法：① 关掉你自己的浏览器（点『再试一次』即可，只需一次）；"
                "② 直接在内置浏览器里登录一次 XIVModArchive —— 登录一次长期有效。" % locked)
        raise SystemExit("没能复制到 Cookie（源目录里没有可读文件）")
    if was:             # 同步完把内置浏览器重新拉起来
        try:
            launch_browser(cfg, BROWSER_HOME)
        except Exception:
            pass
    return {"from": str(best), "to": str(dst_net), "files": n, "was_running": was}


def browser_goto(cfg, url, timeout=45, prefer=""):
    """让内置浏览器停在指定网址上（坑见下）。

    launch_browser() 只在浏览器"没在运行"时才带 URL 启动；CentBrowser 长期开着，
    所以以前下载时其实只是"随便读一个 xivmodarchive 标签页"，读到哪页看运气 ——
    这正是"页面上没找到下载链接"的来源。现在：优先复用/导航已有的 Mod 标签页，
    没有才新开一个（不动你别的页面）。
    """
    import time
    url = (url or "").strip()
    if not url:
        return
    launch_browser(cfg, url)                       # 没运行就先起（带 URL）
    pref = prefer
    m = re.search(r"/modid/(\d+)", url)
    if m:
        pref = pref or ("modid/%s" % m.group(1))
    try:
        pages = [x for x in _http_json("http://127.0.0.1:%d/json/list" % browser_port(cfg))
                 if x.get("type") == "page" and (x.get("url") or "").startswith("http")]
        mine = [x for x in pages if (pref and pref in (x.get("url") or ""))
                or ("modid" in url and "xivmodarchive" in (x.get("url") or ""))]
        if not mine:
            mine = [x for x in pages if "xivmodarchive" in (x.get("url") or "")]
        if mine:
            cdp = CDP(mine[0]["webSocketDebuggerUrl"])
            try:
                cdp.call("Page.enable")
                cdp.call("Page.navigate", url=url)
            finally:
                try:
                    cdp.close()
                except Exception:
                    pass
            time.sleep(1.2)
        else:                                      # 没有 XMA 标签页 → 新开一个
            _cdp_get(cfg, False).call("Target.createTarget", url=url)
            time.sleep(2.0)
    except SystemExit:
        raise
    except Exception:
        pass
    return browser_wait_ready(cfg, timeout=timeout, prefer=pref or "xivmodarchive")


def browser_download(cfg, dest_dir, timeout=900, on_tick=None, should_cancel=None, url=""):
    """让浏览器下载当前 Mod 的文件，并把它搬进 dest_dir。

    返回 (下载直链, 落地文件路径)；失败返回 ("", "")。
    说明：部分 Edge 版本会忽略 CDP 的下载路径设置，所以这里做了兜底——
    监视浏览器真实下载目录，检测到新文件下载完成后自动搬进 dest_dir。
    """
    dest_dir = Path(dest_dir)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dl_dir = Path(resolve_dirs(cfg)[0])
    dl_dir.mkdir(parents=True, exist_ok=True)
    import time as _t0
    t_click = _t0.time()

    try:                                   # 能直接下到目标目录最好
        cdpb = browser_connect(cfg, page_only=False)      # 复用连接，别关
        cdpb.call("Browser.setDownloadBehavior", behavior="allow",
                  downloadPath=str(dest_dir), eventsEnabled=True)
    except Exception:
        pass

    pref = "xivmodarchive"
    _m = re.search(r"/modid/(\d+)", url or "")
    if _m:
        pref = "modid/%s" % _m.group(1)
    if url:
        try:
            browser_goto(cfg, url, prefer=pref)      # 先精确停到这条 Mod 的页面
        except SystemExit:
            raise
        except Exception:
            pass
    info = json.loads(browser_eval(cfg, PAGE_INFO_JS, prefer=pref) or "{}")
    log("浏览器当前页：%s（标题 %s，dl=%s，filesTab=%s，login=%s，hidden=%s）"
              % ((info.get("url") or "?")[:80], (info.get("title") or "?")[:30],
                 "有" if info.get("dl") else "无", info.get("filesTab") or "-",
                 info.get("login"), info.get("hidden")))
    if not re.search(r"/modid/\d+", info.get("url") or ""):
        raise SystemExit("当前页面不是 Mod 详情页（网址里没有 /modid/）")
    if not info.get("dl"):
        # 主页面没链接：有些 Mod 的文件只在 Files 标签页里（页内标签，点一下再读）
        try:
            browser_eval(cfg, "var t=document.querySelector('a[data-toggle=\"tab\"][href$=\"#files\"],a[href=\"#files\"]');"
                              "(t&&t.click&&t.click(),t?'clicked':'no-tab')")
            time.sleep(1.5)
            again = json.loads(browser_eval(cfg, PAGE_INFO_JS, prefer=pref) or "{}")
            if again.get("dl"):
                info = again
        except Exception:
            pass
    if not info.get("dl"):
        why = []
        if not info.get("dl"):
            why.append("这条 Mod 的页面里没有下载链接 —— 最常见的原因是**内置浏览器没有登录**"
                       "（NSFW 的 Mod 不登录就看不到下载）。"
                       "去内置浏览器里登录一次 XIVModArchive，或用「同步登录状态」把你浏览器的登录搬过去")
        if info.get("login"):
            why.append("当前页面像是登录页 → 去内置浏览器里登录 XIVModArchive（登录一次会记住）")
        if info.get("hidden"):
            why.append("页面提示内容被隐藏 → NSFW 的 Mod 要在站点账号设置里打开成人内容")
        why.append("若文件在 Files 标签页里，请在内置浏览器里点一下 Files 再重试")
        why.append("或在浏览器里手动点下载，再用「监视下载目录」")
        _c = info.get("dlCands") or []
        if _c:
            why.append("页面上确实有这些像下载地址的链接：" + " | ".join(x[:70] for x in _c[:3]))
        raise SystemExit("页面上没找到下载链接。%s（页面标题：%s）"
                         % ("；".join(why), (info.get("title") or "?")[:40]))
    browser_eval(cfg, DL_CLICK_JS)
    href = info.get("dl", "")

    import time
    t0, last, done, cancelled = time.time(), {}, None, False
    while time.time() - t0 < timeout:
        time.sleep(0.8)
        if should_cancel and should_cancel():
            cancelled = True
            break
        if on_tick:
            try:
                on_tick(int(time.time() - t0))
            except Exception:
                pass
        for f in dest_dir.iterdir():                       # 已经在目标目录
            if f.is_file() and f.suffix.lower() not in PARTIAL_EXT:
                done = f
                break
        if done:
            break
        news = []
        for f in dl_dir.iterdir():          # 按修改时间判定"这次下载的"
            try:
                if (f.is_file() and f.suffix.lower() not in PARTIAL_EXT
                        and f.stat().st_mtime >= t_click - 15):
                    news.append(f)
            except OSError:
                pass
        if news:
            f = max(news, key=lambda x: x.stat().st_mtime)
            try:
                size = f.stat().st_size
            except OSError:
                size = -1
            if size > 0 and last.get(f.name) == size:       # 大小不再变化 = 下载完成
                done = f
                break
            last[f.name] = size
    if cancelled:                       # 取消：顺手清掉没下完的临时文件
        for f in list(dl_dir.iterdir()):
            try:
                if (f.is_file() and f.suffix.lower() in PARTIAL_EXT
                        and f.stat().st_mtime >= t_click - 15):
                    f.unlink(missing_ok=True)
            except OSError:
                pass
        return href, ""
    if not done:
        return href, ""
    if done.parent.resolve() != dest_dir.resolve():
        target = dest_dir / done.name
        n = 1
        while target.exists():
            target = dest_dir / ("%s (%d)%s" % (done.stem, n, done.suffix))
            n += 1
        try:
            shutil.move(str(done), str(target))
            done = target
            log("下载完成并搬入暂存： %s" % done)
        except Exception as e:
            log("搬移下载文件失败（文件留在下载目录）：%s" % e)
    return href, str(done)


# ------------------------------------------------------------------- library
def next_seq(parent: Path) -> int:
    """同一目录里下一个可用序号（现有最大序号 + 1）"""
    n = 0
    if parent.is_dir():
        for d in parent.iterdir():
            if d.is_dir():
                m = MOD_RE.match(d.name)
                if m:
                    n = max(n, int(m.group(1)))
    return n + 1


def split_mod_name(text: str):
    """从 '3.[作者] 名称' / '[作者] 名称' / '名称' 里拆出 (作者, 名称)"""
    t = re.sub(r"^\s*\d+\s*\.\s*", "", (text or "").strip())
    a, n = split_author_name(t)
    return a, re.sub(r"^-\s*", "", n)


def mod_folder_name(seq: int, author: str, name: str) -> str:
    return "%d.[%s] %s" % (seq, author, name) if author else "%d.%s" % (seq, name)


def import_mod(cfg, src, category, zone="SFW", subdir="", author=None, name=None,
               seq=None, addr="", move=False) -> Path:
    """把来源文件夹/压缩包导入 分类/类型[子目录]/序号.[作者] 名称 下"""
    src = Path(src)
    if not src.exists():
        raise SystemExit("来源不存在：%s" % src)
    if author is None or name is None:
        a, n = split_mod_name(src.name if src.is_dir() else src.stem)
        author = a if author is None else author
        name = n if name is None else name
    author, name = (author or "").strip(), (name or "").strip()
    if not name:
        raise SystemExit("Mod 名称不能为空")
    parent = Path(cfg["root"]) / category
    if zone:
        parent = parent / zone
    for part in re.split(r"[\\/]+", (subdir or "").strip()):
        if part:
            parent = parent / part
    seq = next_seq(parent) if seq in (None, "", 0) else int(seq)
    target = parent / mod_folder_name(seq, author, name)
    if target.exists():
        raise SystemExit("目标已存在：%s" % target)
    parent.mkdir(parents=True, exist_ok=True)
    if src.is_dir():
        (shutil.move if move else shutil.copytree)(str(src), str(target))
    else:
        target.mkdir(parents=True, exist_ok=True)
        dst = target / src.name
        (shutil.move if move else shutil.copy2)(str(src), str(dst))
    if addr:
        (target / ADDR_NAME).write_text(addr.strip() + "\n", encoding="utf-8")
    return target


def parse_mod_folder(root, folder: Path):
    """从 Mod 文件夹里解析出 (分类, 类型, 子目录, 序号, 作者, 名称)；路径相对 Mod 根目录"""
    folder = Path(folder)
    try:
        parts = ("_",) + folder.resolve().relative_to(Path(root).resolve()).parts
    except ValueError:
        parts = ("_", folder.name)
    base = folder.name
    m = MOD_RE.match(base)
    seq = int(m.group(1)) if m else 0
    rest = base[m.end():] if m else base
    author, _nm2 = split_author_name(rest)
    name = re.sub(r"^-\s*", "", _nm2)
    zone, mid = "SFW", []
    if len(parts) > 2:
        mid = list(parts[2:-1])
        for q in mid:
            if q.upper() in ("SFW", "NSFW"):
                zone = q.upper()
        mid = [q for q in mid if q.upper() not in ("SFW", "NSFW")]
    return (parts[1] if len(parts) > 1 else "", zone, os.sep.join(mid), seq, author, name)


def mod_parent(cfg, category, zone, subdir):
    parent = Path(cfg["root"]) / (category or "").strip()
    if zone:
        parent = parent / zone
    for part in re.split(r"[\\/]+", (subdir or "").strip()):
        if part:
            parent = parent / part
    return parent


def rename_mod(cfg, folder, category=None, zone=None, subdir=None, author=None, name=None,
               seq=None, addr=None) -> Path:
    """改名 / 改分类·类型 / 移动 / 改地址；folder 是现有 Mod 文件夹"""
    old = Path(folder)
    if not old.is_dir():
        raise SystemExit("文件夹不存在：%s" % old)
    root = Path(cfg["root"]).resolve()
    try:
        old.resolve().relative_to(root)
    except ValueError:
        raise SystemExit("只能操作 Mod 根目录下的内容：%s" % old)
    c_cat, c_zone, c_sub, c_seq, c_author, c_name = parse_mod_folder(cfg["root"], old)

    category = (c_cat if category is None else category.strip())
    zone = (c_zone if zone is None else (zone or "SFW").strip())
    subdir = (c_sub if subdir is None else subdir.strip())
    author = (c_author if author is None else author.strip())
    name = (c_name if name is None else name.strip())
    seq = c_seq if seq in (None, "", 0) else int(seq)
    if not name:
        raise SystemExit("Mod 名称不能为空")

    parent = mod_parent(cfg, category, zone, subdir)
    target = parent / mod_folder_name(seq, author, name)
    if target.resolve() == old.resolve():
        target = old                                   # 只有地址变了
    elif target.exists():
        raise SystemExit("目标已存在：%s" % target)
    else:
        parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(old), str(target))
        log("已移动到： %s" % target)
        try:                                        # 标签跟着文件夹走
            st = Store()
            moved = st.move_tags(str(old), str(target))
            st.cx.close()
            if moved:
                log("已带过 %d 个标签" % moved)
        except Exception as e:
            log("标签迁移失败：%s" % e)

    if addr is not None:
        addr = addr.strip()
        f = target / ADDR_NAME
        if addr:
            f.write_text(addr + "\n", encoding="utf-8")
            log("已写入地址： %s" % f)
        elif f.exists():
            f.unlink()
            log("已删除地址文件： %s" % f)
    return target


def send_to_recycle_bin(path) -> bool:
    """把文件/文件夹移入回收站（可还原，不依赖第三方库）"""
    import ctypes
    from ctypes import wintypes

    class SHFILEOPSTRUCTW(ctypes.Structure):
        _fields_ = [("hwnd", wintypes.HWND),
                    ("wFunc", wintypes.UINT),
                    ("pFrom", wintypes.LPCWSTR),
                    ("pTo", wintypes.LPCWSTR),
                    ("fFlags", ctypes.c_uint16),
                    ("fAnyOperationsAborted", wintypes.BOOL),
                    ("hNameMappings", ctypes.c_void_p),
                    ("lpszProgressTitle", wintypes.LPCWSTR)]

    fp = os.path.abspath(str(path))
    buf = ctypes.create_unicode_buffer(fp, len(fp) + 2)      # 需要双 NUL 结尾
    op = SHFILEOPSTRUCTW()
    op.wFunc = 3                                             # FO_DELETE
    op.pFrom = ctypes.cast(buf, wintypes.LPCWSTR)
    op.pTo = None
    op.fFlags = 0x0040 | 0x0010 | 0x0004 | 0x0400            # ALLOWUNDO|NOCONFIRM|SILENT|NOERRORUI
    rc = ctypes.windll.shell32.SHFileOperationW(ctypes.byref(op))
    return rc == 0 and not op.fAnyOperationsAborted


def delete_mod(folder, to_recycle=True) -> str:
    folder = Path(folder)
    if not folder.is_dir():
        raise SystemExit("文件夹不存在：%s" % folder)
    if to_recycle:
        if send_to_recycle_bin(folder):
            return "已移入回收站"
        log("  ! 移入回收站失败，改为直接删除")
    shutil.rmtree(str(folder))
    return "已直接删除"


# ---------------------------------------------------------------------- store
TAG_MAX_LEN = 24
TAG_MAX_PER_MOD = 30


def norm_tag(t) -> str:
    """单个标签：去空白、收紧内部空格、限长"""
    s = re.sub(r"\s+", " ", str(t or "")).strip()
    return s[:TAG_MAX_LEN]


AFFECTS_MAX = 200          # 「影响/替换」原文照抄，只做长度上限保护


def norm_affects(text) -> str:
    """规范化「影响/替换」：去首尾空白、把连续空白压成一个空格、限长。

    值来自 XIV Mod Archive 的 `Affects / Replaces` 一栏，可能是
    "Eerie Tights" 这种单件，也可能是 "Eastern Technogaskins, Eastern Technojacket"
    这种多件 —— 原样保留（不做结构拆分），只清掉换行和多余空格。
    """
    s = re.sub(r"\s+", " ", str(text or "")).strip().strip(",;、；")
    return s[:AFFECTS_MAX]


def norm_tags(tags) -> list:
    """一组标签：去空、大小写不敏感去重（保留输入写法）、限个数"""
    out, seen = [], set()
    for t in (tags or []):
        s = norm_tag(t)
        if not s:
            continue
        k = s.lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(s)
        if len(out) >= TAG_MAX_PER_MOD:
            break
    return out


class Store:
    def __init__(self, path=None):
        # 注意：默认值不能写成 path=DB_PATH —— 那样会在函数定义时就绑定，
        # 之后改 mm.DB_PATH（比如 --root 临时模式）根本不生效，会去动真实索引库。
        self.cx = sqlite3.connect(str(path or DB_PATH))
        self.cx.row_factory = sqlite3.Row
        self.cx.executescript(SCHEMA)
        have = {r["name"] for r in self.cx.execute("PRAGMA table_info(mods)")}
        for col, typ in (("img_source", "TEXT"), ("subcat", "TEXT"), ("affects", "TEXT"),
                         ("site_updated", "TEXT"), ("site_latest", "TEXT"), ("site_version", "TEXT"),
                         ("site_checked", "TEXT"), ("update_avail", "INTEGER"),
                         ("races", "TEXT"), ("genders", "TEXT"), ("released", "TEXT"), ("desc", "TEXT"),
                         ("cloud_backend", "TEXT"), ("cloud_path", "TEXT"), ("cloud_state", "TEXT"),
                         ("cloud_size", "INTEGER"), ("cloud_synced", "TEXT"), ("payload_mtime", "REAL")):
            if col not in have:
                self.cx.execute("ALTER TABLE mods ADD COLUMN %s %s" % (col, typ))
        self.cx.commit()

    def upsert(self, m: dict, reuse_hash: bool = True):
        old = self.cx.execute("SELECT img_hash, file_mtime, file_size FROM mods WHERE folder=?",
                              (m["folder"],)).fetchone()
        img_hash, mtime, size = "", None, None
        img = Path(m["img"]) if m["img"] else None
        if img and img.is_file():
            st = img.stat()
            mtime, size = st.st_mtime, st.st_size
            if (reuse_hash and old and old["img_hash"] and old["file_mtime"] == mtime
                    and old["file_size"] == size):
                img_hash = old["img_hash"]
            else:
                img_hash = md5_of(img)
        self.cx.execute(
            """INSERT INTO mods (folder,category,seq,author,nsfw,subcat,name,addr,addr_source,
                                 img,img_source,img_hash,file_mtime,file_size,updated_at,affects)
               VALUES (:folder,:category,:seq,:author,:nsfw,:subcat,:name,:addr,:addr_source,
                       :img,:img_source,:img_hash,:file_mtime,:file_size,:updated_at,:affects)
               ON CONFLICT(folder) DO UPDATE SET
                 category=excluded.category, seq=excluded.seq, author=excluded.author,
                 nsfw=excluded.nsfw, subcat=excluded.subcat, name=excluded.name, addr=excluded.addr,
                 addr_source=excluded.addr_source, img=excluded.img,
                 img_source=excluded.img_source, img_hash=excluded.img_hash,
                 file_mtime=excluded.file_mtime, file_size=excluded.file_size,
                 -- 重扫描时记录里没有 affects（它不在文件夹名里），给 NULL 就保留库里已有的；
                 -- 手动清空传的是 ""（不是 NULL），所以照样能清掉。
                 affects=COALESCE(excluded.affects, mods.affects),
                 site_updated=COALESCE(excluded.site_updated, mods.site_updated),
                 site_latest=COALESCE(excluded.site_latest, mods.site_latest),
                 site_version=COALESCE(excluded.site_version, mods.site_version),
                 site_checked=COALESCE(excluded.site_checked, mods.site_checked),
                 update_avail=COALESCE(excluded.update_avail, mods.update_avail),
                 updated_at=excluded.updated_at""",
            {**m, "img_hash": img_hash, "file_mtime": mtime, "file_size": size,
             "affects": m.get("affects"), "site_updated": m.get("site_updated"),
             "site_latest": m.get("site_latest"), "site_version": m.get("site_version"),
             "site_checked": m.get("site_checked"), "update_avail": m.get("update_avail"),
             "updated_at": _dt.datetime.now().isoformat(timespec="seconds")})
        return img_hash

    def all(self) -> list:
        return [dict(r) for r in self.cx.execute("SELECT * FROM mods ORDER BY category, seq, name")]

    # ------------------------------------------------------------ 标签
    def tags_of(self, folder) -> list:
        return [r["tag"] for r in self.cx.execute(
            "SELECT tag FROM mod_tags WHERE folder=? ORDER BY tag", (str(folder),))]

    def tags_map(self) -> dict:
        d = {}
        for r in self.cx.execute("SELECT folder, tag FROM mod_tags ORDER BY tag"):
            d.setdefault(r["folder"], []).append(r["tag"])
        return d

    def all_tags(self) -> list:
        return [{"tag": r["tag"], "count": r["n"]} for r in self.cx.execute(
            "SELECT tag, COUNT(*) AS n FROM mod_tags GROUP BY tag ORDER BY n DESC, tag")]

    # ------------------------------------------------------- 影响/替换（affects）
    def set_affects(self, folder, text) -> str:
        """手工设置「影响/替换」（空串 = 清空）"""
        f = str(folder)
        val = norm_affects(text)
        self.cx.execute("UPDATE mods SET affects=? WHERE folder=?", (val, f))
        self.cx.commit()
        return val

    def set_affects_many(self, folders, text) -> int:
        """批量写「影响/替换」（空串 = 清空），返回改了几条"""
        val = norm_affects(text)
        n = 0
        for f in folders:
            n += self.cx.execute("UPDATE mods SET affects=? WHERE folder=?", (val, str(f))).rowcount
        self.cx.commit()
        return n

    def set_site_info(self, folder, updated=None, latest=None, version=None,
                      checked=None, avail=None) -> dict:
        """写「检查更新」相关字段（None = 这一项不动）。

        updated = 站点上 Last Version Update（下载/更新时记基线）
        latest  = 最近一次检查看到的站点值；avail = 1/0 有没有新版
        """
        pairs = (("site_updated", updated), ("site_latest", latest), ("site_version", version),
                 ("site_checked", checked), ("update_avail", avail))
        sets = [(c, v) for c, v in pairs if v is not None]
        if sets:
            self.cx.execute("UPDATE mods SET %s WHERE folder=?"
                            % ", ".join("%s=?" % c for c, _ in sets),
                            [v for _, v in sets] + [str(folder)])
            self.cx.commit()
        return dict(sets)

    def set_site_meta(self, folder, races=None, genders=None, released=None) -> dict:
        """写站点元信息：种族 / 性别 / 首发日期（None = 这一项不动）"""
        pairs = (("races", races), ("genders", genders), ("released", released))
        sets = [(c, v) for c, v in pairs if v is not None]
        if sets:
            self.cx.execute("UPDATE mods SET %s WHERE folder=?"
                            % ", ".join("%s=?" % c for c, _ in sets),
                            [v for _, v in sets] + [str(folder)])
            self.cx.commit()
        return dict(sets)

    def set_desc(self, folder, text) -> str:
        """写「内容描述」（空串 = 清空）"""
        val = str(text or "").strip()
        self.cx.execute("UPDATE mods SET desc=? WHERE folder=?", (val, str(folder)))
        self.cx.commit()
        return val

    # ------------------------------------------------- 云存储（归档）
    def set_cloud(self, folder, backend=None, path=None, state=None, size=None,
                  synced=None, payload_mtime=None) -> dict:
        """写云存储相关字段（None = 这一项不动）"""
        pairs = (("cloud_backend", backend), ("cloud_path", path), ("cloud_state", state),
                 ("cloud_size", size), ("cloud_synced", synced), ("payload_mtime", payload_mtime))
        sets = [(c, v) for c, v in pairs if v is not None]
        if sets:
            self.cx.execute("UPDATE mods SET %s WHERE folder=?"
                            % ", ".join("%s=?" % c for c, _ in sets),
                            [v for _, v in sets] + [str(folder)])
            self.cx.commit()
        return dict(sets)

    def set_payload_files(self, folder, items, state="archived") -> int:
        """整批写入载荷清单（items: [{rel_path,size,md5,sha1,cloud_fid}]）"""
        f = str(folder)
        now = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        n = 0
        for it in items:
            self.cx.execute(
                "INSERT INTO payload_files(folder,rel_path,size,md5,sha1,cloud_fid,state,checked) "
                "VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(folder,rel_path) DO UPDATE SET "
                "size=excluded.size, md5=excluded.md5, sha1=excluded.sha1, "
                "cloud_fid=excluded.cloud_fid, state=excluded.state, checked=excluded.checked",
                (f, str(it.get("rel_path")), int(it.get("size") or 0), str(it.get("md5") or ""),
                 str(it.get("sha1") or ""), str(it.get("cloud_fid") or ""), str(state), now))
            n += 1
        self.cx.commit()
        return n

    def payload_files_of(self, folder, states=None) -> list:
        """某条 mod 的载荷清单（可按状态过滤）"""
        sql = "SELECT * FROM payload_files WHERE folder=?"
        args = [str(folder)]
        if states:
            sql += " AND state IN (%s)" % ",".join("?" * len(states))
            args += [str(x) for x in states]
        return [dict(r) for r in self.cx.execute(sql + " ORDER BY rel_path", args)]

    def payload_files_map(self, states=None) -> dict:
        """全部载荷清单：folder -> [rows]（列表页做状态标记用）"""
        sql = "SELECT * FROM payload_files"
        args = []
        if states:
            sql += " WHERE state IN (%s)" % ",".join("?" * len(states))
            args = [str(x) for x in states]
        out = {}
        for r in self.cx.execute(sql, args):
            out.setdefault(r["folder"], []).append(dict(r))
        return out

    def clear_payload_files(self, folder, state=None) -> int:
        """删掉某条 mod 的载荷清单（可按状态）"""
        if state:
            n = self.cx.execute("DELETE FROM payload_files WHERE folder=? AND state=?",
                                (str(folder), str(state))).rowcount
        else:
            n = self.cx.execute("DELETE FROM payload_files WHERE folder=?", (str(folder),)).rowcount
        self.cx.commit()
        return n

    def cloud_counts(self) -> dict:
        """归档状态统计（界面顶部显示）"""
        d = {"archived": 0, "local": 0, "missing": 0, "uploading": 0, "none": 0}
        for r in self.cx.execute("SELECT COALESCE(cloud_state,'') AS s, COUNT(*) AS n FROM mods GROUP BY s"):
            k = r["s"] or "none"
            d[k if k in d else "none"] = r["n"]
        return d

    def affects_all(self) -> list:
        """已有取值 + 条数（给界面做建议/筛选）"""
        return [{"affects": r["affects"], "count": r["n"]} for r in self.cx.execute(
            "SELECT affects, COUNT(*) AS n FROM mods WHERE affects IS NOT NULL AND affects<>'' "
            "GROUP BY affects ORDER BY n DESC, affects")]

    def set_tags(self, folder, tags) -> list:
        f = str(folder)
        want = norm_tags(tags)
        self.cx.execute("DELETE FROM mod_tags WHERE folder=?", (f,))
        for t in want:
            self.cx.execute("INSERT OR IGNORE INTO mod_tags (folder, tag) VALUES (?,?)", (f, t))
        self.cx.commit()
        return want

    def add_tags(self, folders, tags) -> int:
        want = norm_tags(tags)
        n = 0
        for f in folders:
            for t in want:
                cur = self.cx.execute(
                    "INSERT OR IGNORE INTO mod_tags (folder, tag) VALUES (?,?)", (str(f), t))
                n += cur.rowcount or 0
        self.cx.commit()
        return n

    def remove_tags(self, folders, tags) -> int:
        want = [t.lower() for t in norm_tags(tags)]
        n = 0
        for f in folders:
            for t in want:
                cur = self.cx.execute(
                    "DELETE FROM mod_tags WHERE folder=? AND lower(tag)=?", (str(f), t))
                n += cur.rowcount or 0
        self.cx.commit()
        return n

    def move_tags(self, old_folder, new_folder) -> int:
        """文件夹改名/移动时把标签带过去"""
        if str(old_folder) == str(new_folder):
            return 0
        n = 0
        for r in self.cx.execute("SELECT tag FROM mod_tags WHERE folder=?", (str(old_folder),)):
            cur = self.cx.execute("INSERT OR IGNORE INTO mod_tags (folder, tag) VALUES (?,?)",
                                  (str(new_folder), r["tag"]))
            n += cur.rowcount or 0
        self.cx.execute("DELETE FROM mod_tags WHERE folder=?", (str(old_folder),))
        self.cx.commit()
        return n

    def remap_folder_prefix(self, old_prefix, new_prefix) -> int:
        """整目录改名（比如改分类名）时，把这一批 Mod 的标签一起搬到新路径"""
        old_p = str(old_prefix).rstrip("\\/")
        new_p = str(new_prefix).rstrip("\\/")
        if old_p == new_p:
            return 0
        n = 0
        rows = list(self.cx.execute("SELECT folder, tag FROM mod_tags"))
        sep = ("\\", "/")
        for r in rows:
            f = r["folder"]
            if f == old_p or f.startswith(old_p + sep[0]) or f.startswith(old_p + sep[1]):
                cur = self.cx.execute("INSERT OR IGNORE INTO mod_tags (folder, tag) VALUES (?,?)",
                                      (new_p + f[len(old_p):], r["tag"]))
                n += cur.rowcount or 0
        self.cx.execute(
            "DELETE FROM mod_tags WHERE folder=? OR folder LIKE ? OR folder LIKE ?",
            (old_p, old_p + "\\%", old_p + "/%"))
        self.cx.commit()
        return n

    def prune_tags(self, alive: set) -> int:
        """删掉已经不存在的 Mod 的标签"""
        dead = [r["folder"] for r in self.cx.execute("SELECT DISTINCT folder FROM mod_tags")
                if r["folder"] not in alive]
        for f in dead:
            self.cx.execute("DELETE FROM mod_tags WHERE folder=?", (f,))
        if dead:
            self.cx.commit()
        return len(dead)

    def prune(self, alive: set) -> int:
        dead = [r["folder"] for r in self.cx.execute("SELECT folder FROM mods")
                if r["folder"] not in alive]
        for f in dead:
            self.cx.execute("DELETE FROM mods WHERE folder=?", (f,))
        if dead:
            self.prune_tags(alive)
        return len(dead)

    def commit(self):
        self.cx.commit()


# --------------------------------------------------------------------- thumbs
def ensure_thumb(img_path: str, width: int, img_hash: str, cache: dict) -> str:
    """按需生成缩略图，返回用于嵌入的图片路径"""
    try:
        from PIL import Image
    except Exception:
        return img_path
    if not img_hash:
        img_hash = md5_of(Path(img_path))
    key = "%s_%d" % (img_hash[:16], width)
    if key in cache:
        return cache[key]
    THUMB_DIR.mkdir(parents=True, exist_ok=True)
    out = THUMB_DIR / (key + ".jpg")
    if not out.exists():
        try:
            im = Image.open(img_path)
            im = im.convert("RGB")
            if im.width > width:
                im = im.resize((width, max(1, round(im.height * width / im.width))), Image.LANCZOS)
            im.save(out, "JPEG", quality=86, optimize=True)
        except Exception:
            return img_path
    cache[key] = str(out)
    return str(out)


def fit_size(w, h, box_w, box_h):
    if not w or not h:
        return box_w, box_h
    s = min(box_w / w, box_h / h)
    return max(1, round(w * s)), max(1, round(h * s))


# --------------------------------------------------------------------- export
def write_excel(mods: list, cfg: dict, out_path: Path, use_store_hash=None, progress=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    try:
        from openpyxl.drawing.image import Image as XLImage
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        have_img = True
    except Exception:
        have_img = False

    embed = bool(cfg.get("embed_images", True)) and have_img
    box_w, box_h = cfg.get("preview_px", [178, 100])
    thumb_w = int(cfg.get("thumb_width", 800))
    order = cfg.get("category_order") or []

    def cat_key(c):
        return (order.index(c) if c in order else len(order), c)

    cats = sorted({m["category"] for m in mods}, key=cat_key)
    thin = Side(style="thin", color="FF9C9C9C")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    f_head = Font(name="等线", size=10, bold=True)
    f_body = Font(name="等线", size=10)

    wb = Workbook()
    wb.remove(wb.active)
    cache = {}
    n_img = 0
    done = 0
    total_n = max(1, len(mods))
    for cat in cats:
        rows = sorted([m for m in mods if m["category"] == cat],
                      key=lambda x: ((x.get("subcat") or ""), x["seq"], x["name"]))
        # 这个分类里有子分类时，表格里多加一列「子分类」
        show_sub = any((m.get("subcat") or "") for m in rows)
        if show_sub:
            heads = ["序号", "子分类", "作者", "NSFW/SFW", "影响/替换", "更新时间", "Mod名称", "Mod地址", "预览图"]
            cws = [7.7, 15.857, 15.857, 15.857, 26.0, 17.0, 48.0, 63.0, 28.0]
        else:
            heads = list(HEADERS)
            cws = list(COL_WIDTHS)
        nail = len(heads)                                   # 预览图所在的 1-based 列号
        nail_px = round(cws[-1] * 7 + 5)
        ws = wb.create_sheet(title=cat[:31])
        ws.sheet_view.showGridLines = True
        for i, h in enumerate(heads, 1):
            c = ws.cell(1, i, h)
            c.font, c.alignment, c.border = f_head, center, border
        ws.row_dimensions[1].height = 18
        for i, w in enumerate(cws, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for r, m in enumerate(rows, start=2):
            vals = ([m["seq"], m.get("subcat") or "", m["author"], m["nsfw"],
                     m.get("affects") or "", m.get("site_updated") or "",
                     m["name"], m["addr"], None] if show_sub else
                    [m["seq"], m["author"], m["nsfw"], m.get("affects") or "",
                     m.get("site_updated") or "", m["name"], m["addr"], None])
            for i, v in enumerate(vals, 1):
                c = ws.cell(r, i, v)
                c.font, c.alignment, c.border = f_body, center, border
            ws.row_dimensions[r].height = ROW_HEIGHT_PT
            if embed and m.get("img") and Path(m["img"]).is_file():
                try:
                    src = ensure_thumb(m["img"], thumb_w, use_store_hash.get(m["folder"], "") if use_store_hash else "", cache)
                    xi = XLImage(src)
                    w, h = fit_size(xi.width, xi.height, box_w, box_h)
                    xoff = int((nail_px - w) / 2 * EMU_PER_PX)
                    yoff = int((ROW_HEIGHT_PT * 96 / 72 - h) / 2 * EMU_PER_PX)
                    xi.anchor = OneCellAnchor(
                        _from=AnchorMarker(col=nail - 1,
                                           colOff=max(0, xoff),
                                           row=r - 1,
                                           rowOff=max(0, yoff)),
                        ext=XDRPositiveSize2D(int(w * EMU_PER_PX), int(h * EMU_PER_PX)))
                    ws.add_image(xi)
                    n_img += 1
                except Exception as e:
                    log("  ! 预览图嵌入失败 %s : %s" % (m["name"], e))
            done += 1
            if progress:
                progress(done, total_n, "%s ｜ %s" % (cat, m["name"]))
        if cfg.get("autofilter") and rows:
            ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(heads)), len(rows) + 1)
        ws.freeze_panes = "A2"

    if progress:
        progress(total_n, total_n, "正在写入 Excel 文件…")
    tmp = out_path.with_name(out_path.stem + ".tmp.xlsx")
    wb.save(tmp)
    if out_path.exists():
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M%S")
        try:
            shutil.copy2(out_path, BACKUP_DIR / ("%s.%s.xlsx" % (out_path.stem, stamp)))
            olds = sorted(BACKUP_DIR.glob(out_path.stem + ".*.xlsx"), key=lambda q: q.stat().st_mtime)
            for q in olds[:-5]:                 # 只保留最近 5 份备份
                try:
                    q.unlink()
                except OSError:
                    pass
        except Exception:
            pass
        try:
            os.remove(out_path)
        except OSError:
            pass
    shutil.move(str(tmp), str(out_path))
    return len(cats), len(mods), n_img


def _remove_part(tmp: Path, tries: int = 8) -> bool:
    """删掉打包中途的半成品；被杀软/索引器短暂占用时重试几次"""
    for _ in range(tries):
        try:
            Path(tmp).unlink()
            return True
        except FileNotFoundError:
            return True
        except OSError:
            time.sleep(0.25)
    return False


def clean_stale_parts(cfg=None, older_than=600) -> int:
    """清掉备份目录里没人管的 *.part 半成品（上次被强杀/断电留下的）"""
    try:
        d = resolve_backup_dir(cfg) if cfg else BACKUP_DIR
        if not d.is_dir():
            return 0
        now = time.time()
        n = 0
        for q in d.glob("*.part"):
            try:
                if now - q.stat().st_mtime > older_than:
                    _remove_part(q)
                    n += 1
            except OSError:
                pass
        if n:
            log("  清理了 %d 个上次没删掉的半成品备份(.part)" % n)
        return n
    except Exception:
        return 0


# --------------------------------------------------------------- 备份 / 恢复
BACKUP_MANIFEST = "ModManager备份.json"
BACKUP_FORMAT = 1
BACKUP_APP = "FFXIV Mod 管理工具"
BACKUP_MOD_PREFIX = "mods/"
BACKUP_DB_PREFIX = "index/"
BACKUP_XL_PREFIX = "excel/"


class BackupCancelled(Exception):
    """用户中途取消了备份 / 恢复"""


def backup_stamp() -> str:
    return _dt.datetime.now().strftime("%Y%m%d_%H%M%S")


def resolve_backup_dir(cfg) -> Path:
    d = (cfg.get("backup_dir") or "").strip()
    return Path(d) if d else BACKUP_DIR


def excel_path(cfg) -> Path:
    p = (cfg.get("excel") or "").strip()
    if p:
        return Path(p)
    return Path(cfg.get("root") or ".").parent / "Mod信息汇总表.xlsx"


def default_backup_path(cfg) -> Path:
    root = Path(cfg.get("root") or "MOD")
    return resolve_backup_dir(cfg) / ("Mod备份_%s_%s.zip"
                                      % (root.name or "MOD", backup_stamp()))


def _bsize(p) -> int:
    try:
        return Path(p).stat().st_size
    except OSError:
        return 0


def iter_tree_files(root: Path):
    """遍历目录下所有文件，产出 (相对路径, 绝对路径)"""
    root = Path(root)
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames.sort()
        for fn in sorted(filenames):
            p = Path(dirpath) / fn
            try:
                yield p.relative_to(root), p
            except ValueError:
                continue


def collect_backup_items(cfg, include_mods=True, include_meta=True):
    """收集要打包的内容，返回 (entries, stats)；entries = [(zip 内路径, 磁盘路径)]"""
    entries, root = [], Path(cfg.get("root") or "")
    n_mods = mod_bytes = 0
    if include_mods and root.is_dir():
        n_mods = len(scan_root(root))
        for rel, p in iter_tree_files(root):
            entries.append((Path(BACKUP_MOD_PREFIX) / rel, p))
            mod_bytes += _bsize(p)
    meta_files = []
    if include_meta:
        for arc, src in ((Path(BACKUP_DB_PREFIX) / DB_PATH.name, DB_PATH),
                         (Path(BACKUP_DB_PREFIX) / CONFIG_PATH.name, CONFIG_PATH),
                         (Path(BACKUP_XL_PREFIX) / excel_path(cfg).name, excel_path(cfg))):
            if src.is_file():
                meta_files.append((arc, src))
    entries += meta_files
    stats = {"mods": n_mods, "files": len(entries), "mod_bytes": mod_bytes,
             "meta_files": len(meta_files),
             "meta_bytes": sum(_bsize(p) for _, p in meta_files),
             "bytes": sum(_bsize(p) for _, p in entries)}
    return entries, stats


def _archived_count() -> int:
    """有多少条 Mod 的载荷已经归档到云端（写进备份 manifest，方便日后看懂这个包）"""
    try:
        st = Store()
        n = st.cx.execute("SELECT COUNT(*) AS n FROM mods WHERE cloud_state='archived'").fetchone()["n"]
        st.cx.close()
        return int(n or 0)
    except Exception:
        return 0


def make_backup(cfg, dest=None, include_mods=True, include_meta=True, compress=False,
                progress=None, should_cancel=None):
    """打包备份。返回统计 dict。"""
    if not include_mods and not include_meta:
        raise SystemExit("「Mod 文件夹」和「索引库+汇总表」至少要勾一个。")
    dest = Path(dest) if dest else default_backup_path(cfg)
    clean_stale_parts(cfg)                     # 顺手清掉上次没删干净的半成品
    entries, stats = collect_backup_items(cfg, include_mods, include_meta)
    if not entries:
        raise SystemExit("没有可备份的内容（Mod 目录是空的？）")
    dest.parent.mkdir(parents=True, exist_ok=True)
    root = Path(cfg.get("root") or "")
    manifest = {
        "工具": BACKUP_APP, "format": BACKUP_FORMAT, "版本": APP_VERSION,
        "备份时间": _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "Mod根目录": str(root), "Mod根目录名": root.name,
        "汇总表": excel_path(cfg).name,
        "包含Mod文件夹": bool(include_mods), "包含索引和汇总表": bool(include_meta),
        "Mod数": stats["mods"], "文件数": stats["files"], "原始大小": stats["bytes"],
        "压缩": bool(compress),
        # ---- 云存储：让备份包自己说清楚「载荷在哪」 ----
        "已归档到云端": _archived_count(),
        "归档说明": ("已归档的 Mod 载荷不在本包里（长期副本在夸克网盘）："
                 "恢复后本地只有元数据+图，需要时点「从云盘取回」把载荷拉回来。"),
        "分类": sorted(d.name for d in root.iterdir() if d.is_dir()) if root.is_dir() else [],
    }
    comp = zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED
    tmp = dest.with_name(dest.name + ".part")
    total = stats["bytes"] or 1
    done_b = done_n = 0
    CHUNK = 4 << 20                      # 4 MB 一块：取消能立刻响应，进度也更顺
    last_tick = [0.0]

    def tick(name, force=False):
        now = time.time()
        if progress and (force or now - last_tick[0] > 0.15):
            last_tick[0] = now
            progress(done_n, stats["files"], min(done_b + carry[0], total), total, name)

    carry = [0]                       # 当前文件已拷的字节，用来算滚动进度
    try:
        with zipfile.ZipFile(str(tmp), "w", compression=comp, allowZip64=True) as zf:
            zf.writestr(BACKUP_MANIFEST, json.dumps(manifest, ensure_ascii=False, indent=2))
            for arc, src in entries:
                if should_cancel and should_cancel():
                    raise BackupCancelled()
                stt = src.stat()
                zi = zipfile.ZipInfo(arc.as_posix(), date_time=time.localtime(stt.st_mtime)[:6])
                zi.compress_type = comp
                carry[0] = 0
                with zf.open(zi, "w") as out, open(str(src), "rb") as fh:
                    while True:
                        if should_cancel and should_cancel():
                            raise BackupCancelled()
                        buf = fh.read(CHUNK)
                        if not buf:
                            break
                        out.write(buf)
                        carry[0] += len(buf)
                        tick(arc.name)
                done_b += carry[0]
                carry[0] = 0
                done_n += 1
                tick(arc.name, force=True)
        if should_cancel and should_cancel():
            raise BackupCancelled()
        os.replace(str(tmp), str(dest))
    except BaseException:
        _remove_part(tmp)                      # 取消/出错时别留下上 GB 的半成品
        raise
    stats.update({"path": str(dest), "zip_bytes": _bsize(dest), "manifest": manifest})
    return stats


def backup_mod_dirs(names):
    """从 zip 条目名里找出所有 Mod 文件夹（相对 mods/ 的路径）"""
    cand = set()
    for n in names:
        if not n.startswith(BACKUP_MOD_PREFIX):
            continue
        parts = n.split("/")[1:-1]
        for i, part in enumerate(parts):
            if MOD_RE.match(part):
                cand.add("/".join(parts[:i + 1]))
                break
    return sorted(d for d in cand
                  if not any(d != o and d.startswith(o + "/") for o in cand))


def owner_of(rel: str, moddirs):
    """判断 zip 里的条目属于哪个 Mod 文件夹（同级同名的预览图也算它的）"""
    best = None
    for md in moddirs:
        hit = (rel == md or rel.startswith(md + "/"))
        if not hit:                      # 同级目录里与 Mod 文件夹同名的图片/文件
            d, _, fn = rel.rpartition("/")
            if d == Path(md).parent.as_posix() and Path(fn).stem == Path(md).name:
                hit = True
        if hit and (best is None or len(md) > len(best)):
            best = md
    return best


def inspect_backup(zip_path):
    """读取备份内容说明（不解压）"""
    zip_path = Path(zip_path)
    info = {"path": str(zip_path), "bytes": _bsize(zip_path), "ok": False, "error": "",
            "manifest": None, "n_files": 0, "raw_bytes": 0, "mods": [],
            "has_db": False, "has_excel": False, "excel_name": "", "cats": []}
    if not zip_path.is_file():
        info["error"] = "文件不存在"
        return info
    try:
        with zipfile.ZipFile(str(zip_path)) as zf:
            names = zf.namelist()
            info["n_files"] = len(names)
            info["raw_bytes"] = sum(i.file_size for i in zf.infolist())
            if BACKUP_MANIFEST in names:
                try:
                    info["manifest"] = json.loads(
                        zf.read(BACKUP_MANIFEST).decode("utf-8-sig"))
                except Exception:
                    info["manifest"] = None
            info["mods"] = backup_mod_dirs(names)
            info["has_db"] = any(
                n.startswith(BACKUP_DB_PREFIX) and n.lower().endswith(".db")
                for n in names)
            info["has_excel"] = any(n.startswith(BACKUP_XL_PREFIX)
                                    and n.lower().endswith(".xlsx") for n in names)
            info["excel_name"] = next(
                (Path(n).name for n in names if n.startswith(BACKUP_XL_PREFIX)
                 and n.lower().endswith(".xlsx")), "")
            cats = []
            for n in names:
                if n.startswith(BACKUP_MOD_PREFIX):
                    seg = n.split("/")[1:2]
                    if seg and seg[0] and seg[0] not in cats:
                        cats.append(seg[0])
            info["cats"] = cats
            info["ok"] = bool(info["mods"] or info["has_db"] or info["has_excel"])
            if not info["ok"]:
                info["error"] = "里面既没有 Mod 文件夹，也没有索引库/汇总表"
    except zipfile.BadZipFile:
        info["error"] = "不是有效的 zip 文件（可能没下载完或已损坏）"
    except Exception as e:
        info["error"] = "读取失败：%s" % e
    return info


def merge_db_from(path: Path) -> int:
    """把备份里的索引库合并进当前索引库（不删已有记录），返回合并的条数"""
    n = 0
    con = sqlite3.connect(str(DB_PATH))
    try:
        con.execute("PRAGMA busy_timeout=8000")
        con.execute("ATTACH DATABASE ? AS bk", (str(path),))
        try:
            tables = [r[0] for r in con.execute(
                "SELECT name FROM bk.sqlite_master WHERE type='table'").fetchall()]
            for t in tables:
                if t.startswith("sqlite_"):
                    continue
                try:
                    before = con.total_changes
                    con.execute("INSERT OR REPLACE INTO main.%s SELECT * FROM bk.%s" % (t, t))
                    n += con.total_changes - before
                except Exception:
                    pass
            con.commit()
        finally:
            con.execute("DETACH DATABASE bk")
    finally:
        con.close()
    return n


def restore_backup(cfg, zip_path, include_mods=True, include_meta=True, mode="skip",
                   progress=None, should_cancel=None, dry_run=False):
    """从备份 zip 恢复到 Mod 根目录 / 索引库。返回统计 dict。"""
    zip_path = Path(zip_path)
    root = Path(cfg.get("root") or "")
    if mode not in ("skip", "overwrite", "rename"):
        raise SystemExit("未知的冲突处理方式：%s" % mode)
    if include_mods and not root.is_dir():
        raise SystemExit("Mod 根目录不存在或没设置：%s" % root)
    if not include_mods and not include_meta:
        raise SystemExit("「Mod 文件夹」和「索引库+汇总表」至少要勾一个。")
    st = {"mods_new": 0, "mods_skipped": 0, "mods_overwritten": 0, "mods_renamed": 0,
          "files": 0, "bytes": 0, "db": False, "excel": False, "db_rows": 0,
          "renamed_list": [], "db_sidecar": ""}
    with zipfile.ZipFile(str(zip_path)) as zf:
        names = zf.namelist()
        infos = {i.filename: i for i in zf.infolist()}

        # ---------------- Mod 文件夹 ----------------
        moddirs = backup_mod_dirs(names) if include_mods else []
        renames = {}
        if include_mods:
            for md in moddirs:
                target = root / md
                if not target.exists():
                    st["mods_new"] += 1
                    continue
                if mode == "overwrite":
                    st["mods_overwritten"] += 1
                    if not dry_run:
                        if not send_to_recycle_bin(target):
                            shutil.rmtree(target, ignore_errors=True)
                elif mode == "rename":
                    new, i = md + " (备份恢复)", 2
                    while (root / new).exists() or new in renames.values():
                        new = "%s (备份恢复 %d)" % (md, i)
                        i += 1
                    renames[md] = new
                    st["mods_renamed"] += 1
                    st["renamed_list"].append((md, new))
                else:
                    renames[md] = None
                    st["mods_skipped"] += 1

        todo, total_bytes = [], 0
        if include_mods:
            for n in names:
                if n == BACKUP_MANIFEST or n.endswith("/") \
                        or not n.startswith(BACKUP_MOD_PREFIX):
                    continue
                rel = n[len(BACKUP_MOD_PREFIX):]
                top = owner_of(rel, moddirs)
                if top is not None:
                    if top in renames and renames[top] is None:
                        continue                      # 跳过：连同它的预览图一起不动
                    if renames.get(top):
                        rel = renames[top] + rel[len(top):]
                dst = root / Path(rel)
                if mode == "skip" and dst.exists():
                    continue          # 「跳过已存在」：连零散文件也不动，只补齐缺的
                todo.append((n, dst))
                total_bytes += infos[n].file_size
        done_b = 0
        for i, (n, dst) in enumerate(todo, 1):
            if should_cancel and should_cancel():
                raise BackupCancelled()
            if not dry_run:
                dst.parent.mkdir(parents=True, exist_ok=True)
                with zf.open(n) as src, open(str(dst), "wb") as fh:
                    shutil.copyfileobj(src, fh)
                st["files"] += 1
            done_b += infos[n].file_size
            if progress:
                progress(i, len(todo), done_b, total_bytes or 1, Path(n).name)
        st["bytes"] = done_b

        # ---------------- 索引库 / 汇总表 ----------------
        if include_meta and not dry_run:
            BACKUP_DIR.mkdir(parents=True, exist_ok=True)
            stamp = backup_stamp()
            for n in names:
                nm = Path(n).name
                if n.startswith(BACKUP_DB_PREFIX) and nm == DB_PATH.name:
                    if DB_PATH.exists():
                        try:
                            shutil.copy2(str(DB_PATH),
                                         str(BACKUP_DIR / ("mod_manager.%s.db" % stamp)))
                        except OSError:
                            pass
                    sidecar = Path(str(DB_PATH) + ".restored")
                    with zf.open(n) as src, open(str(sidecar), "wb") as fh:
                        shutil.copyfileobj(src, fh)
                    try:
                        st["db_rows"] = merge_db_from(sidecar)
                        st["db"] = True
                    finally:
                        try:
                            sidecar.unlink()
                        except OSError:
                            pass
                elif n.startswith(BACKUP_XL_PREFIX) and nm.lower().endswith(".xlsx"):
                    dst = excel_path(cfg)
                    dst.parent.mkdir(parents=True, exist_ok=True)
                    if dst.exists():
                        try:
                            shutil.copy2(str(dst), str(BACKUP_DIR / ("%s.%s.xlsx"
                                                                     % (dst.stem, stamp))))
                        except OSError:
                            pass
                    with zf.open(n) as src, open(str(dst), "wb") as fh:
                        shutil.copyfileobj(src, fh)
                    st["excel"] = True
    return st


# ------------------------------------------------------------------- commands
def cmd_scan(cfg, quiet=False):
    root = Path(cfg["root"])
    if not root.is_dir():
        raise SystemExit("Mod 根目录不存在：%s" % root)
    store = Store()
    found = scan_root(root)
    for m in found:
        m["img_hash"] = store.upsert(m)
    store.commit()
    pruned = store.prune({m["folder"] for m in found})
    store.commit()
    if not quiet:
        log("扫描完成：%d 个 Mod（清理失效记录 %d 条）" % (len(found), pruned))
        miss_addr = [m for m in found if not m["addr"]]
        miss_img = [m for m in found if not m["img"]]
        if miss_addr:
            log("  ! 缺地址 %d 个：%s" % (len(miss_addr), ", ".join(m["rel"] for m in miss_addr)))
        if miss_img:
            log("  ! 缺预览图 %d 个：%s" % (len(miss_img), ", ".join(m["rel"] for m in miss_img)))
    return found


def cmd_export(cfg):
    store = Store()
    mods = store.all()
    if not mods:
        cmd_scan(cfg, quiet=True)
        mods = store.all()
    out = Path(cfg["excel"])
    out.parent.mkdir(parents=True, exist_ok=True)
    hashes = {m["folder"]: m.get("img_hash") or "" for m in mods}
    n_cat, n_mod, n_img = write_excel(mods, cfg, out, hashes)
    size = out.stat().st_size / 1024 / 1024
    log("已生成 %s" % out)
    log("  工作表 %d 个 / 记录 %d 条 / 内嵌预览图 %d 张 / 文件 %.1f MB" % (n_cat, n_mod, n_img, size))
    return out


def cmd_check(cfg):
    store = Store()
    mods = store.all()
    root = cfg.get("root", "")

    def rel(m):
        try:
            return os.path.relpath(m["folder"], root)
        except Exception:
            return m["folder"]

    problems = 0
    for m in mods:
        issues = []
        if not Path(m["folder"]).is_dir():
            issues.append("文件夹不存在")
        if not m["addr"]:
            issues.append("缺 Mod地址")
        elif m.get("addr_source") != "地址.txt":
            issues.append("地址来自子目录（%s）" % m["addr_source"])
        if not m["img"]:
            issues.append("缺预览图")
        elif m.get("img_source") != "同级同名":
            issues.append("预览图取自%s" % m["img_source"])
        if issues:
            problems += 1
            log("  %-42s %s" % (rel(m)[:42], "；".join(issues)))
    dup = {}
    for m in mods:
        dup.setdefault((m["category"], m["seq"]), []).append(rel(m))
    for k, v in sorted(dup.items()):
        if len(v) > 1:
            problems += 1
            log("  ! 分类「%s」序号 %s 重复：%s" % (k[0], k[1], ", ".join(v)))
    log("检查完成：共 %d 条，%d 处需要关注" % (len(mods), problems))


def cmd_stats(cfg):
    store = Store()
    mods = store.all()
    log("Mod 索引： %s" % DB_PATH)
    by_cat = {}
    for m in mods:
        by_cat.setdefault(m["category"], []).append(m)
    for c in sorted(by_cat, key=lambda x: (cfg["category_order"].index(x) if x in cfg["category_order"] else 99, x)):
        rows = by_cat[c]
        nsfw = sum(1 for r in rows if r["nsfw"] == "NSFW")
        log("  %-6s %2d 条（NSFW %d / SFW %d）" % (c, len(rows), nsfw, len(rows) - nsfw))
    log("  合计   %d 条" % len(mods))


def cmd_add(cfg, src, category, zone="SFW", subdir="", author=None, name=None,
            seq=None, addr="", move=False, rescan=True, export=False):
    target = import_mod(cfg, src, category, zone, subdir, author, name, seq,
                        norm_addr(addr), move)
    log("已导入： %s" % target)
    if rescan:
        cmd_scan(cfg, quiet=True)
    if export:
        cmd_export(cfg)
    return target


def cmd_delete(cfg, folder, to_recycle=True, rescan=True, export=False):
    root = Path(cfg["root"]).resolve()
    f = Path(folder)
    if not f.is_absolute():
        f = root / f
    f = f.resolve()
    try:
        f.relative_to(root)                     # 只允许删 Mod 根目录内的东西
    except ValueError:
        raise SystemExit("出于安全考虑，只能删除 Mod 根目录下的内容：%s" % f)
    if f == root:
        raise SystemExit("不能删除 Mod 根目录本身")
    msg = delete_mod(f, to_recycle)
    log("%s：%s" % (msg, f))
    if rescan:
        cmd_scan(cfg, quiet=True)
    if export:
        cmd_export(cfg)
    return f


def cmd_edit(cfg, folder, category=None, zone=None, subdir=None, author=None, name=None,
             seq=None, addr=None, rescan=True, export=False):
    target = rename_mod(cfg, folder, category, zone, subdir, author, name, seq, addr)
    if rescan:
        cmd_scan(cfg, quiet=True)
    if export:
        cmd_export(cfg)
    return target


# ------------------------------------------------------------------------ gui
def gui(cfg):
    """图形界面。

    布局（为以后继续加功能预留了位置）：
        菜单栏      文件 / Mod / 视图 / 设置 / 帮助   —— 功能变多时优先往这里加
        工具栏      最常用的 6 个动作 + 「更多 ▾」
        路径行      一行显示 Mod 目录 / Excel，点「设置…」改
        左侧        分类·类型·搜索 筛选条 + 列表 + 统计
        右侧        预览图 + 字段详情 + 快捷操作
        状态栏      运行状态
    新增功能只要在下面的 ACTIONS 里加一行（菜单和工具栏会自动出现）。
    """
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    store = Store()
    win = tk.Tk()
    win.title("FFXIV Mod 管理工具")
    sw, sh = win.winfo_screenwidth(), win.winfo_screenheight()
    geo = str(cfg.get("window") or "")
    if re.match(r"^\d+x\d+[+-]-?\d+[+-]-?\d+$", geo):
        win.geometry(geo)
    else:
        _w, _h = min(1280, sw - 80), min(760, sh - 100)
        win.geometry("%dx%d+%d+%d" % (_w, _h, max(0, (sw - _w) // 2), max(0, (sh - _h) // 3)))
    win.minsize(900, 560)
    for cand in (Path(getattr(sys, "_MEIPASS", APP_DIR)) / "app.ico", APP_DIR / "app.ico"):
        if cand.is_file():
            try:
                win.iconbitmap(default=str(cand))
            except Exception:
                pass
            break
    if not (cfg.get("root") and Path(cfg["root"]).is_dir()):
        win.withdraw()                       # 首次使用：先让用户选 Mod 根目录
        if not first_run_setup(win, cfg):
            win.destroy()
            return
        win.deiconify()
    try:
        from PIL import Image, ImageTk
        HAVE_PIL = True
    except Exception:
        HAVE_PIL = False

    st = {"mods": [], "photo": None, "busy": False, "last": "", "tip": None,
          "inst_dir": "", "installed": {}}
    status_var = tk.StringVar(value="就绪")
    count_var = tk.StringVar(value="")
    v_search = tk.StringVar()
    v_cat = tk.StringVar(value="全部分类")
    v_sub = tk.StringVar(value="全部子分类")
    v_zone = tk.StringVar(value="全部")

    # ------------------------------------------------------------------ 通用
    def status(msg):
        status_var.set(msg)
        win.update_idletasks()

    def tip(widget, text):
        if not text:
            return

        def enter(_e):
            leave(_e)
            t = tk.Toplevel(widget)
            t.wm_overrideredirect(True)
            t.wm_geometry("+%d+%d" % (widget.winfo_rootx() + 10,
                                      widget.winfo_rooty() + widget.winfo_height() + 6))
            tk.Label(t, text=text, background="#ffffe0", relief="solid", borderwidth=1,
                     font=("Microsoft YaHei", 9), padx=6, pady=2).pack()
            st["tip"] = t

        def leave(_e):
            t = st.pop("tip", None)
            if t:
                try:
                    t.destroy()
                except Exception:
                    pass

        widget.bind("<Enter>", enter, add="+")
        widget.bind("<Leave>", leave, add="+")

    def place_dialog(dlg):
        """把对话框放到主窗口正中间，并保证完全落在屏幕内"""
        dlg.update_idletasks()
        w, h = dlg.winfo_reqwidth(), dlg.winfo_reqheight()
        scw, sch = dlg.winfo_screenwidth(), dlg.winfo_screenheight()
        try:
            px, py = win.winfo_rootx(), win.winfo_rooty()
            pw, ph = win.winfo_width(), win.winfo_height()
            if pw <= 1 or ph <= 1:
                px = py = 0
                pw, ph = scw, sch
        except Exception:
            px = py = 0
            pw, ph = scw, sch
        x = max(8, min(int(px + (pw - w) / 2), scw - w - 8))
        y = max(8, min(int(py + (ph - h) / 2), sch - h - 48))
        dlg.geometry("+%d+%d" % (x, y))
        dlg.lift()
        try:
            dlg.focus_force()
        except Exception:
            pass

    def run_task(title, fn, done=None):
        if st["busy"]:
            messagebox.showinfo("请稍候", "上一个任务还在运行")
            return
        st["busy"] = True
        status(title + "…")

        def worker():
            try:
                fn()
            except Exception:
                err = traceback.format_exc()
                log(err)
                win.after(0, lambda: messagebox.showerror("出错了", err[-1500:]))
            finally:
                st["busy"] = False

                def finish():
                    try:
                        if done:
                            done()
                    finally:
                        status("就绪")
                win.after(0, finish)

        import threading
        threading.Thread(target=worker, daemon=True).start()

    def brief(p, n=42):
        p = str(p)
        return p if len(p) <= n else "…" + p[-n:]

    # ------------------------------------------------------------- 列表相关
    def visible_mods():
        kw = v_search.get().strip().lower()
        cat, zone, sub = v_cat.get(), v_zone.get(), v_sub.get()
        out = []
        for m in st["mods"]:
            if cat != "全部分类" and m["category"] != cat:
                continue
            if sub != "全部子分类" and (m.get("subcat") or "") != sub:
                continue
            if zone != "全部" and m["nsfw"] != zone:
                continue
            if kw and kw not in ("%s %s %s" % (m["author"], m["name"], m["folder"])).lower():
                continue
            out.append(m)
        return out

    def fill_tree():
        keep = (tree.selection() or [None])[0]
        tree.delete(*tree.get_children())
        for m in visible_mods():
            ok = bool(st["installed"].get(m["folder"]))
            tree.insert("", "end", iid=m["folder"],
                        values=(m["category"], m.get("subcat") or "—", m["seq"],
                                m["author"], m["nsfw"], m["name"],
                                ("已安装" if ok else
                                 ("未安装" if st.get("inst_dir") else "?"))),
                        tags=("inst_yes" if ok else "inst_no",))
        if keep and tree.exists(keep):
            tree.selection_set(keep)

    def update_counts():
        by = {}
        for m in st["mods"]:
            by[m["category"]] = by.get(m["category"], 0) + 1
        order = cfg.get("category_order") or []
        cats = " / ".join("%s %d" % (c, by[c]) for c in
                          sorted(by, key=lambda x: (order.index(x) if x in order else 99, x)))
        count_var.set("共 %d 条 ｜ 显示 %d 条 ｜ %s" % (len(st["mods"]), len(visible_mods()), cats))

    def refresh_installed():
        try:
            names, raw, d = list_installed(cfg)
            st["inst_dir"] = d if (d and Path(d).is_dir()) else ""
            st["installed"] = {m["folder"]: installed_match(m, names, raw) for m in st["mods"]}
        except Exception:
            st["inst_dir"] = ""
            st["installed"] = {}

    def refresh(*_):
        st["mods"] = store.all()
        refresh_installed()
        cat_box["values"] = ["全部分类"] + cat_list()      # 含磁盘上还没放 Mod 的空分类
        cur_cat = v_cat.get()
        subs = sorted({(m.get("subcat") or "") for m in st["mods"]
                       if (m.get("subcat") or "")
                       and (cur_cat == "全部分类" or m["category"] == cur_cat)})
        sub_box["values"] = ["全部子分类"] + subs
        if v_sub.get() not in sub_box["values"]:
            v_sub.set("全部子分类")
        fill_tree()
        update_counts()
        return st["mods"]

    def selected_mod():
        sel = tree.selection()
        if not sel:
            return None
        return next((x for x in st["mods"] if x["folder"] == sel[0]), None)


    def cat_list():
        """所有分类：磁盘上的目录 + 索引库 + 配置顺序"""
        order = list(cfg.get("category_order") or [])
        cats = set(order) | {m["category"] for m in st["mods"]}
        try:
            cats |= {d.name for d in Path(cfg["root"]).iterdir() if d.is_dir()}
        except Exception:
            pass
        return sorted([c for c in cats if c],
                      key=lambda c: (order.index(c) if c in order else 99, c))

    def sub_list(cat):
        """某个分类下已有的子分类（如 皮肤 下的 纹身）"""
        out = set()
        try:
            for d in (Path(cfg["root"]) / (cat or "")).iterdir():
                if not d.is_dir():
                    continue
                if d.name.upper() in ("SFW", "NSFW"):
                    for e in d.iterdir():
                        if e.is_dir() and not MOD_RE.match(e.name):
                            out.add(e.name)
                elif not MOD_RE.match(d.name):
                    out.add(d.name)
        except Exception:
            pass
        out |= {(m.get("subcat") or "") for m in st["mods"] if m["category"] == cat}
        return sorted(x for x in out if x)

    # --------------------------------------------------------------- 详情面板
    def set_detail(m):
        if not m:
            for k in detail_vars:
                detail_vars[k].set("")
            img_label.configure(image="", text="← 选中左侧一条查看预览")
            st["photo"] = None
            return
        vals = {"分类": m["category"], "序号": str(m["seq"]), "作者": m["author"],
                "类型": m["nsfw"], "影响/替换": m.get("affects") or "—",
                "Mod 名称": m["name"],
                "Mod 地址": m["addr"] or "(缺失)", "文件夹": m["folder"]}
        for k, v in vals.items():
            detail_vars[k].set(v)
        detail_vars["预览图"].set(Path(m["img"]).name if m["img"] else "(缺失)")
        if HAVE_PIL and m["img"] and Path(m["img"]).is_file():
            try:
                im = Image.open(m["img"])
                im.thumbnail((460, 300))
                photo = ImageTk.PhotoImage(im)
                st["photo"] = photo
                img_label.configure(image=photo, text="")
            except Exception as e:
                img_label.configure(image="", text="预览加载失败：%s" % e)
        else:
            img_label.configure(image="", text="(无预览图)")

    def on_select(_evt=None):
        set_detail(selected_mod())

    def open_addr(_evt=None):
        m = selected_mod()
        if m and m["addr"]:
            import webbrowser
            webbrowser.open(m["addr"])

    def open_folder(_evt=None):
        m = selected_mod()
        if m:
            open_path(m["folder"])

    def copy_path():
        m = selected_mod()
        if m:
            win.clipboard_clear()
            win.clipboard_append(m["folder"])
            status("已复制路径：%s" % brief(m["folder"]))

    # ------------------------------------------------------------------ 动作
    def open_selected():
        m = selected_mod()
        if not m:
            messagebox.showinfo("提示", "请先在左边选一条 Mod")
            return
        subprocess.Popen(["explorer", "/select,", os.path.normpath(m["folder"])])

    def mod_subdir(m):
        if not m:
            return ""
        try:
            rel = Path(m["folder"]).relative_to(Path(cfg["root"])).parts
        except Exception:
            return ""
        return os.sep.join([q for q in rel[1:-1] if q.upper() not in ("SFW", "NSFW")])

    def do_scan():
        run_task("扫描中", lambda: cmd_scan(cfg, quiet=True),
                 lambda: (refresh(), status("扫描完成：共 %d 条" % len(st["mods"]))))

    def do_export():
        run_task("导出 Excel", lambda: cmd_export(cfg),
                 lambda: messagebox.showinfo("完成", "Excel 已生成：\n%s" % cfg["excel"]))

    def do_run():
        def work():
            cmd_scan(cfg, quiet=True)
            cmd_export(cfg)
        run_task("扫描并导出", work,
                 lambda: (refresh(), messagebox.showinfo("完成", "已扫描并重新生成 Excel：\n%s" % cfg["excel"])))

    def do_check():
        run_task("检查", lambda: cmd_check(cfg),
                 lambda: messagebox.showinfo("检查完成", "报告已写入：\n%s" % LOG_PATH))

    def do_delete():
        m = selected_mod()
        if not m:
            messagebox.showinfo("提示", "请先在左边选一条 Mod")
            return
        if not messagebox.askyesno("确认删除",
                                   "删除这条 Mod？\n\n名称：%s\n文件夹：%s\n\n"
                                   "文件夹会移入回收站，需要的话可以还原。"
                                   % (m["name"], m["folder"])):
            return
        run_task("删除中", lambda: cmd_delete(cfg, m["folder"], True, True, False),
                 lambda: (refresh(), messagebox.showinfo("完成", "已移入回收站：\n%s" % m["folder"])))

    def do_add(preset=None, on_done=None, export_default=None, local_cover=""):
        dlg = tk.Toplevel(win)
        dlg.title("添加 Mod" + ("（来自下载）" if preset else ""))
        dlg.transient(win)
        dlg.grab_set()
        cur = selected_mod()
        pre = dict(preset or {})
        v_src = tk.StringVar(value=pre.get("src", ""))
        v_cat2 = tk.StringVar(value=pre.get("category") or (cur["category"] if cur else
                              ((cfg.get("category_order") or [""])[0])))
        v_zone2 = tk.StringVar(value=pre.get("zone") or (cur["nsfw"] if cur else "SFW"))
        v_sub2 = tk.StringVar(value=pre.get("subdir") if pre else mod_subdir(cur))
        v_author2 = tk.StringVar(value=pre.get("author") or (cur["author"] if cur else ""))
        v_name2 = tk.StringVar(value=pre.get("name") or (cur["name"] if cur else ""))
        v_seq2 = tk.StringVar()
        v_addr2 = tk.StringVar(value=pre.get("addr", ""))
        v_mode2 = tk.StringVar(value="move" if pre.get("src") else "copy")
        v_exp2 = tk.BooleanVar(value=True if export_default is None else bool(export_default))
        v_cv2 = tk.BooleanVar(value=True)
        existing = cat_list()

        def target_parent():
            return mod_parent(cfg, v_cat2.get(), v_zone2.get(), v_sub2.get())

        def update_seq(*_):
            try:
                v_seq2.set(str(next_seq(target_parent())))
            except Exception:
                pass

        def fill_from_src(*_):
            src = v_src.get().strip()
            if src:
                q = Path(src)
                a, n = split_mod_name(q.name if q.is_dir() else q.stem)
                v_author2.set(a)
                v_name2.set(n)
            update_seq()

        def browse_dir():
            f = filedialog.askdirectory(parent=dlg, title="选择 Mod 文件夹")
            if f:
                v_src.set(os.path.normpath(f))
                fill_from_src()

        def browse_zip():
            f = filedialog.askopenfilename(parent=dlg, title="选择 Mod 压缩包",
                                           initialdir=resolve_dirs(cfg)[0],
                                           filetypes=[("Mod 压缩包", "*.zip *.7z *.rar *.pmp *.pcp"),
                                                      ("所有文件", "*.*")])
            if f:
                v_src.set(os.path.normpath(f))
                fill_from_src()

        fields = (("来源", v_src, "src"),
                  ("分类", v_cat2, "combo"),
                  ("类型", v_zone2, "zone"),
                  ("子分类目录（可留空，如 纹身）", v_sub2, "subcombo"),
                  ("作者", v_author2, "entry"),
                  ("Mod 名称", v_name2, "entry"),
                  ("序号（留空自动）", v_seq2, "entry"),
                  ("Mod 地址（可选，会写成 地址.txt）", v_addr2, "entry"))
        for r, (label, var, kind) in enumerate(fields):
            ttk.Label(dlg, text=label).grid(row=r, column=0, sticky="e", padx=8, pady=4)
            if kind == "combo":
                w = ttk.Combobox(dlg, textvariable=var, values=existing, width=42)
            elif kind == "zone":
                w = ttk.Combobox(dlg, textvariable=var, values=["SFW", "NSFW"], width=42, state="readonly")
            elif kind == "subcombo":
                w = ttk.Combobox(dlg, textvariable=var, values=sub_list(v_cat2.get()), width=42)
                sub_widget = w
            else:
                w = ttk.Entry(dlg, textvariable=var, width=44)
            w.grid(row=r, column=1, sticky="we", padx=8, pady=4)
            if kind == "src":
                bf2 = ttk.Frame(dlg)
                bf2.grid(row=r, column=2, padx=6)
                ttk.Button(bf2, text="选文件夹…", command=browse_dir).pack(side="left", padx=2)
                ttk.Button(bf2, text="选压缩包…", command=browse_zip).pack(side="left", padx=2)
        r = len(fields)
        ttk.Radiobutton(dlg, text="复制进来（原文件保留）", variable=v_mode2, value="copy")\
            .grid(row=r, column=1, sticky="w", padx=8)
        ttk.Radiobutton(dlg, text="移动进来（原位置不再保留）", variable=v_mode2, value="move")\
            .grid(row=r + 1, column=1, sticky="w", padx=8)
        ttk.Checkbutton(dlg, text="导入后自动重新生成 Excel", variable=v_exp2)\
            .grid(row=r + 2, column=1, sticky="w", padx=8, pady=(4, 0))
        ttk.Checkbutton(dlg, text="自动把封面存为该 Mod 的预览图", variable=v_cv2)\
            .grid(row=r + 3, column=1, sticky="w", padx=8, pady=(2, 0))
        dlg.columnconfigure(1, weight=1)

        def ok():
            if not v_src.get().strip():
                messagebox.showwarning("提示", "请选择要导入的文件夹或压缩包", parent=dlg)
                return
            if not v_cat2.get().strip() or not v_name2.get().strip():
                messagebox.showwarning("提示", "分类和 Mod 名称不能为空", parent=dlg)
                return
            seq = v_seq2.get().strip()
            addr_in = norm_addr(v_addr2.get())
            kw = dict(src=v_src.get().strip(), category=v_cat2.get().strip(), zone=v_zone2.get(),
                      subdir=v_sub2.get().strip(), author=v_author2.get().strip(),
                      name=v_name2.get().strip(), seq=int(seq) if seq.isdigit() else None,
                      addr=addr_in, move=(v_mode2.get() == "move"),
                      rescan=True, export=False)
            dlg.destroy()
            st["last"] = ""

            def work():
                st["last"] = str(cmd_add(cfg, **kw))
                st["preview"] = ""
                target = Path(st.get("last") or "")
                if v_cv2.get() and target.is_dir():
                    st["preview"] = ensure_preview(cfg, target, v_author2.get().strip(),
                                                   v_name2.get().strip(), local_cover, addr_in)
                if v_exp2.get() or st.get("preview"):
                    cmd_scan(cfg, quiet=True)
                    cmd_export(cfg)

            def finish():
                refresh()
                if on_done:
                    try:
                        on_done(st.get("last"))
                    except Exception:
                        pass
                msg = "已导入到：\n%s" % st.get("last")
                if st.get("preview"):
                    msg += "\n\n预览图已自动放好。"
                elif v_cv2.get():
                    msg += ("\n\n这个 Mod 暂时没有预览图：\n"
                            "· 内置浏览器开着、且地址里有 modid 时，工具会自动去抓；\n"
                            "· 也可以之后选中它点「补预览图」。")
                messagebox.showinfo("完成", msg)
            run_task("导入中", work, finish)

        bf = ttk.Frame(dlg)
        bf.grid(row=r + 4, column=0, columnspan=3, pady=10)
        ttk.Button(bf, text="确定", width=12, command=ok).pack(side="left", padx=8)
        ttk.Button(bf, text="取消", width=12, command=dlg.destroy).pack(side="left", padx=8)
        dlg.bind("<Return>", lambda *_: ok())
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        for var in (v_cat2, v_zone2, v_sub2):
            var.trace_add("write", update_seq)
        v_cat2.trace_add("write", lambda *_: sub_widget.configure(values=sub_list(v_cat2.get())))
        update_seq()
        place_dialog(dlg)

    def do_edit():
        m = selected_mod()
        if not m:
            messagebox.showinfo("提示", "请先在左边选一条 Mod")
            return
        dlg = tk.Toplevel(win)
        dlg.title("编辑 Mod")
        dlg.transient(win)
        dlg.grab_set()
        cur_parent = Path(m["folder"]).parent
        v_cat2 = tk.StringVar(value=m["category"])
        v_zone2 = tk.StringVar(value=m["nsfw"])
        v_sub2 = tk.StringVar(value=mod_subdir(m))
        v_author2 = tk.StringVar(value=m["author"])
        v_name2 = tk.StringVar(value=m["name"])
        v_seq2 = tk.StringVar(value=str(m["seq"]))
        v_addr2 = tk.StringVar(value=m["addr"] or "")
        v_exp2 = tk.BooleanVar(value=True)
        v_new = tk.StringVar()
        existing = cat_list()

        def target_parent():
            return mod_parent(cfg, v_cat2.get(), v_zone2.get(), v_sub2.get())

        def preview_target():
            try:
                q = target_parent() / mod_folder_name(
                    int(v_seq2.get()) if v_seq2.get().strip().isdigit() else m["seq"],
                    v_author2.get().strip(), v_name2.get().strip())
                v_new.set(str(q))
            except Exception:
                v_new.set("")

        def update_seq(*_):
            try:
                if target_parent().resolve() != cur_parent.resolve():
                    v_seq2.set(str(next_seq(target_parent())))
            except Exception:
                pass
            preview_target()

        ttk.Label(dlg, text="当前文件夹").grid(row=0, column=0, sticky="e", padx=8, pady=(10, 2))
        ttk.Label(dlg, text=m["folder"], wraplength=460, foreground="#555")\
            .grid(row=0, column=1, columnspan=2, sticky="w", padx=8, pady=(10, 2))
        fields = (("分类", v_cat2, "combo"),
                  ("类型", v_zone2, "zone"),
                  ("子分类目录（可留空，如 纹身）", v_sub2, "subcombo"),
                  ("作者", v_author2, "entry"),
                  ("Mod 名称", v_name2, "entry"),
                  ("序号", v_seq2, "entry"),
                  ("Mod 地址（清空 = 删除 地址.txt）", v_addr2, "entry"))
        for r, (label, var, kind) in enumerate(fields, start=1):
            ttk.Label(dlg, text=label).grid(row=r, column=0, sticky="e", padx=8, pady=4)
            if kind == "combo":
                w = ttk.Combobox(dlg, textvariable=var, values=existing, width=46)
            elif kind == "zone":
                w = ttk.Combobox(dlg, textvariable=var, values=["SFW", "NSFW"], width=46, state="readonly")
            elif kind == "subcombo":
                w = ttk.Combobox(dlg, textvariable=var, values=sub_list(v_cat2.get()), width=46)
                sub_widget = w
            else:
                w = ttk.Entry(dlg, textvariable=var, width=48)
            w.grid(row=r, column=1, columnspan=2, sticky="we", padx=8, pady=4)
        r = len(fields) + 1
        ttk.Label(dlg, text="将变为").grid(row=r, column=0, sticky="e", padx=8, pady=2)
        ttk.Label(dlg, textvariable=v_new, wraplength=460, foreground="#0a5")\
            .grid(row=r, column=1, columnspan=2, sticky="w", padx=8, pady=2)
        ttk.Checkbutton(dlg, text="保存后自动重新生成 Excel", variable=v_exp2)\
            .grid(row=r + 1, column=1, sticky="w", padx=8, pady=(4, 0))
        dlg.columnconfigure(1, weight=1)

        def ok():
            if not v_cat2.get().strip() or not v_name2.get().strip():
                messagebox.showwarning("提示", "分类和 Mod 名称不能为空", parent=dlg)
                return
            seq = v_seq2.get().strip()
            kw = dict(category=v_cat2.get().strip(), zone=v_zone2.get(),
                      subdir=v_sub2.get().strip(), author=v_author2.get().strip(),
                      name=v_name2.get().strip(), seq=int(seq) if seq.isdigit() else None,
                      addr=v_addr2.get())
            dlg.destroy()
            st["last"] = ""

            def work():
                st["last"] = str(cmd_edit(cfg, m["folder"], rescan=True,
                                          export=v_exp2.get(), **kw))
            run_task("保存中", work,
                     lambda: (refresh(), messagebox.showinfo("完成", "已更新：\n%s" % st.get("last"))))

        bf = ttk.Frame(dlg)
        bf.grid(row=r + 2, column=0, columnspan=3, pady=10)
        ttk.Button(bf, text="保存", width=12, command=ok).pack(side="left", padx=8)
        ttk.Button(bf, text="取消", width=12, command=dlg.destroy).pack(side="left", padx=8)
        dlg.bind("<Return>", lambda *_: ok())
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        for var in (v_cat2, v_zone2, v_sub2, v_author2, v_name2, v_seq2):
            var.trace_add("write", lambda *_: preview_target())
        v_zone2.trace_add("write", update_seq)
        v_cat2.trace_add("write", update_seq)
        v_sub2.trace_add("write", update_seq)
        v_cat2.trace_add("write", lambda *_: sub_widget.configure(values=sub_list(v_cat2.get())))
        preview_target()
        place_dialog(dlg)

    # ------------------------------------------------------------ 下载 / 网站
    def do_open_site():
        import webbrowser
        webbrowser.open("https://www.xivmodarchive.com/")

    def do_open_downloads():
        open_path(resolve_dirs(cfg)[0])

    def do_open_inbox():
        q = Path(resolve_dirs(cfg)[1])
        q.mkdir(parents=True, exist_ok=True)
        open_path(q)

    # ------------------------------------------------------------ 下载工作台
    def do_browser_open():
        try:
            launch_browser(cfg)
            status("内置浏览器已启动；首次需要在那个窗口里登录一次")
        except SystemExit as e:
            messagebox.showerror("打开浏览器失败", str(e))

    def browser_next_page():
        try:
            cdp = browser_connect(cfg)
            try:
                r = cdp.js(NEXT_PAGE_JS)
            finally:
                cdp.close()
            status("已翻到下一页" if r else "没找到「下一页」按钮")
        except SystemExit as e:
            messagebox.showinfo("提示", str(e))

    def do_workbench():
        dlg = tk.Toplevel(win)
        dlg.title("下载工作台 —— 内置浏览器 + 下载监视")
        dlg.transient(win)
        st2 = {"rows": [], "cap": {}, "cover": "", "photo": None, "dl_map": {}, "last_dl": "", "cancel": False}
        v_dl = tk.StringVar(value=resolve_dirs(cfg)[0])
        v_ib = tk.StringVar(value=resolve_dirs(cfg)[1])
        v_all = tk.BooleanVar(value=False)
        v_use = tk.BooleanVar(value=True)
        v_cap = tk.StringVar(value="（还没抓取。请在浏览器里打开某个 Mod 页面，再点「抓取当前页面」）")
        hint_var = tk.StringVar(value="")

        # ---------------- 处理函数（先定义，按钮要用） ----------------
        def show_cover():
            p = st2.get("cover")
            if p and Path(p).is_file() and HAVE_PIL:
                try:
                    im = Image.open(p)
                    im.thumbnail((300, 170))
                    st2["photo"] = ImageTk.PhotoImage(im)
                    cover_label.configure(image=st2["photo"], text="")
                    return
                except Exception:
                    pass
            cover_label.configure(image="", text="（无封面）")

        def save_dirs():
            cfg["download_dir"] = v_dl.get().strip()
            cfg["inbox_dir"] = v_ib.get().strip()
            save_config(cfg)

        def cap_of(r):
            m = st2["dl_map"].get(r["file"]) or {}
            return m.get("cover") or ""

        def info_for_row(r):
            m = dict(st2["dl_map"].get(r["file"]) or {})
            if not m.get("addr") and st2.get("cap"):
                m = dict(st2["cap"])
                m["cover"] = st2.get("cover", "")
            return m

        def load(*_):
            save_dirs()
            st2["rows"] = list_pending(cfg, all_files=v_all.get())
            tv.delete(*tv.get_children())
            for i, r in enumerate(st2["rows"]):
                tv.insert("", "end", iid=str(i),
                          values=(r["where"], r["file"],
                                  "✔" if cap_of(r) else "",
                                  fmt_size(r["size"]),
                                  _dt.datetime.fromtimestamp(r["mtime"]).strftime("%m-%d %H:%M"),
                                  r["author"], r["name"]))
            hint_var.set("发现 %d 个待导入文件" % len(st2["rows"]))

        def show_cap():
            c = st2.get("cap") or {}
            if c.get("challenge"):
                v_cap.set("页面还在 Cloudflare 人机验证中。\n"
                          "请在浏览器窗口里等几秒（或点一下验证框），然后再点「抓取当前页面」。")
                cover_label.configure(image="", text="（无封面）")
                return
            if c.get("is_mod"):
                v_cap.set("modid/%s　作者：%s　名称：%s\n%s%s"
                          % (c.get("modid"), c.get("author") or "-", c.get("name") or "-",
                             c.get("url", ""),
                             "\n（封面已抓取）" if st2.get("cover") else "\n（这个页面没抓到封面）"))
            else:
                v_cap.set("当前页面不是 Mod 详情页：%s\n（请点进某个 Mod 再抓取）" % c.get("url", ""))
            show_cover()

        def grab():
            def work():
                browser_wait_ready(cfg, timeout=30,
                                   on_tick=lambda s: status("等待页面打开/人机验证… %ds" % s))
                st2["cap"] = browser_capture(cfg)
                c = st2["cap"]
                st2["cover"] = ""
                if c.get("cover"):
                    dest = APP_DIR / "cover_cache" / ("mod_%s.jpg" % (c.get("modid") or "tmp"))
                    try:
                        browser_fetch(cfg, c["cover"], dest)
                        st2["cover"] = str(dest)
                    except Exception as e:
                        log("封面抓取失败：%s" % e)

            run_task("抓取页面", work, show_cap)

        def cancel_dl():
            st2["cancel"] = True
            status("正在取消…")

        def dl_mod():
            def work():
                st2["cancel"] = False
                status("读取当前页面…")
                cap = {}
                try:                                   # 顺手把 地址/名称/作者/封面 一起抓下来
                    browser_wait_ready(cfg, timeout=25,
                                       on_tick=lambda x: status("等待页面/人机验证… %d 秒" % x))
                    cap = browser_capture(cfg)
                    st2["cap"] = cap
                    st2["cover"] = ""
                    if cap.get("cover"):
                        dest = APP_DIR / "cover_cache" / ("mod_%s.jpg" % (cap.get("modid") or "tmp"))
                        try:
                            browser_fetch(cfg, cap["cover"], dest)
                            st2["cover"] = str(dest)
                        except Exception as e:
                            log("封面抓取失败：%s" % e)
                except SystemExit as e:
                    log("抓取页面失败（继续下载）：%s" % e)
                status("下载中…")
                href, path = browser_download(cfg, resolve_dirs(cfg)[1],
                                              on_tick=lambda x: status("下载中… 已等 %d 秒（点「取消下载」可中止）" % x),
                                              should_cancel=lambda: st2.get("cancel"))
                st2["last_dl"] = path
                if path:
                    st2["dl_map"][Path(path).name] = {
                        "addr": cap.get("addr", ""), "name": cap.get("name", ""),
                        "author": cap.get("author", ""), "cover": st2.get("cover", "")}

            def done():
                load()
                show_cap()
                if st2.get("cancel"):
                    messagebox.showinfo("已取消", "已停止等待下载。\n"
                                                  "（浏览器里已经开始的这一次下载可能还在继续，"
                                                  "下载目录里会留一个没下完的临时文件，可以手动删掉）")
                    return
                p2 = st2.get("last_dl")
                if p2:
                    c = st2.get("cap") or {}
                    extra = ""
                    if c.get("addr"):
                        extra = "\n\n已记下本条 Mod 的地址%s。" % ("和封面图" if st2.get("cover") else "")
                    messagebox.showinfo("下载完成",
                                        "已下载并放入暂存目录：\n%s%s\n\n"
                                        "现在选中它点「导入选中…」，地址和封面会自动带上。" % (p2, extra))
                else:
                    messagebox.showwarning("没等到文件",
                                           "下载可能还没完成，或没触发。\n"
                                           "可以过一会儿点「刷新」；大文件请耐心等（状态栏有计时）。")
            run_task("下载中", work, done)

        def cur_row():
            sel = tv.selection()
            return st2["rows"][int(sel[0])] if sel else None

        def do_import():
            r = cur_row()
            if not r:
                messagebox.showinfo("提示", "请先选一个文件", parent=dlg)
                return
            info = info_for_row(r) if v_use.get() else {}
            use = bool(info.get("addr"))
            cover = info.get("cover") or ""
            preset = {"src": r["src"], "author": r["author"], "name": r["name"]}
            if use:
                preset["addr"] = info["addr"]
                if info.get("name"):
                    preset["name"] = info["name"]
                if info.get("author"):
                    preset["author"] = info["author"]

            def after(_target):
                load()
            do_add(preset=preset, on_done=after, export_default=True, local_cover=cover)

        def open_src():
            r = cur_row()
            if r:
                subprocess.Popen(["explorer", "/select,", os.path.normpath(r["src"])])

        def del_src():
            r = cur_row()
            if not r:
                return
            if messagebox.askyesno("确认", "删除这个文件（移入回收站）？\n%s" % r["file"], parent=dlg):
                if send_to_recycle_bin(r["src"]):
                    load()

        # ---------------- 浏览器区 ----------------
        brow = ttk.LabelFrame(dlg, text="内置浏览器（真实 Edge/Chrome，站点的人机验证能正常通过）")
        brow.pack(fill="x", padx=8, pady=(8, 4))
        r1 = ttk.Frame(brow)
        r1.pack(fill="x", padx=6, pady=(6, 2))
        for text, fn, tt in (("打开/显示浏览器", do_browser_open, "首次需要在那个窗口里登录一次，之后保持登录"),
                             ("抓取当前页面", grab, "读取网址 / 名称 / 作者 / 封面 / 下载直链"),
                             ("下载到暂存区", dl_mod, "让浏览器把当前 Mod 的文件下载到暂存目录"),
                             ("下一页", browser_next_page, "浏览列表翻页"),
                             ("取消下载", cancel_dl, "中止等待下载（不影响浏览器里已开始的那次）"),
                             ("打开暂存文件夹", do_open_inbox, "")):
            b = ttk.Button(r1, text=text, command=fn)
            b.pack(side="left", padx=(0, 6))
            tip(b, tt)
        r2 = ttk.Frame(brow)
        r2.pack(fill="x", padx=6, pady=(2, 2))
        ttk.Label(r2, textvariable=v_cap, wraplength=640, justify="left").pack(side="left", anchor="w")
        cover_label = ttk.Label(r2, text="（无封面）", anchor="center")
        cover_label.pack(side="right", padx=6)
        r3 = ttk.Frame(brow)
        r3.pack(fill="x", padx=6, pady=(0, 6))
        ttk.Checkbutton(r3, text="导入时带上刚抓取的「地址 + 封面图」", variable=v_use).pack(side="left")
        tip(r3.winfo_children()[-1], "勾上后：自动填 Mod 地址、并把封面图作为该 Mod 的预览图")

        # ---------------- 监视目录 ----------------
        top = ttk.LabelFrame(dlg, text="监视的文件夹（下载目录 / 暂存目录）")
        top.pack(fill="x", padx=8, pady=(4, 2))
        for i, (label, var) in enumerate((("下载目录", v_dl), ("暂存目录", v_ib))):
            ttk.Label(top, text=label).grid(row=i, column=0, sticky="e", padx=6, pady=3)
            ttk.Entry(top, textvariable=var, width=62).grid(row=i, column=1, sticky="we", padx=6)
            ttk.Button(top, text="浏览…", width=8,
                       command=lambda v=var: v.set(filedialog.askdirectory(parent=dlg) or v.get()))\
                .grid(row=i, column=2, padx=6)
        ttk.Checkbutton(top, text="显示所有类型的文件（不止 .pmp/.zip/.7z/.rar/.ttmp2）",
                        variable=v_all, command=lambda: load())\
            .grid(row=2, column=1, sticky="w", padx=6, pady=(0, 4))
        top.columnconfigure(1, weight=1)

        # ---------------- 待导入列表 ----------------
        bar = ttk.Frame(dlg)                      # 底部按钮先摆，免得被列表挤走
        bar.pack(side="bottom", fill="x", padx=8, pady=(2, 8))
        for text, fn, tt in (("刷新", load, "重新扫描这两个文件夹"),
                             ("导入选中…", do_import, "选分类/类型后导入到库里"),
                             ("打开所在文件夹", open_src, ""),
                             ("删除文件", del_src, "移入回收站")):
            b = ttk.Button(bar, text=text, command=fn)
            b.pack(side="left", padx=(0, 6))
            tip(b, tt)
        ttk.Label(bar, textvariable=hint_var, foreground="#555").pack(side="right")

        body = ttk.Frame(dlg)
        body.pack(fill="both", expand=True, padx=8, pady=4)
        cols = ("where", "file", "cap", "size", "time", "author", "name")
        tv = ttk.Treeview(body, columns=cols, show="headings", selectmode="browse")
        for c, t, w, anch in (("where", "来源", 56, "center"), ("file", "文件", 280, "w"),
                              ("cap", "封面", 52, "center"),
                              ("size", "大小", 76, "center"), ("time", "修改时间", 120, "center"),
                              ("author", "作者", 110, "w"), ("name", "识别出的名称", 230, "w")):
            tv.heading(c, text=t)
            tv.column(c, width=w, anchor=anch, stretch=(c in ("file", "name")))
        ys = ttk.Scrollbar(body, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=ys.set)
        tv.pack(side="left", fill="both", expand=True)
        ys.pack(side="right", fill="y")


        dlg.geometry("1080x720")
        dlg.minsize(880, 520)
        place_dialog(dlg)
        load()

    # ----------------------------------------------- 序号重排 / 查重 / 安装检查
    def do_renumber():
        dlg = tk.Toplevel(win)
        dlg.title("自动分配序号")
        dlg.transient(win)
        dlg.grab_set()
        v_c = tk.StringVar(value=("全部分类" if v_cat.get() == "全部分类" else v_cat.get()))
        v_s = tk.StringVar(value=v_sub.get() if v_sub.get() != "全部子分类" else "")
        v_st = tk.StringVar(value="1")
        cats = ["全部分类"] + cat_list()

        top = ttk.Frame(dlg)
        top.pack(fill="x", padx=10, pady=(10, 4))
        ttk.Label(top, text="分类").pack(side="left")
        cb1 = ttk.Combobox(top, textvariable=v_c, values=cats, width=12, state="readonly")
        cb1.pack(side="left", padx=(4, 10))
        ttk.Label(top, text="子分类").pack(side="left")
        cb2 = ttk.Combobox(top, textvariable=v_s, width=12)
        cb2.pack(side="left", padx=(4, 10))
        ttk.Label(top, text="起始序号").pack(side="left")
        ttk.Entry(top, textvariable=v_st, width=6).pack(side="left", padx=4)

        def upd_subs(*_):
            cb2.configure(values=[""] + sub_list(v_c.get()) if v_c.get() != "全部分类" else [""])
        cb1.bind("<<ComboboxSelected>>", upd_subs)
        upd_subs()

        tv = ttk.Treeview(dlg, columns=("old", "new", "name"), show="headings", height=12)
        for c, t, w in (("old", "原序号", 70), ("new", "新序号", 70), ("name", "分类 / Mod 名称 / 作者", 430)):
            tv.heading(c, text=t)
            tv.column(c, width=w, anchor="center" if c != "name" else "w")
        tv.pack(fill="both", expand=True, padx=10, pady=4)
        info = tk.StringVar(value="点「预览」看看要改哪些")
        ttk.Label(dlg, textvariable=info, foreground="#444").pack(fill="x", padx=12)

        plan_holder = {"plan": []}

        def preview():
            try:
                start = int(v_st.get().strip() or 1)
            except ValueError:
                start = 1
            plan = renumber_plan(cfg, v_c.get(), v_s.get().strip(), start)
            plan_holder["plan"] = plan
            tv.delete(*tv.get_children())
            for m, n in plan:
                tv.insert("", "end", values=(m["seq"], n, "[%s] %s（%s）"
                                             % (m["category"], m["name"], m["author"] or "-")))
            info.set("范围：%s%s ｜ 需要改 %d 条%s"
                     % (v_c.get(), ("\\" + v_s.get().strip()) if v_s.get().strip() else "",
                        len(plan),
                        "（已经连续，不用改）" if not plan
                        else "；每个分类各自从 %d 开始编号" % start))

        def apply():
            plan = plan_holder["plan"]
            if not plan:
                messagebox.showinfo("提示", "先点「预览」，没有需要改的就直接关掉。", parent=dlg)
                return
            if not messagebox.askyesno("确认", "按上面的对照把 %d 个文件夹改名？\n"
                                              "（只是改文件夹名，Mod 内容不动）" % len(plan), parent=dlg):
                return
            try:
                start = int(v_st.get().strip() or 1)
            except ValueError:
                start = 1
            cat, sub = v_c.get(), v_s.get().strip()
            dlg.destroy()

            def work():
                renumber(cfg, cat, sub, start, dry_run=False, export=True)
            run_task("重排序号", work,
                     lambda: (refresh(), messagebox.showinfo(
                         "完成", "序号已重排并重新生成了 Excel。")))

        bar = ttk.Frame(dlg)
        bar.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(bar, text="预览", width=10, command=preview).pack(side="left", padx=(0, 6))
        ttk.Button(bar, text="应用", width=10, command=apply).pack(side="left")
        ttk.Button(bar, text="关闭", width=10, command=dlg.destroy).pack(side="right")
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        dlg.geometry("640x460")
        place_dialog(dlg)
        preview()

    def do_dupes():
        dlg = tk.Toplevel(win)
        dlg.title("重复 Mod 检查")
        dlg.transient(win)
        groups = find_duplicates(st["mods"])

        # 先摆底部按钮和说明，再放列表——反过来按钮会被挤到旁边
        bar = ttk.Frame(dlg)
        bar.pack(side="bottom", fill="x", padx=10, pady=(0, 10))
        ttk.Label(dlg, foreground="#666", justify="left",
                  text="按「同一个 Mod 地址」和「作者+名称相同」找重复；%s"
                       % ("没有发现重复 ✅" if not groups
                          else "共 %d 组，请人工确认后处理" % len(groups)))\
            .pack(side="bottom", fill="x", padx=12, pady=(0, 6))

        body = ttk.Frame(dlg)
        body.pack(fill="both", expand=True, padx=10, pady=10)
        tv = ttk.Treeview(body, columns=("info",), show="tree headings", selectmode="browse")
        tv.heading("#0", text="组 / Mod")
        tv.column("#0", width=430)
        tv.heading("info", text="说明")
        tv.column("info", width=260, anchor="w")
        ys = ttk.Scrollbar(body, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=ys.set)
        tv.pack(side="left", fill="both", expand=True)
        ys.pack(side="right", fill="y")

        mapping = {}
        for i, (why, info_txt, g) in enumerate(groups, 1):
            node = tv.insert("", "end", text="第 %d 组：%s" % (i, why),
                             values=(info_txt[:60],), tags=("grp",), open=True)
            for m in g:
                iid = tv.insert(node, "end", text="   %s（%s）" % (m["name"], m["author"] or "-"),
                                values=(os.path.relpath(m["folder"], cfg["root"])[:60],),
                                tags=("mod",))
                mapping[iid] = m

        def sel_mod():
            s2 = tv.selection()
            return mapping.get(s2[0]) if s2 else None

        def open_one():
            m = sel_mod()
            if m:
                subprocess.Popen(["explorer", "/select,", os.path.normpath(m["folder"])])

        def del_one():
            m = sel_mod()
            if not m:
                messagebox.showinfo("提示", "请选中组里的某一条 Mod", parent=dlg)
                return
            if not messagebox.askyesno("确认", "把「%s」移入回收站？" % m["name"], parent=dlg):
                return

            def work():
                cmd_delete(cfg, m["folder"], to_recycle=True, rescan=True, export=False)
            run_task("删除中", work,
                     lambda: (refresh(), dlg.destroy(),
                              messagebox.showinfo("完成", "已移入回收站。再开一次「查重」看结果。")))

        ttk.Button(bar, text="打开所在文件夹", command=open_one).pack(side="left", padx=(0, 6))
        ttk.Button(bar, text="删除选中（回收站）", command=del_one).pack(side="left", padx=(0, 6))
        ttk.Button(bar, text="关闭", command=dlg.destroy).pack(side="right")
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        dlg.geometry("780x500")
        place_dialog(dlg)

    def do_installed():
        dlg = tk.Toplevel(win)
        dlg.title("安装检查")
        dlg.transient(win)
        v_dir = tk.StringVar(value=find_install_dir(cfg) or "")
        only_missing = tk.BooleanVar(value=False)
        sum_var = tk.StringVar(value="")

        top = ttk.Frame(dlg)
        top.pack(side="top", fill="x", padx=10, pady=(10, 4))
        ttk.Label(top, text="安装目录（Penumbra 的 Mods 目录）").pack(side="left")
        ttk.Entry(top, textvariable=v_dir, width=52).pack(side="left", padx=6, fill="x", expand=True)
        ttk.Button(top, text="浏览…", command=lambda: v_dir.set(
            filedialog.askdirectory(parent=dlg) or v_dir.get())).pack(side="left")

        # 底部按钮先摆
        bar = ttk.Frame(dlg)
        bar.pack(side="bottom", fill="x", padx=10, pady=(0, 10))
        ttk.Label(dlg, foreground="#666", justify="left",
                  text="判断方式：拿安装目录里的文件夹名和 Mod 名称/作者比对"
                       "（支持「名称_作者」这种 Penumbra 命名）。")\
            .pack(side="bottom", fill="x", padx=12, pady=(0, 6))

        body = ttk.Frame(dlg)
        body.pack(fill="both", expand=True, padx=10, pady=6)
        cols = ("st", "cat", "seq", "author", "name")
        tv = ttk.Treeview(body, columns=cols, show="headings", selectmode="browse")
        for c, t, w, a in (("st", "是否安装", 80, "center"), ("cat", "分类", 66, "center"),
                           ("seq", "序号", 46, "center"), ("author", "作者", 100, "w"),
                           ("name", "Mod 名称", 320, "w")):
            tv.heading(c, text=t)
            tv.column(c, width=w, anchor=a, stretch=(c == "name"))
        tv.tag_configure("inst_yes", background="#D8F5D8")
        tv.tag_configure("inst_no", background="#FFFFFF")
        ys = ttk.Scrollbar(body, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=ys.set)
        tv.pack(side="left", fill="both", expand=True)
        ys.pack(side="right", fill="y")
        mapping = {}

        def load():
            d = v_dir.get().strip()
            cfg["install_dir"] = d
            save_config(cfg)
            names, raw, _ = list_installed(cfg)
            tv.delete(*tv.get_children())
            mapping.clear()
            shown = ins = 0
            missing = []
            for m in st["mods"]:
                ok = installed_match(m, names, raw)
                ins += 1 if ok else 0
                if not ok:
                    missing.append("　".join(x for x in (m["category"], m["name"], m["author"]) if x))
                if only_missing.get() and ok:
                    continue
                iid = tv.insert("", "end", values=("已安装" if ok else "未安装",
                                                   m["category"], m["seq"],
                                                   m["author"], m["name"]),
                                tags=("inst_yes" if ok else "inst_no",))
                mapping[iid] = m
                shown += 1
            st["inst_dir"] = d if (d and Path(d).is_dir()) else ""
            st["installed"] = {m["folder"]: installed_match(m, names, raw) for m in st["mods"]}
            fill_tree()
            sum_var.set("目录里 %d 个文件夹 ｜ 已安装 %d ｜ 未安装 %d ｜ 共 %d 条%s"
                        % (len(raw), ins, len(st["mods"]) - ins, len(st["mods"]),
                           "" if d and Path(d).is_dir() else "（目录无效，请重新选择）"))
            dlg.missing = missing

        def open_one():
            s2 = tv.selection()
            if s2 and s2[0] in mapping:
                subprocess.Popen(["explorer", "/select,",
                                  os.path.normpath(mapping[s2[0]]["folder"])])

        def copy_missing():
            miss = getattr(dlg, "missing", [])
            win.clipboard_clear()
            win.clipboard_append("\n".join(miss))
            messagebox.showinfo("已复制", "未安装清单（%d 条）已复制到剪贴板。" % len(miss), parent=dlg)

        def open_dir():
            d = v_dir.get().strip()
            if d and Path(d).is_dir():
                open_path(d)

        for text, fn in (("重新扫描", load), ("打开安装目录", open_dir),
                         ("复制未安装清单", copy_missing), ("打开所在文件夹", open_one)):
            ttk.Button(bar, text=text, command=fn).pack(side="left", padx=(0, 6))
        ttk.Checkbutton(bar, text="只看未安装", variable=only_missing,
                        command=load).pack(side="left", padx=8)
        ttk.Label(bar, textvariable=sum_var, foreground="#444").pack(side="right")
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        dlg.geometry("900x540")
        place_dialog(dlg)
        load()

    def do_categories():
        """分类 / 子分类管理：新建、重命名、删除（空目录）、排序"""
        import tkinter as tk
        from tkinter import messagebox, simpledialog, ttk

        dlg = tk.Toplevel(win)
        dlg.title("分类管理")
        dlg.transient(win)
        dlg.grab_set()
        tv = ttk.Treeview(dlg, columns=("kind", "n"), show="tree headings", selectmode="browse")
        tv.heading("#0", text="分类 / 子分类")
        tv.column("#0", width=300)
        tv.heading("kind", text="类型")
        tv.column("kind", width=70, anchor="center")
        tv.heading("n", text="数量")
        tv.column("n", width=60, anchor="center")
        ys = ttk.Scrollbar(dlg, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=ys.set)
        tv.pack(side="top", fill="both", expand=True, padx=8, pady=(8, 4))
        ys.place(in_=tv, relx=1.0, rely=0, relheight=1.0, anchor="ne")

        def save_order(order):
            cfg["category_order"] = order
            save_config(cfg)

        def cur():
            sel = tv.selection()
            if not sel:
                return None, None
            tags = tv.item(sel[0], "tags") or ()
            txt = tv.item(sel[0], "text")
            if "cat" in tags:
                return ("cat", txt)
            if "sub" in tags:
                parent = tv.parent(sel[0])
                return ("sub", txt, tv.item(parent, "text"))
            return None, None

        def load(*_):
            tv.delete(*tv.get_children())
            for c in cat_list():
                n = sum(1 for m in st["mods"] if m["category"] == c)
                node = tv.insert("", "end", text=c, values=("分类", n), tags=("cat",), open=True)
                for sc in sub_list(c):
                    n2 = sum(1 for m in st["mods"]
                             if m["category"] == c and (m.get("subcat") or "") == sc)
                    tv.insert(node, "end", text=sc, values=("子分类", n2), tags=("sub",))

        def new_cat():
            name = simpledialog.askstring("新建分类", "分类名（就是 Mod 根目录下的文件夹名）：", parent=dlg)
            if not name:
                return
            name = name.strip()
            if not name or any(x in name for x in '\\/:*?"<>|'):
                messagebox.showwarning("提示", "名字不合法", parent=dlg)
                return
            p = Path(cfg["root"]) / name
            if p.exists():
                messagebox.showinfo("提示", "这个分类已经存在了", parent=dlg)
                return
            try:
                p.mkdir(parents=True)
                for z in ("SFW", "NSFW"):
                    (p / z).mkdir(exist_ok=True)
            except Exception as e:
                messagebox.showerror("出错", str(e), parent=dlg)
                return
            order = list(cfg.get("category_order") or [])
            if name not in order:
                order.append(name)
            save_order(order)
            cmd_scan(cfg, quiet=True)
            refresh()
            load()
            status("已新建分类：%s" % name)

        def new_sub():
            kind = cur()
            if not kind or kind[0] != "cat":
                messagebox.showinfo("提示", "请先选中一个分类，再新建子分类", parent=dlg)
                return
            cat = kind[1]
            name = simpledialog.askstring("新建子分类",
                                          "子分类名（会建在 %s 下，如 纹身）：" % cat, parent=dlg)
            if not name:
                return
            name = name.strip()
            if not name or any(x in name for x in '\\/:*?"<>|'):
                messagebox.showwarning("提示", "名字不合法", parent=dlg)
                return
            try:
                (Path(cfg["root"]) / cat / name).mkdir(parents=True, exist_ok=True)
            except Exception as e:
                messagebox.showerror("出错", str(e), parent=dlg)
                return
            cmd_scan(cfg, quiet=True)
            refresh()
            load()
            status("已新建子分类：%s\\%s（默认是 SFW，可在编辑里改成 NSFW）" % (cat, name))

        def rename():
            kind = cur()
            if not kind:
                messagebox.showinfo("提示", "请先选中一个分类或子分类", parent=dlg)
                return
            if kind[0] == "cat":
                old = Path(cfg["root"]) / kind[1]
            else:
                old = Path(cfg["root"]) / kind[2] / kind[1]
            new = simpledialog.askstring("重命名", "新名字：", initialvalue=old.name, parent=dlg)
            if not new or new.strip() == old.name:
                return
            new = new.strip()
            if any(x in new for x in '\\/:*?"<>|'):
                messagebox.showwarning("提示", "名字不合法", parent=dlg)
                return
            target = old.parent / new
            if target.exists():
                messagebox.showinfo("提示", "已经有了这个名字", parent=dlg)
                return
            try:
                old.rename(target)
            except Exception as e:
                messagebox.showerror("出错", str(e), parent=dlg)
                return
            if kind[0] == "cat":
                order = [new if x == kind[1] else x for x in (cfg.get("category_order") or [])]
                save_order(order)
            cmd_scan(cfg, quiet=True)
            refresh()
            load()
            status("已重命名")

        def delete():
            kind = cur()
            if not kind:
                messagebox.showinfo("提示", "请先选中一个分类或子分类", parent=dlg)
                return
            if kind[0] == "cat":
                p2 = Path(cfg["root"]) / kind[1]
                n = sum(1 for m in st["mods"] if m["category"] == kind[1])
            else:
                p2 = Path(cfg["root"]) / kind[2] / kind[1]
                n = sum(1 for m in st["mods"]
                        if m["category"] == kind[2] and (m.get("subcat") or "") == kind[1])
            if n:
                messagebox.showwarning("不能删除",
                                       "这个%s下面还有 %d 个 Mod。\n"
                                       "请先把它们移走或删掉，再来删目录。"
                                       % ("分类" if kind[0] == "cat" else "子分类", n), parent=dlg)
                return
            if not messagebox.askyesno("确认", "把空目录「%s」移入回收站？" % p2.name, parent=dlg):
                return
            if send_to_recycle_bin(p2):
                if kind[0] == "cat":
                    save_order([x for x in (cfg.get("category_order") or []) if x != kind[1]])
                cmd_scan(cfg, quiet=True)
                refresh()
                load()
                status("已删除空目录：%s" % p2.name)

        def move(delta):
            kind = cur()
            if not kind or kind[0] != "cat":
                messagebox.showinfo("提示", "请选中一个分类（排序只对分类有效）", parent=dlg)
                return
            order = list(cfg.get("category_order") or [])
            for c in cat_list():                     # 保证所有分类都在列表里
                if c not in order:
                    order.append(c)
            i = order.index(kind[1])
            j = max(0, min(len(order) - 1, i + delta))
            if i == j:
                return
            order.insert(j, order.pop(i))
            save_order(order)
            load()
            status("已调整顺序（影响下拉框和 Excel 工作表顺序）")

        def open_cat():
            kind = cur()
            if not kind:
                return
            p2 = (Path(cfg["root"]) / kind[1]) if kind[0] == "cat" \
                else (Path(cfg["root"]) / kind[2] / kind[1])
            if p2.is_dir():
                open_path(p2)

        bar = ttk.Frame(dlg)
        bar.pack(side="bottom", fill="x", padx=8, pady=(4, 8))
        row1 = ttk.Frame(bar)
        row1.pack(fill="x", pady=(0, 4))
        row2 = ttk.Frame(bar)
        row2.pack(fill="x")
        for text, fn in (("新建分类", new_cat), ("新建子分类", new_sub),
                         ("重命名", rename), ("删除（空）", delete)):
            ttk.Button(row1, text=text, width=12, command=fn).pack(side="left", padx=(0, 6))
        for text, fn in (("上移", lambda: move(-1)), ("下移", lambda: move(1)),
                         ("打开所在文件夹", open_cat), ("刷新", load)):
            ttk.Button(row2, text=text, width=13, command=fn).pack(side="left", padx=(0, 6))
        ttk.Label(dlg, foreground="#666", justify="left",
                  text="分类 = Mod 根目录下的文件夹；子分类 = 分类里的下一级目录（如 皮肤\\纹身）。\n"
                       "删除只允许空目录（会移入回收站）；排序会影响下拉框和 Excel 工作表顺序。").pack(
            side="bottom", fill="x", padx=10, pady=(0, 4))
        dlg.geometry("660x500")
        place_dialog(dlg)
        load()

    def do_settings():
        dlg = tk.Toplevel(win)
        dlg.title("路径与浏览器设置")
        dlg.transient(win)
        dlg.grab_set()
        v_root2 = tk.StringVar(value=cfg["root"])
        v_excel2 = tk.StringVar(value=cfg["excel"])
        v_thumb2 = tk.StringVar(value=str(cfg.get("thumb_width", 800)))
        v_dl2 = tk.StringVar(value=resolve_dirs(cfg)[0])
        v_ib2 = tk.StringVar(value=cfg.get("inbox_dir") or resolve_dirs(cfg)[1])
        browsers = detect_browsers()
        names = [n for n, _ in browsers] + ["（自定义 / 手动选）"]
        path_of = dict(browsers)
        v_bname = tk.StringVar(value=next((n for n, p in browsers
                                           if p == (cfg.get("browser_path") or "")),
                                          names[0]))
        v_brow = tk.StringVar(value=cfg.get("browser_path") or (browsers[0][1] if browsers else ""))
        v_bdir = tk.StringVar(value=cfg.get("browser_dir") or "")
        v_inst = tk.StringVar(value=cfg.get("install_dir") or "")

        def pick_browser(*_):
            if v_bname.get() in path_of:
                v_brow.set(path_of[v_bname.get()])

        def row(i, label, var, browse=None):
            ttk.Label(dlg, text=label).grid(row=i, column=0, sticky="e", padx=8, pady=5)
            ttk.Entry(dlg, textvariable=var, width=62).grid(row=i, column=1, sticky="we", padx=8)
            if browse is not None:
                ttk.Button(dlg, text="浏览…", command=browse).grid(row=i, column=2, padx=8)

        row(0, "Mod 根目录", v_root2,
            lambda: v_root2.set(filedialog.askdirectory(parent=dlg) or v_root2.get()))
        row(1, "Excel 输出", v_excel2,
            lambda: v_excel2.set(filedialog.asksaveasfilename(parent=dlg, defaultextension=".xlsx",
                                                              filetypes=[("Excel", "*.xlsx")])
                                 or v_excel2.get()))
        row(2, "预览图宽度(px)", v_thumb2)
        row(3, "下载目录", v_dl2,
            lambda: v_dl2.set(filedialog.askdirectory(parent=dlg) or v_dl2.get()))
        row(4, "暂存目录", v_ib2,
            lambda: v_ib2.set(filedialog.askdirectory(parent=dlg) or v_ib2.get()))

        ttk.Label(dlg, text="内置浏览器").grid(row=5, column=0, sticky="e", padx=8, pady=5)
        cb = ttk.Combobox(dlg, textvariable=v_bname, values=names, width=59, state="readonly")
        cb.grid(row=5, column=1, sticky="we", padx=8)
        cb.bind("<<ComboboxSelected>>", pick_browser)
        row(6, "浏览器程序", v_brow,
            lambda: v_brow.set(filedialog.askopenfilename(parent=dlg, title="选择浏览器程序",
                                                          initialdir=r"C:\Program Files",
                                                          filetypes=[("可执行文件", "*.exe")])
                               or v_brow.get()))
        row(7, "Mod 安装目录（Penumbra 的 Mods）", v_inst,
            lambda: v_inst.set(filedialog.askdirectory(parent=dlg) or v_inst.get()))
        row(8, "浏览器配置目录（留空=工具专用）", v_bdir,
            lambda: v_bdir.set(filedialog.askdirectory(parent=dlg) or v_bdir.get()))
        ttk.Label(dlg, foreground="#666", wraplength=560, justify="left",
                  text="提示：工具用的是上面这个「配置目录」，和你日常浏览器的配置互不影响，"
                       "第一次需要在弹出的浏览器窗口里登录一次。\n"
                       "想复用日常登录：把这里改成你浏览器的 User Data 目录即可，"
                       "但用之前要先完全退出那个浏览器。\n"
                       "「Mod 安装目录」填 Penumbra 的 Mods 目录（默认在 "
                       "%APPDATA%\\XIVLauncher\\pluginConfigs\\Penumbra\\Mods），"
                       "填了才能用「安装检查」。").grid(row=9, column=0, columnspan=3,
                                                                sticky="w", padx=10, pady=(2, 0))
        dlg.columnconfigure(1, weight=1)

        def ok():
            cfg["root"] = v_root2.get().strip()
            cfg["excel"] = v_excel2.get().strip()
            try:
                cfg["thumb_width"] = int(v_thumb2.get().strip() or 800)
            except ValueError:
                pass
            cfg["download_dir"] = v_dl2.get().strip()
            cfg["inbox_dir"] = v_ib2.get().strip()
            cfg["browser_path"] = v_brow.get().strip()
            cfg["browser_dir"] = v_bdir.get().strip()
            cfg["install_dir"] = v_inst.get().strip()
            save_config(cfg)
            refresh()
            update_paths_label()
            dlg.destroy()
            status("设置已保存")

        bf = ttk.Frame(dlg)
        bf.grid(row=10, column=0, columnspan=3, pady=10)
        ttk.Button(bf, text="保存", width=12, command=ok).pack(side="left", padx=8)
        ttk.Button(bf, text="取消", width=12, command=dlg.destroy).pack(side="left", padx=8)
        dlg.bind("<Escape>", lambda *_: dlg.destroy())
        place_dialog(dlg)

    def do_fix_cover():
        """给选中的 Mod 补预览图：打开它的 Mod 页面抓封面，存成 <文件夹名>.jpg"""
        m = selected_mod()
        if not m:
            messagebox.showinfo("提示", "请先在左边选一条 Mod")
            return
        addr = (m.get("addr") or "").strip()
        got = re.search(r"/modid/(\d+)", addr)
        if not got:
            messagebox.showinfo(
                "提示",
                "这条 Mod 没有 xivmodarchive 的 Mod 地址，没法自动补封面。\n\n"
                "也可以手动放一张图进它的文件夹，命名成：\n%s.jpg" % Path(m["folder"]).name)
            return
        if has_preview(m["folder"], m["author"], m["name"]):
            messagebox.showinfo("提示", "这个 Mod 已经有预览图了。\n想换一张的话，"
                                        "先把文件夹里那张同名图片删掉再点这里。")
            return
        if not browser_running(cfg):
            if not messagebox.askyesno("需要内置浏览器",
                                       "补封面要用内置浏览器打开这个 Mod 页面。\n现在打开吗？"):
                return
            try:
                launch_browser(cfg, url=addr)
            except SystemExit as e:
                messagebox.showerror("打不开浏览器", str(e))
                return

        def work():
            p2 = ensure_preview(cfg, m["folder"], m["author"], m["name"], "", m["addr"])
            if not p2:
                raise SystemExit("没能抓到封面：这个 Mod 页面上可能没有封面图。")
            cmd_scan(cfg, quiet=True)
            cmd_export(cfg)

        run_task("补预览图", work,
                 lambda: (refresh(), messagebox.showinfo(
                     "完成", "封面已保存为该 Mod 的预览图，并重新生成了 Excel。")))

    # ------------------------------------------------------- 打包备份 / 导入备份
    def _backup_size_hint(v_mods, v_meta, est):
        """后台估算体积，更新提示"""
        def work():
            try:
                _, s_mod = collect_backup_items(cfg, True, False)
                _, s_meta = collect_backup_items(cfg, False, True)
            except Exception:
                return

            def show():
                if not win.winfo_exists():
                    return
                est.set("Mod 文件夹 %d 个 Mod / %d 个文件 / %s　｜　索引库+汇总表 %s"
                        % (s_mod["mods"], s_mod["files"], fmt_size(s_mod["bytes"]),
                           fmt_size(s_meta["bytes"])))
            win.after(0, show)

        import threading
        threading.Thread(target=work, daemon=True).start()

    def do_backup():
        dlg = tk.Toplevel(win)
        dlg.title("打包备份")
        dlg.transient(win)
        v_mods = tk.BooleanVar(value=True)
        v_meta = tk.BooleanVar(value=True)
        v_zip = tk.BooleanVar(value=False)
        v_out = tk.StringVar(value=str(default_backup_path(cfg)))
        est = tk.StringVar(value="正在统计体积…")
        msg = tk.StringVar(value="")
        # 这些变量只被后台线程/闭包引用，必须挂到窗口上，否则会被 GC 掉、控件变空
        dlg._keep = (v_mods, v_meta, v_zip, v_out, est, msg)

        # 底部按钮先摆
        bar = ttk.Frame(dlg)
        bar.pack(side="bottom", fill="x", padx=10, pady=(0, 10))
        ttk.Button(bar, text="关闭", width=10, command=dlg.destroy).pack(side="right")
        ttk.Button(bar, text="打开备份文件夹", command=lambda: open_path(
            resolve_backup_dir(cfg))).pack(side="right", padx=6)
        b_go = ttk.Button(bar, text="开始备份", width=12)
        b_go.pack(side="right", padx=6)
        pb = ttk.Progressbar(dlg, mode="determinate", maximum=100)
        pb.pack(side="bottom", fill="x", padx=10, pady=(0, 6))
        ttk.Label(dlg, textvariable=msg, foreground="#555").pack(side="bottom", fill="x", padx=12)

        body = ttk.Frame(dlg)
        body.pack(fill="both", expand=True, padx=12, pady=(10, 4))
        ttk.Label(body, text="备份内容", font=("Microsoft YaHei", 10, "bold")).pack(anchor="w")
        ttk.Checkbutton(body, text="Mod 文件夹（含各家预览图、地址.txt）",
                        variable=v_mods).pack(anchor="w", padx=12, pady=(4, 0))
        ttk.Checkbutton(body, text="索引库 + 汇总表 + 配置（很小，建议勾上）",
                        variable=v_meta).pack(anchor="w", padx=12)
        ttk.Checkbutton(body, text="压缩（体积小一点，但慢很多；Mod 本身多半已是压缩包）",
                        variable=v_zip).pack(anchor="w", padx=12)
        ttk.Label(body, textvariable=est, foreground="#0a6").pack(anchor="w", padx=12,
                                                                  pady=(8, 0))
        row = ttk.Frame(body)
        row.pack(fill="x", pady=(10, 0))
        ttk.Label(row, text="保存到").pack(side="left")
        ttk.Entry(row, textvariable=v_out).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row, text="浏览…", command=lambda: v_out.set(
            filedialog.asksaveasfilename(parent=dlg, title="备份保存到",
                                         initialdir=str(resolve_backup_dir(cfg)),
                                         initialfile=Path(v_out.get()).name or "Mod备份.zip",
                                         defaultextension=".zip",
                                         filetypes=[("Mod 备份", "*.zip")]) or v_out.get())
                   ).pack(side="left")
        ttk.Label(body, foreground="#666", justify="left",
                  text="备份是一整个 zip，换电脑/重装后可以直接用「导入备份」还原。").pack(
            anchor="w", padx=12, pady=(8, 0))

        run = {"on": False, "cancel": False, "n": 0, "tn": 0, "b": 0, "tb": 1, "name": ""}

        def tick():
            if not dlg.winfo_exists():
                return
            if run["on"]:
                pb["value"] = min(100.0, run["b"] * 100.0 / max(1, run["tb"]))
                msg.set("打包中 %d/%d　%s　%s / %s"
                        % (run["n"], run["tn"], run["name"][:26],
                           fmt_size(run["b"]), fmt_size(run["tb"])))
            dlg.after(150, tick)

        def start():
            if run["on"]:
                run["cancel"] = True
                msg.set("正在取消…")
                return
            out = v_out.get().strip() or str(default_backup_path(cfg))
            if not out.lower().endswith(".zip"):
                out += ".zip"
            v_out.set(out)
            if Path(out).exists() and not messagebox.askyesno(
                    "覆盖？", "这个文件已经存在，覆盖它吗？\n%s" % out, parent=dlg):
                return
            run.update({"on": True, "cancel": False, "n": 0, "tn": 0, "b": 0, "tb": 1,
                        "name": ""})
            pb["value"] = 0
            b_go.configure(text="取消")

            def prog(n, tn, b, tb, name):
                run.update({"n": n, "tn": tn, "b": b, "tb": tb or 1, "name": name or ""})

            def work():
                try:
                    s2 = make_backup(cfg, out, v_mods.get(), v_meta.get(), v_zip.get(),
                                     progress=prog, should_cancel=lambda: run["cancel"])
                except BackupCancelled:
                    win.after(0, lambda: msg.set("已取消，没有生成备份文件"))
                except Exception as e:
                    tb = traceback.format_exc()
                    log(tb)
                    win.after(0, lambda: (msg.set("失败：%s" % e),
                                          messagebox.showerror("备份失败", tb[-1200:],
                                                               parent=dlg)))
                else:
                    def ok():
                        pb["value"] = 100
                        msg.set("完成：%s（%s / %d 个文件）"
                                % (s2["path"], fmt_size(s2["zip_bytes"]), s2["files"]))
                        status("备份完成：%s" % s2["path"])
                    win.after(0, ok)
                finally:
                    def reset():
                        run["on"] = False
                        b_go.configure(text="再备份一次")
                    win.after(0, reset)

            import threading
            threading.Thread(target=work, daemon=True).start()

        b_go.configure(command=start)
        place_dialog(dlg)
        _backup_size_hint(v_mods, v_meta, est)
        dlg.after(150, tick)

    def do_restore():
        dlg = tk.Toplevel(win)
        dlg.title("导入备份")
        dlg.transient(win)
        v_file = tk.StringVar(value="")
        v_mods = tk.BooleanVar(value=True)
        v_meta = tk.BooleanVar(value=False)
        v_mode = tk.StringVar(value="skip")
        info = tk.StringVar(value="挑一个备份 zip，下面会显示里面有什么。")
        msg = tk.StringVar(value="")
        dlg._keep = (v_file, v_mods, v_meta, v_mode, info, msg)

        bar = ttk.Frame(dlg)
        bar.pack(side="bottom", fill="x", padx=10, pady=(0, 10))
        ttk.Button(bar, text="关闭", width=10, command=dlg.destroy).pack(side="right")
        ttk.Button(bar, text="打开备份文件夹", command=lambda: open_path(
            resolve_backup_dir(cfg))).pack(side="right", padx=6)
        b_go = ttk.Button(bar, text="开始恢复", width=12)
        b_go.pack(side="right", padx=6)
        pb = ttk.Progressbar(dlg, mode="determinate", maximum=100)
        pb.pack(side="bottom", fill="x", padx=10, pady=(0, 6))
        ttk.Label(dlg, textvariable=msg, foreground="#555").pack(side="bottom", fill="x", padx=12)

        body = ttk.Frame(dlg)
        body.pack(fill="both", expand=True, padx=12, pady=(10, 4))
        row = ttk.Frame(body)
        row.pack(fill="x")
        ttk.Label(row, text="备份文件").pack(side="left")
        ttk.Entry(row, textvariable=v_file).pack(side="left", fill="x", expand=True, padx=6)
        ttk.Button(row, text="浏览…", command=lambda: pick()).pack(side="left")
        ttk.Label(body, textvariable=info, foreground="#0a6", justify="left",
                  wraplength=520).pack(anchor="w", padx=2, pady=(8, 0))
        ttk.Separator(body, orient="horizontal").pack(fill="x", pady=10)
        ttk.Label(body, text="恢复内容", font=("Microsoft YaHei", 10, "bold")).pack(anchor="w")
        ttk.Checkbutton(body, text="Mod 文件夹（解压回 Mod 根目录）",
                        variable=v_mods).pack(anchor="w", padx=12, pady=(4, 0))
        ttk.Checkbutton(body, text="索引库 + 汇总表（一般不用勾，恢复完会自动重扫重建）",
                        variable=v_meta).pack(anchor="w", padx=12)
        ttk.Label(body, text="如果已经存在同名的 Mod", font=("Microsoft YaHei", 10, "bold")
                  ).pack(anchor="w", pady=(10, 2))
        for text, val in (("跳过（保持现状，最保险）", "skip"),
                          ("覆盖（旧的先移入回收站，再写回备份里的版本）", "overwrite"),
                          ("另存一份（文件夹名后加「(备份恢复)」）", "rename")):
            ttk.Radiobutton(body, text=text, value=val, variable=v_mode).pack(anchor="w", padx=12)
        ttk.Label(body, foreground="#666", justify="left", wraplength=520,
                  text="恢复 Mod 文件夹后会自动重扫索引并重新生成 Excel。").pack(
            anchor="w", pady=(8, 0))

        state = {"path": None, "run": False, "cancel": False, "n": 0, "tn": 0, "b": 0, "tb": 1,
                 "name": ""}

        def load_info(*_):
            p = v_file.get().strip()
            if not p or not Path(p).is_file():
                info.set("还没有选备份文件。")
                state["path"] = None
                return
            d = inspect_backup(p)
            state["path"] = d
            if not d["ok"]:
                info.set("读不出内容：%s" % (d["error"] or "格式不认识"))
                return
            mf = d.get("manifest") or {}
            info.set("备份时间 %s　｜　%d 个 Mod　%d 个文件　%s\n"
                     "包含：%s%s　｜　分类：%s"
                     % (mf.get("备份时间", "未知"), len(d["mods"]), d["n_files"],
                        fmt_size(d["raw_bytes"]),
                        "Mod 文件夹" if d["mods"] else "",
                        "　索引库" if d["has_db"] else "",
                        "、".join(d["cats"]) or "—"))

        def pick():
            p = filedialog.askopenfilename(parent=dlg, title="选择备份文件",
                                           initialdir=str(resolve_backup_dir(cfg)),
                                           filetypes=[("Mod 备份", "*.zip"), ("所有文件", "*.*")])
            if p:
                v_file.set(p)
                load_info()

        def tick():
            if not dlg.winfo_exists():
                return
            if state["run"]:
                pb["value"] = min(100.0, state["b"] * 100.0 / max(1, state["tb"]))
                msg.set("恢复中 %d/%d　%s　%s / %s"
                        % (state["n"], state["tn"], state["name"][:26],
                           fmt_size(state["b"]), fmt_size(state["tb"])))
            dlg.after(150, tick)

        def start():
            if state["run"]:
                state["cancel"] = True
                msg.set("正在取消…")
                return
            p = v_file.get().strip()
            if not p or not Path(p).is_file():
                messagebox.showinfo("提示", "先选一个备份 zip。", parent=dlg)
                return
            if not (v_mods.get() or v_meta.get()):
                messagebox.showinfo("提示", "「Mod 文件夹」和「索引库+汇总表」至少要勾一个。",
                                    parent=dlg)
                return
            d = state.get("path") or inspect_backup(p)
            if not d["ok"]:
                messagebox.showerror("打不开", "这个文件不是本工具的备份，或者已经损坏。\n%s"
                                     % (d["error"] or ""), parent=dlg)
                return
            if v_mods.get() and not d["mods"]:
                messagebox.showinfo("提示", "这个备份里没有 Mod 文件夹（是「只备份索引」那种）。",
                                    parent=dlg)
                return
            if v_mods.get() and v_mode.get() == "overwrite":
                if not messagebox.askyesno(
                        "确认覆盖？", "已存在的同名 Mod 会先被移入回收站，再写入备份里的版本。\n"
                        "确定要覆盖吗？", parent=dlg):
                    return
            state.update({"run": True, "cancel": False, "n": 0, "tn": 0, "b": 0, "tb": 1,
                          "name": ""})
            pb["value"] = 0
            b_go.configure(text="取消")

            def prog(n, tn, b, tb, name):
                state.update({"n": n, "tn": tn, "b": b, "tb": tb or 1, "name": name or ""})

            def work():
                try:
                    r = restore_backup(cfg, p, v_mods.get(), v_meta.get(), v_mode.get(),
                                       progress=prog, should_cancel=lambda: state["cancel"])
                except BackupCancelled:
                    win.after(0, lambda: msg.set("已取消（已经写进去的文件不会删，可以再恢复一次）"))
                except Exception as e:
                    tb = traceback.format_exc()
                    log(tb)
                    win.after(0, lambda: (msg.set("失败：%s" % e),
                                          messagebox.showerror("恢复失败", tb[-1200:],
                                                               parent=dlg)))
                else:
                    if v_mods.get() and not state["cancel"]:
                        win.after(0, lambda: msg.set("正在重扫索引并重新生成 Excel…"))
                        try:
                            cmd_scan(cfg, quiet=True)
                            cmd_export(cfg)
                        except Exception:
                            log(traceback.format_exc())

                    def ok():
                        pb["value"] = 100
                        bits = ["新增 %d" % r["mods_new"], "跳过 %d" % r["mods_skipped"],
                                "覆盖 %d" % r["mods_overwritten"], "另存 %d" % r["mods_renamed"],
                                "解出 %d 个文件" % r["files"]]
                        if r["db"]:
                            bits.append("合并索引 %d 条" % r["db_rows"])
                        if r["excel"]:
                            bits.append("恢复汇总表")
                        msg.set("完成：" + "，".join(bits))
                        status("备份已恢复：%s" % Path(p).name)
                        refresh()
                        if r["renamed_list"]:
                            messagebox.showinfo(
                                "提示", "有 %d 个 Mod 是另存进来的，名字后面带了「(备份恢复)」：\n%s"
                                % (len(r["renamed_list"]),
                                   "\n".join("  " + b for _, b in r["renamed_list"][:8])), parent=dlg)
                    win.after(0, ok)
                finally:
                    def reset():
                        state["run"] = False
                        b_go.configure(text="再恢复一次")
                    win.after(0, reset)

            import threading
            threading.Thread(target=work, daemon=True).start()

        b_go.configure(command=start)
        v_file.trace_add("write", load_info)
        place_dialog(dlg)
        load_info()
        dlg.after(150, tick)

    def do_help():
        messagebox.showinfo("使用说明", "完整说明见程序目录下的「使用说明.txt」\n\n"
                                      "常用流程：\n"
                                      "  · 加新 Mod：添加 Mod… → 填分类/类型 → 确定\n"
                                      "  · 改内容：选中后点编辑选中…\n"
                                      "  · 加完记得点「扫描并生成 Excel」\n\n"
                                      "程序目录：%s\nMod 目录：%s" % (APP_DIR, cfg["root"]))

    def do_about():
        messagebox.showinfo("关于", "FFXIV Mod 管理工具  v" + APP_VERSION + "\n\n"
                                    "程序目录：%s\n数据目录：%s\n汇总表：%s\n索引库：%s"
                                    % (APP_DIR, cfg["root"], cfg["excel"], DB_PATH))

    # ------------------------------------------------------- 动作注册表（扩展点）
    ACTIONS = (
        # (标签, 分组, 处理函数, 是否放工具栏, 悬停提示)
        ("添加 Mod…", "mod", do_add, True, "导入新 Mod：文件夹或压缩包，自动编号并写 地址.txt"),
        ("编辑选中…", "mod", do_edit, True, "改名 / 换分类 / 换 SFW·NSFW / 改子目录 / 改地址"),
        ("删除选中", "mod", do_delete, True, "删除选中的 Mod（移入回收站，可还原）"),
        ("打开所在文件夹", "mod", open_selected, False, "在资源管理器里定位这条 Mod"),
        ("复制文件夹路径", "mod", copy_path, False, "把选中 Mod 的路径复制到剪贴板"),
        ("补预览图", "mod", do_fix_cover, False, "用内置浏览器打开该 Mod 页面，抓封面做预览图"),
        ("扫描目录", "table", do_scan, True, "重新扫描 Mod 文件夹，更新索引库"),
        ("生成 Excel", "table", do_export, True, "用索引库生成 Excel（不重新扫描）"),
        ("扫描并生成 Excel", "table", do_run, True, "扫描 + 重新生成 Excel（加完新 Mod 点这个）"),
        ("检查", "table", do_check, False, "检查缺地址 / 缺预览图 / 序号重复"),
        ("分类管理…", "table", do_categories, False, "新建 / 重命名 / 排序 分类和子分类"),
        ("自动分配序号…", "table", do_renumber, False, "把某个分类的序号重排成连续的 1..N"),
        ("查重（重复 Mod）…", "table", do_dupes, False, "按 Mod 地址 / 作者+名称 找重复"),
        ("安装检查…", "table", do_installed, False, "拿安装目录里的文件夹名判断哪些已安装"),
        ("打开 Mod 根目录", "table", lambda: open_path(cfg["root"]), False, ""),
        ("打开 Excel", "table", lambda: open_path(cfg["excel"]), False, ""),
        ("下载工作台…", "web", do_workbench, True,
         "内置浏览器抓取网址/封面 + 监视下载 + 一键导入"),
        ("用内置浏览器打开网站", "web", do_browser_open, False, "打开可被工具读取的浏览器窗口"),
        ("用系统浏览器打开网站", "web", do_open_site, False, "不想用内置浏览器时的备选"),
        ("打开下载文件夹", "web", do_open_downloads, False, ""),
        ("打开暂存文件夹", "web", do_open_inbox, False, ""),
        ("打包备份…", "backup", do_backup, True,
         "把 Mod 文件夹 + 索引库 + 汇总表 打成一个 zip，换电脑能整体还原"),
        ("导入备份…", "backup", do_restore, True,
         "从备份 zip 还原 Mod / 索引库（自动重扫并重新生成 Excel）"),
        ("打开备份文件夹", "backup", lambda: open_path(resolve_backup_dir(cfg)), False, ""),
    )

    def menu_add(menu, group):
        for label, g, fn, _tb, _tt in ACTIONS:
            if g == group and fn:
                menu.add_command(label=label, command=fn)

    def fill_menu(menu, everything=False):
        menu_add(menu, "mod")
        menu.add_separator()
        menu_add(menu, "table")
        menu.add_separator()
        menu_add(menu, "backup")
        if everything:
            menu.add_separator()
            menu.add_command(label="路径设置…", command=do_settings)
            menu.add_separator()
            menu.add_command(label="退出", command=on_close)

    # ------------------------------------------------------------------ 菜单栏
    def on_close():
        try:
            cfg["window"] = win.geometry()
            save_config(cfg)
        except Exception:
            pass
        win.destroy()

    menubar = tk.Menu(win)
    m_file = tk.Menu(menubar, tearoff=0)
    m_file.add_command(label="扫描目录", command=do_scan)
    m_file.add_command(label="生成 Excel", command=do_export)
    m_file.add_command(label="扫描并生成 Excel", command=do_run)
    m_file.add_separator()
    m_file.add_command(label="打开 Mod 根目录", command=lambda: open_path(cfg["root"]))
    m_file.add_command(label="打开 Excel", command=lambda: open_path(cfg["excel"]))
    m_file.add_separator()
    m_file.add_command(label="自动分配序号…", command=do_renumber)
    m_file.add_command(label="查重（重复 Mod）…", command=do_dupes)
    m_file.add_command(label="安装检查…", command=do_installed)
    m_file.add_separator()
    m_file.add_command(label="打包备份…", command=do_backup)
    m_file.add_command(label="导入备份…", command=do_restore)
    m_file.add_separator()
    m_file.add_command(label="退出", command=on_close)
    menubar.add_cascade(label="文件", menu=m_file)

    m_mod = tk.Menu(menubar, tearoff=0)
    menu_add(m_mod, "mod")
    menubar.add_cascade(label="Mod", menu=m_mod)

    m_web = tk.Menu(menubar, tearoff=0)
    m_web.add_command(label="下载工作台…", command=do_workbench)
    m_web.add_command(label="用内置浏览器打开网站", command=do_browser_open)
    m_web.add_command(label="用系统浏览器打开网站", command=do_open_site)
    m_web.add_separator()
    menu_add(m_web, "web")
    menubar.add_cascade(label="网站", menu=m_web)

    m_view = tk.Menu(menubar, tearoff=0)
    m_view.add_command(label="刷新列表", command=refresh)
    m_view.add_separator()
    for label, val in (("全部", "全部"), ("只看 SFW", "SFW"), ("只看 NSFW", "NSFW")):
        m_view.add_radiobutton(label=label, variable=v_zone, value=val, command=lambda: refresh())
    menubar.add_cascade(label="视图", menu=m_view)

    m_bak = tk.Menu(menubar, tearoff=0)
    m_bak.add_command(label="打包备份…", command=do_backup)
    m_bak.add_command(label="导入备份…", command=do_restore)
    m_bak.add_separator()
    m_bak.add_command(label="打开备份文件夹",
                      command=lambda: open_path(resolve_backup_dir(cfg)))
    menubar.add_cascade(label="备份", menu=m_bak)

    m_set = tk.Menu(menubar, tearoff=0)
    m_set.add_command(label="路径设置…", command=do_settings)
    m_set.add_command(label="分类管理…", command=do_categories)
    m_set.add_separator()
    m_set.add_command(label="打开配置文件", command=lambda: open_path(CONFIG_PATH))
    m_set.add_command(label="打开日志文件", command=lambda: open_path(LOG_PATH))
    m_set.add_command(label="打开程序目录", command=lambda: open_path(APP_DIR))
    menubar.add_cascade(label="设置", menu=m_set)

    m_help = tk.Menu(menubar, tearoff=0)
    m_help.add_command(label="使用说明", command=do_help)
    m_help.add_command(label="关于", command=do_about)
    menubar.add_cascade(label="帮助", menu=m_help)
    win.config(menu=menubar)

    # ------------------------------------------------------- 工具栏（单排）
    bar = ttk.Frame(win)
    bar.pack(fill="x", padx=8, pady=(6, 2))
    for label, group, fn, in_tb, ttext in ACTIONS:
        if in_tb:
            b = ttk.Button(bar, text=label, command=fn)
            b.pack(side="left", padx=3)
            tip(b, ttext)
    mb = ttk.Menubutton(bar, text="更多 ▾")
    mb.pack(side="left", padx=(10, 3))
    more = tk.Menu(mb, tearoff=0)
    fill_menu(more, everything=True)
    mb["menu"] = more
    tip(mb, "其余功能都在这里（和菜单栏一致）")

    # --------------------------------------------------------------- 路径行
    path_row = ttk.Frame(win)
    path_row.pack(fill="x", padx=8, pady=(0, 4))
    root_lbl = ttk.Label(path_row, text="", foreground="#555")
    root_lbl.pack(side="left")
    excel_lbl = ttk.Label(path_row, text="", foreground="#555")
    excel_lbl.pack(side="left", padx=(12, 0))
    ttk.Button(path_row, text="设置…", width=8, command=do_settings).pack(side="left", padx=8)

    def update_paths_label():
        root_lbl.configure(text="Mod 目录：%s" % brief(cfg["root"], 46))
        excel_lbl.configure(text="Excel：%s" % brief(cfg["excel"], 40))
        tip(root_lbl, cfg["root"])
        tip(excel_lbl, cfg["excel"])

    # ------------------------------------------------------------- 状态栏
    # 铺满窗口最底部：左边是运行状态，右边是条数统计
    sbar = tk.Frame(win, bd=1, relief="sunken")
    sbar.pack(fill="x", side="bottom")
    ttk.Label(sbar, textvariable=status_var, anchor="w")\
        .pack(side="left", fill="x", expand=True, padx=8, pady=2)
    ttk.Label(sbar, textvariable=count_var, anchor="e", foreground="#444")\
        .pack(side="right", padx=8, pady=2)

    # ------------------------------------------------------------- 主体分栏
    paned = ttk.Panedwindow(win, orient="horizontal")
    paned.pack(fill="both", expand=True, padx=8, pady=4)

    left = ttk.Frame(paned)
    paned.add(left, weight=3)
    right = ttk.Frame(paned)
    paned.add(right, weight=2)

    # 左侧：筛选条
    fbar = ttk.Frame(left)
    fbar.pack(fill="x")
    ttk.Label(fbar, text="分类").pack(side="left")
    cat_box = ttk.Combobox(fbar, textvariable=v_cat, width=10, state="readonly")
    cat_box.pack(side="left", padx=(4, 6))
    ttk.Label(fbar, text="子分类").pack(side="left")
    sub_box = ttk.Combobox(fbar, textvariable=v_sub, width=9, state="readonly")
    sub_box.pack(side="left", padx=(4, 10))
    ttk.Label(fbar, text="类型").pack(side="left")
    zone_box = ttk.Combobox(fbar, textvariable=v_zone, width=7, state="readonly",
                            values=["全部", "SFW", "NSFW"])
    zone_box.pack(side="left", padx=(4, 10))
    ttk.Label(fbar, text="搜索").pack(side="left")
    search_entry = ttk.Entry(fbar, textvariable=v_search)
    search_entry.pack(side="left", fill="x", expand=True, padx=4)
    ttk.Button(fbar, text="清空", width=6,
               command=lambda: (v_search.set(""), v_cat.set("全部分类"),
                                v_sub.set("全部子分类"), v_zone.set("全部")))\
        .pack(side="left", padx=(6, 0))

    # 左侧：列表
    cols = ("cat", "sub", "seq", "author", "nsfw", "name", "inst")
    tree = ttk.Treeview(left, columns=cols, show="headings", selectmode="browse")
    for c, t, w in (("cat", "分类", 62), ("sub", "子分类", 70), ("seq", "序号", 44),
                    ("author", "作者", 92), ("nsfw", "NSFW/SFW", 78),
                    ("name", "Mod名称", 300), ("inst", "是否安装", 76)):
        tree.heading(c, text=t)
        tree.column(c, width=w, stretch=(c == "name"),
                    anchor="w" if c in ("author", "name") else "center")
    tree.tag_configure("inst_yes", background="#D8F5D8")   # 已安装 = 淡绿
    tree.tag_configure("inst_no", background="#FFFFFF")    # 未安装 = 白
    ysb = ttk.Scrollbar(left, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=ysb.set)
    tree.pack(side="left", fill="both", expand=True, pady=(4, 0))
    ysb.pack(side="right", fill="y", pady=(4, 0))

    # 右侧：预览 + 详情 + 快捷操作
    img_label = ttk.Label(right, text="← 选中左侧一条查看预览", anchor="center")
    img_label.pack(fill="both", expand=True, padx=6, pady=(4, 6))

    df = ttk.LabelFrame(right, text="详情")
    df.pack(fill="x", padx=6)
    detail_vars = {}
    order = ("分类", "序号", "作者", "类型", "影响/替换", "Mod 名称", "Mod 地址", "文件夹", "预览图")
    for i, k in enumerate(order):
        ttk.Label(df, text=k, foreground="#666").grid(row=i, column=0, sticky="e", padx=(6, 4), pady=1)
        var = tk.StringVar()
        detail_vars[k] = var
        lbl = ttk.Label(df, textvariable=var, wraplength=330, justify="left")
        lbl.grid(row=i, column=1, sticky="w", padx=(0, 6), pady=1)
        if k == "Mod 地址":
            lbl.configure(foreground="#0645ad", cursor="hand2")
            lbl.bind("<Button-1>", open_addr)
            tip(lbl, "点一下用浏览器打开")
        elif k == "文件夹":
            lbl.configure(foreground="#0645ad", cursor="hand2")
            lbl.bind("<Button-1>", open_folder)
            tip(lbl, "点一下在资源管理器里打开")
    df.columnconfigure(1, weight=1)

    quick = ttk.Frame(right)
    quick.pack(fill="x", padx=6, pady=6)
    for label, fn, ttext in (("编辑…", do_edit, "改名 / 换分类 / 换类型 / 改地址"),
                             ("打开文件夹", open_selected, "在资源管理器里定位"),
                             ("复制路径", copy_path, "复制文件夹路径"),
                             ("补预览图", do_fix_cover, "用内置浏览器打开该 Mod 页面抓封面"),
                             ("删除", do_delete, "移入回收站")):
        b = ttk.Button(quick, text=label, command=fn)
        b.pack(side="left", padx=(0, 4))
        tip(b, ttext)

    # ------------------------------------------------------------ 事件绑定
    def popup(event):
        row = tree.identify_row(event.y)
        if row:
            tree.selection_set(row)
            on_select()
        menu = tk.Menu(win, tearoff=0)
        menu.add_command(label="编辑（改名 / 分类 / 类型 / 地址）…", command=do_edit)
        menu.add_command(label="打开所在文件夹", command=open_selected)
        menu.add_command(label="复制文件夹路径", command=copy_path)
        menu.add_command(label="补预览图（按地址抓封面）", command=do_fix_cover)
        menu.add_command(label="自动分配序号…", command=do_renumber)
        menu.add_command(label="安装检查…", command=do_installed)
        menu.add_separator()
        menu.add_command(label="删除（移入回收站）", command=do_delete)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()

    tree.bind("<<TreeviewSelect>>", on_select)
    tree.bind("<Button-3>", popup)
    tree.bind("<Double-1>", lambda *_: open_selected())
    v_search.trace_add("write", lambda *_: (fill_tree(), update_counts()))
    cat_box.bind("<<ComboboxSelected>>", lambda *_: (refresh(), ))
    sub_box.bind("<<ComboboxSelected>>", lambda *_: (fill_tree(), update_counts()))
    zone_box.bind("<<ComboboxSelected>>", lambda *_: (fill_tree(), update_counts()))
    win.bind("<F5>", lambda *_: do_scan())
    win.bind("<Control-n>", lambda *_: do_add())
    win.bind("<Delete>", lambda *_: do_delete())
    win.protocol("WM_DELETE_WINDOW", on_close)

    update_paths_label()
    refresh()
    on_select()
    win.mainloop()


# ----------------------------------------------------------------------- main
def main(argv=None):
    cfg = load_config()
    ap = argparse.ArgumentParser(description="FFXIV Mod 管理工具")
    ap.add_argument("command", nargs="?", default="gui",
                    choices=["gui", "scan", "export", "run", "check", "stats", "add", "edit", "delete",
                             "browser", "dupes", "installed", "renumber", "backup", "restore",
                             "backups"])
    ap.add_argument("--root", help="Mod 根目录（覆盖配置）")
    ap.add_argument("--excel", help="Excel 输出路径（覆盖配置）")
    ap.add_argument("--thumb-width", type=int, help="缩略图宽度，像素")
    ap.add_argument("--no-images", action="store_true", help="不内嵌预览图")
    ap.add_argument("--open", action="store_true", help="完成后打开 Excel")
    ap.add_argument("--src", help="add：要导入的文件夹或压缩包")
    ap.add_argument("--category", help="add/edit：分类（衣服/饰品/武器/皮肤…）")
    ap.add_argument("--zone", choices=["SFW", "NSFW"], default=None, help="add/edit：类型（add 默认 SFW；edit 不填=保持不变）")
    ap.add_argument("--subdir", default=None, help="add/edit：子分类目录，如 纹身")
    ap.add_argument("--author", help="add/edit：作者")
    ap.add_argument("--name", help="add/edit：Mod名称")
    ap.add_argument("--seq", type=int, help="add/edit：序号（add 默认自动取最大+1）")
    ap.add_argument("--addr", default=None, help="add/edit：Mod地址（会写入 地址.txt；edit 传空串可清除）")
    ap.add_argument("--move", action="store_true", help="add：移动而不是复制")
    ap.add_argument("--target", help="delete：要删除的 Mod 文件夹（默认用 --src）")
    ap.add_argument("--permanent", action="store_true", help="delete：直接删除，不进回收站")
    ap.add_argument("--yes", action="store_true", help="delete/renumber：跳过确认")
    ap.add_argument("--start", type=int, default=1, help="renumber：起始序号（默认 1）")
    ap.add_argument("--out", help="backup：备份 zip 的保存路径（默认放到 程序目录\\backup）")
    ap.add_argument("--only", choices=["all", "meta"],
                    help="backup/restore：all=含 Mod 文件夹（默认）；meta=只处理索引库+汇总表")
    ap.add_argument("--compress", action="store_true", help="backup：压缩（慢，体积略小）")
    ap.add_argument("--with-index", action="store_true",
                    help="restore：连索引库+汇总表一起恢复（默认只恢复 Mod 文件夹后重扫）")
    ap.add_argument("--mode", choices=["skip", "overwrite", "rename"], default="skip",
                    help="restore：同名 Mod 的处理方式（默认 skip 跳过）")
    a = ap.parse_args(argv)

    if a.root:
        cfg["root"] = a.root
    if a.excel:
        cfg["excel"] = a.excel
    if a.thumb_width:
        cfg["thumb_width"] = a.thumb_width
    if a.no_images:
        cfg["embed_images"] = False
    # 命令行的临时覆盖不写回配置（配置由界面上的「保存路径」写入）
    if a.command != "gui" and not (cfg.get("root") and Path(cfg["root"]).is_dir()):
        raise SystemExit("还没有设置 Mod 根目录。\n"
                         "请先打开图形界面完成首次设置，或用 --root \"<Mod 根目录>\" 指定。")

    if a.command == "gui":
        gui(cfg)
    elif a.command == "scan":
        cmd_scan(cfg)
    elif a.command == "export":
        cmd_export(cfg)
    elif a.command == "run":
        cmd_scan(cfg, quiet=True)
        out = cmd_export(cfg)
        if a.open:
            open_path(out)
    elif a.command == "check":
        cmd_check(cfg)
    elif a.command == "stats":
        cmd_stats(cfg)
    elif a.command == "add":
        if not a.src or not a.category:
            raise SystemExit("用法： run.py add --src \"<文件夹或压缩包>\" --category <分类> "
                             "[--zone SFW|NSFW] [--subdir 纹身] [--author 作者] [--name 名称] "
                             "[--addr <Mod地址>] [--move]")
        t = cmd_add(cfg, a.src, a.category, a.zone or "SFW", a.subdir or "", a.author, a.name,
                    a.seq, a.addr or "", a.move, rescan=True, export=True)
        log("目标文件夹： %s" % t)
        if a.open:
            open_path(Path(cfg["excel"]))
    elif a.command == "dupes":
        mods = Store().all()
        groups = find_duplicates(mods)
        if not groups:
            log("没有发现重复 Mod（按 Mod 地址 / 作者+名称 判断）")
        else:
            log("发现 %d 组可能重复：" % len(groups))
            for why, info, g in groups:
                log("  [%s] %s" % (why, info))
                for m in g:
                    log("      %s（序号 %s）" % (os.path.relpath(m["folder"], cfg["root"]), m["seq"]))

    elif a.command == "installed":
        mods = Store().all()
        names, raw, d = list_installed(cfg)
        log("安装目录： %s" % (d or "(没设置/没找到，可在设置里指定 Penumbra 的 Mods 目录)"))
        if d and Path(d).is_dir():
            ins = [m for m in mods if installed_match(m, names, raw)]
            log("目录里有 %d 个文件夹；已安装 %d / 未安装 %d（共 %d）"
                % (len(raw), len(ins), len(mods) - len(ins), len(mods)))
            for m in mods:
                if not installed_match(m, names, raw):
                    log("    未安装： %s（%s）" % (m["name"], m["author"] or "-"))

    elif a.command == "backup":
        out = a.out or str(default_backup_path(cfg))
        inc_mods = (a.only != "meta")

        def _bp(n, tn, b, tb, nm):
            step = max(1, (tn or 1) // 10)
            if n == tn or n % step == 0:
                log("  打包 %d/%d  %s" % (n, tn, fmt_size(b)))

        s2 = make_backup(cfg, out, inc_mods, True, a.compress, progress=_bp)
        log("备份完成： %s" % s2["path"])
        log("  Mod %d 个 ｜ 文件 %d 个 ｜ 原始 %s ｜ 备份 %s"
            % (s2["mods"], s2["files"], fmt_size(s2["bytes"]), fmt_size(s2["zip_bytes"])))

    elif a.command == "restore":
        if not a.src:
            raise SystemExit("用法： run.py restore --src \"<备份.zip>\" "
                             "[--only meta] [--with-index] [--mode skip|overwrite|rename] [--yes]")
        d = inspect_backup(a.src)
        if not d["ok"]:
            raise SystemExit("这个文件不是本工具的备份，或者已损坏：%s" % (d["error"] or ""))
        inc_mods = (a.only != "meta")
        inc_meta = (a.only == "meta") or a.with_index
        log("备份：%s ｜ %d 个 Mod ｜ %d 个文件 ｜ %s"
            % (Path(a.src).name, len(d["mods"]), d["n_files"], fmt_size(d["raw_bytes"])))
        if inc_mods and a.mode == "overwrite" and not a.yes:
            log("将覆盖已存在的同名 Mod（旧的先移入回收站）。确认请加 --yes 再执行。")
        else:
            def _rp(n, tn, b, tb, nm):
                step = max(1, (tn or 1) // 10)
                if n == tn or n % step == 0:
                    log("  恢复 %d/%d  %s" % (n, tn, fmt_size(b)))

            r = restore_backup(cfg, a.src, inc_mods, inc_meta, a.mode,
                               progress=_rp, should_cancel=None)
            log("恢复完成： 新增 %d ｜ 跳过 %d ｜ 覆盖 %d ｜ 另存 %d ｜ 解出 %d 个文件"
                % (r["mods_new"], r["mods_skipped"], r["mods_overwritten"],
                   r["mods_renamed"], r["files"]))
            for _old, _new in r["renamed_list"]:
                log("  另存为： %s" % _new)
            if r["db"]:
                log("  合并索引库 %d 条" % r["db_rows"])
            if r["excel"]:
                log("  已恢复汇总表")
            if inc_mods:
                cmd_scan(cfg, quiet=True)
                cmd_export(cfg)
                log("  已重新扫描并重新生成 Excel")

    elif a.command == "backups":
        d0 = resolve_backup_dir(cfg)
        log("备份目录： %s" % d0)
        if d0.is_dir():
            rows = sorted(d0.glob("Mod备份_*.zip"), key=lambda q: q.stat().st_mtime, reverse=True)
            if not rows:
                log("  （还是空的）")
            for q in rows:
                mt = _dt.datetime.fromtimestamp(q.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
                log("  %s  %s  %s" % (mt, fmt_size(q.stat().st_size), q.name))
        else:
            log("  （目录还没建，第一次备份时会自动建）")

    elif a.command == "renumber":
        if not a.category:
            raise SystemExit("用法： run.py renumber --category <分类> [--subdir 子分类] "
                             "[--start 1] [--yes]")
        plan = renumber(cfg, a.category, a.subdir or "", a.start or 1, dry_run=True)
        if not plan:
            log("「%s」的序号已经是连续的，不用改。" % a.category)
        elif not a.yes:
            log("将要改 %d 个（确认后加 --yes 执行）：" % len(plan))
            for m, n in plan:
                log("   %s -> %s   %s" % (m["seq"], n, m["name"]))
        else:
            res = renumber(cfg, a.category, a.subdir or "", a.start or 1,
                           dry_run=False, export=True)
            log("已重排 %d 个。" % len(res or []))

    elif a.command == "browser":
        port = launch_browser(cfg)
        log("内置浏览器已就绪（调试端口 %d）" % port)
        log("使用的浏览器： %s" % find_browser(cfg))
        try:
            t, u = browser_wait_ready(cfg, timeout=45,
                                      on_tick=lambda x: log("  等待页面就绪… %ds" % x) if x and x % 15 == 0 else None)
            log("页面： %s" % u)
            info = browser_capture(cfg)
            log("当前页面：%s" % info.get("url"))
            log("modid：%s ｜ 名称：%s ｜ 作者：%s" % (info.get("modid") or "-",
                                                     info.get("name") or "-",
                                                     info.get("author") or "-"))
            if info.get("cover"):
                dest = APP_DIR / "cover_cache" / ("mod_%s.jpg" % (info.get("modid") or "tmp"))
                n = browser_fetch(cfg, info["cover"], dest)
                log("封面已抓取：%s（%.0f KB）" % (dest, n / 1024))
        except SystemExit as e:
            log("提示：%s" % e)
    elif a.command == "edit":
        target = a.target or a.src
        if not target:
            raise SystemExit("用法： run.py edit --target \"<Mod 文件夹>\" [--category 分类] "
                             "[--zone SFW|NSFW] [--subdir 纹身] [--author 作者] [--name 名称] "
                             "[--seq N] [--addr <Mod地址>]")
        t = cmd_edit(cfg, target, a.category, a.zone, a.subdir, a.author, a.name,
                     a.seq, a.addr, rescan=True, export=True)
        log("现在的文件夹： %s" % t)
        if a.open:
            open_path(Path(cfg["excel"]))
    elif a.command == "delete":
        target = a.target or a.src
        if not target:
            raise SystemExit("用法： run.py delete --target \"<Mod 文件夹>\" [--permanent] [--yes]")
        if not a.yes:
            log("将要删除（默认移入回收站）： %s" % target)
            log("确认无误请加 --yes 重新执行。")
        else:
            cmd_delete(cfg, target, to_recycle=not a.permanent, rescan=True, export=True)
            if a.open:
                open_path(Path(cfg["excel"]))


if __name__ == "__main__":
    main()
